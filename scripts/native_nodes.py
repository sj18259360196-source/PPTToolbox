"""Explicit node/control-point edits within a frozen native path frame."""
from __future__ import annotations
import copy
import math
from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn
try:
    from .native_capabilities import obj,number
except ImportError:
    from native_capabilities import obj,number

PAIR={'type':'array','items':number(-1_000_000,1_000_000),'minItems':2,'maxItems':2}
NODES=obj({'op':{'const':'native.nodes'},'id':{'type':'string','minLength':1},
           'max_displacement':number(.001,32),
           'edits':{'type':'array','minItems':1,'maxItems':32,'items':obj({
               'command_index':{'type':'integer','minimum':0,'maximum':511},
               'point_index':{'type':'integer','minimum':0,'maximum':2},
               'expected':PAIR,'delta':PAIR},('command_index','point_index','expected','delta'))}},
          ('op','id','max_displacement','edits'))

def topology(commands):
    from graphics_geometry import check_commands,flatten
    from shapely.geometry import Polygon
    check_commands(commands)
    rings=flatten(commands)
    polys=[Polygon(p) for p in rings]
    if any(not p.is_valid or p.area<.0001 for p in polys):raise ValueError('Invalid, self-intersecting or collapsed ring')
    for i,p in enumerate(polys):
        for q in polys[i+1:]:
            if p.boundary.intersects(q.boundary):raise ValueError('Touching or crossing compound rings')
    depths=[sum(q.contains(p) for q in polys if q is not p) for p in polys]
    return {'rings':len(polys),'holes':sum(d%2 for d in depths),'nesting':depths,
            'winding':[bool(p.exterior.is_ccw) for p in polys]}

def summary(pptx):
    from local_edit import shapes_by_name
    result={}
    for (si,name),(shape,parent) in shapes_by_name(Presentation(pptx)).items():
        paths=shape._element.xpath('./p:spPr/a:custGeom/a:pathLst/a:path')
        if not paths:continue
        rows=[]
        for path in paths:
            rows.append({'width':path.get('w'),'height':path.get('h'),
                         'commands':[{'op':etree.QName(n).localname,
                                      'points':[[p.get('x'),p.get('y')] for p in n]} for n in path]})
        result[f'{si+1}/{name}']={'parent':parent,'paths':rows,
             'path_count':len(rows),'closed_subpaths':sum(n.tag==qn('a:close') for p in paths for n in p),
             'hole_count':'requires_geometry_analysis','coordinate_space':'native path local integers'}
        try:
            from native_topology import _path
            t=topology(_path(shape,dict(width=1,height=1,width_pt=1,height_pt=1)))
            result[f'{si+1}/{name}'].update(hole_count=t['holes'],ring_count=t['rings'])
        except (ValueError,TypeError,AttributeError):
            result[f'{si+1}/{name}']['geometry_analysis']='unsupported'
    return result

def apply(obj,shape,change,canvas,native,si,parent):
    from jsonschema import Draft202012Validator
    from native_topology import _path
    from build_pptx import make_path
    Draft202012Validator(NODES).validate(change)
    if obj['kind']!='path' or not obj.get('closed'):raise ValueError('Closed native path required')
    if shape.rotation or shape._element.xpath('./p:spPr/a:xfrm[@flipH="1" or @flipV="1" or @flipH="true" or @flipV="true"]'):
        raise ValueError('Rotated or flipped path unsupported')
    ancestor=parent
    while ancestor:
        g,next_parent=native[si,ancestor]
        xf=g._element.find(qn('p:grpSpPr')).find(qn('a:xfrm'))
        if (g.rotation or xf.get('flipH') in {'1','true'} or xf.get('flipV') in {'1','true'} or
            dict(xf.find(qn('a:off')).attrib)!=dict(xf.find(qn('a:chOff')).attrib) or
            dict(xf.find(qn('a:ext')).attrib)!=dict(xf.find(qn('a:chExt')).attrib)):
            raise ValueError('Transformed group requires explicit coordinate conversion; unchanged group supported')
        ancestor=next_parent
    observed=_path(shape,canvas)
    before=obj['commands']
    if len(observed)!=len(before) or any(a[0]!=b[0] or len(a)!=len(b) or any(abs(x-y)>.003 for x,y in zip(a[1:],b[1:])) for a,b in zip(observed,before)):
        raise ValueError('Scene path and persisted native path differ')
    after=copy.deepcopy(obj);seen=set()
    for edit in change['edits']:
        key=(edit['command_index'],edit['point_index'])
        if key in seen:raise ValueError('Duplicate node selector')
        seen.add(key);i,j=key
        if i>=len(before) or 1+2*j+1>=len(before[i]):raise ValueError('Node selector outside command')
        pos=1+2*j
        if any(abs(x-y)>1e-8 for x,y in zip(before[i][pos:pos+2],edit['expected'])):
            raise ValueError('Stale node expected coordinate')
        if math.hypot(*edit['delta'])>change['max_displacement']:raise ValueError('Node displacement exceeds declared bound')
        after['commands'][i][pos:pos+2]=[x+y for x,y in zip(edit['expected'],edit['delta'])]
    old_top,new_top=topology(before),topology(after['commands'])
    if old_top!=new_top:raise ValueError('Node edit changes ring, hole, nesting or winding topology')
    sx=canvas['width']/canvas['width_pt'];sy=canvas['height']/canvas['height_pt']
    box=[shape.left/12700*sx,shape.top/12700*sy,shape.width/12700*sx,shape.height/12700*sy]
    for c in after['commands']:
        for j in range(1,len(c),2):
            if not(box[0]-.003<=c[j]<=box[0]+box[2]+.003 and box[1]-.003<=c[j+1]<=box[1]+box[3]+.003):
                raise ValueError('Edited node leaves frozen path frame')
    after['native_path_frame']={'bbox':box,'word_wrap':shape.text_frame.word_wrap,
                              'auto_size':int(shape.text_frame.auto_size) if shape.text_frame.auto_size is not None else None}
    temp=Presentation();slide=temp.slides.add_slide(temp.slide_layouts[6])
    replacement=make_path(slide.shapes,after,canvas)
    old=shape._element.spPr.find(qn('a:custGeom'))
    shape._element.spPr.replace(old,copy.deepcopy(replacement._element.spPr.find(qn('a:custGeom'))))
    return after
