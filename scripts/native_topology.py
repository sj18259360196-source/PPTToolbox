"""Bounded Office topology editing on owned copies, with observed output identity."""
from __future__ import annotations
import copy
import shutil
from pathlib import Path
from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn
from common import read_json, write_json, sha256, walk_objects
from local_edit import properties, shapes_by_name, semantic_nodes
from native_capabilities import TOPOLOGY as SCHEMA

MERGE = {"union": 1, "combine": 2, "intersect": 3, "subtract": 4, "fragment": 5}


def preflight(scene_path, source, slide_id, allowed, change):
    from jsonschema import Draft202012Validator
    Draft202012Validator(SCHEMA).validate(change)
    scene = read_json(scene_path)
    index = next(i for i,s in enumerate(scene["slides"]) if s["id"] == slide_id)
    prs = Presentation(source)
    slide = prs.slides[index]
    ids, action = change["inputs"], change["action"]
    objects = {o["id"]: o for o in scene["slides"][index]["objects"]}
    native = shapes_by_name(prs)
    if not set(ids) <= allowed or not set(ids) <= objects.keys():
        raise ValueError("Topology inputs must be existing authorized top-level objects")
    if change["primary"] != ids[0] or change["style_source"] != change["primary"]:
        raise ValueError("Primary must be first ordered input and format source")
    grouped = action in {"ungroup", "move_group"}
    if (grouped and len(ids) != 1) or (not grouped and len(ids) < 2):
        raise ValueError("Group target count or multi-input operation count invalid")
    if ("delta" in change) != (action == "move_group"):
        raise ValueError("Delta is required only for move_group")
    if change["empty"] == "allow" and action not in MERGE:
        raise ValueError("Empty policy applies only to Boolean operations")
    if slide._element.xpath(".//p:timing|.//p:cxnSp"):
        raise ValueError("Slide animation/connector references cannot be safely maintained")
    if set(ids)&set(scene["slides"][index].get("review_requirements",{}).get("relationship_objects",[])):
        raise ValueError("Topology of explicit semantic relationship objects is unsupported")
    order = [s.name for s in slide.shapes]
    positions = sorted(order.index(n) for n in ids)
    if positions != list(range(min(positions), max(positions)+1)):
        raise ValueError("Unselected objects interleave topology inputs in z-order")
    touched = set()
    for name in ids:
        shape,parent = native[index,name]
        o = objects[name]
        if parent or shape.rotation or o.get("rotation_deg"):
            raise ValueError("Topology requires unrotated top-level inputs")
        if grouped:
            if shape.shape_type != 6 or o["kind"] != "group":
                raise ValueError("Explicit native group required")
            x = shape._element.find(qn("p:grpSpPr")).find(qn("a:xfrm"))
            if x.get("rot") or x.get("flipH") or x.get("flipV"):
                raise ValueError("Rotated/flipped group unsupported")
            a,b = x.find(qn("a:ext")),x.find(qn("a:chExt"))
            if dict(a.attrib) != dict(b.attrib):
                raise ValueError("Scaled group unsupported")
            if any(s.shape_type == 6 or s.rotation for s in shape.shapes):
                raise ValueError("Nested or rotated members unsupported")
        elif o["kind"] not in {"shape", "text", "path"} or shape.shape_type not in {1,5,17}:
            raise ValueError("Only native shapes, closed paths and plain text supported")
        if action=="group" and o["kind"] not in {"shape","text"}:
            raise ValueError("Grouping derived paths is outside the single-level shape/text scope")
        if action in MERGE:
            if o["kind"] == "text" or (shape.has_text_frame and shape.text):
                raise ValueError("Boolean conversion of text is forbidden")
            if o["kind"] == "path" and not o.get("closed"):
                raise ValueError("Boolean requires closed paths")
            if shape._element.xpath(".//a:effectLst/*|.//a:effectDag|.//a:sp3d|.//a:gradFill|.//a:blipFill"):
                raise ValueError("Boolean initially supports flat solid shapes only")
            if shape._element.xpath(".//a:alpha|.//a:alphaMod|.//a:alphaOff"):
                raise ValueError("Boolean alpha styles require explicit typed readback support")
        for member in walk_objects([o]):
            touched.add(member["id"])
        if shape._element.xpath(".//a:hlinkClick|.//a:hlinkMouseOver"):
            raise ValueError("Linked topology input unsupported")
    if grouped and any(v["kind"] not in {"shape","text"} for v in objects[ids[0]]["children"]):
        raise ValueError("Only simple shape/text group children are supported")
    if action=="move_group":
        from native_format import effect_margin
        props=properties(source)
        members={v["id"]:v for v in walk_objects(scene["slides"][index]["objects"])}
        def bounds(name,delta):
            x,y,w,h=props[f"{index}/{name}"]["bbox_pt"]
            pad=effect_margin(members[name].get("style",{}).get("native_format",{}))
            return x+delta[0]-pad,y+delta[1]-pad,x+w+delta[0]+pad,y+h+delta[1]+pad
        def intersects(a,b):
            return min(a[2],b[2])>max(a[0],b[0]) and min(a[3],b[3])>max(a[1],b[1])
        for name in touched:
            if members[name]["kind"]=="group":continue
            for neighbor in members.keys()-touched:
                if members[neighbor]["kind"]=="group":continue
                box=bounds(neighbor,[0,0])
                if intersects(bounds(name,change["delta"]),box):
                    raise ValueError("Group movement may occlude an unselected neighbor; separate authorization required")
    return scene, index, touched


def _owned_operation(source, dest, index, change, prefix):
    from win32com.client import Dispatch
    from runtime_env import office_guard
    if dest.exists():
        raise ValueError("Topology copy must be new")
    shutil.copyfile(source,dest)
    with office_guard():
        app=Dispatch("PowerPoint.Application")
        deck=None
        try:
            deck=app.Presentations.Open(str(dest.resolve()),0,0,0)
            shapes=deck.Slides.Item(index+1).Shapes
            names=[str(shapes.Item(i).Name) for i in range(1,shapes.Count+1)]
            if len(names)!=len(set(names)):
                raise ValueError("Duplicate stable names")
            selected=change["inputs"]
            protected=set(names)-set(selected)
            if any(n.startswith(prefix) for n in names):
                raise ValueError("Output namespace already exists")
            primary=shapes.Item(change["primary"])
            action=change["action"]
            if action=="group":
                group=shapes.Range(tuple(selected)).Group()
                group.Name=prefix+"-group"
            elif action=="ungroup":
                primary.Ungroup()
            elif action=="move_group":
                primary.Left=float(primary.Left)+change["delta"][0]
                primary.Top=float(primary.Top)+change["delta"][1]
            else:
                shapes.Range(tuple(selected)).MergeShapes(MERGE[action],primary)
            outputs=[shapes.Item(i) for i in range(1,shapes.Count+1)
                     if str(shapes.Item(i).Name) not in protected]
            if len(outputs)>32:
                deck.Save()
                raise ValueError("Topology output complexity exceeds32")
            if action in MERGE:
                for i,s in enumerate(outputs,1):
                    s.Name=prefix+"-"+str(i)
            deck.Save()
            count=len(outputs)
            if count==0 and change["empty"]=="reject":
                raise ValueError("Office succeeded with empty result; contract forbids consumption to empty")
        finally:
            if deck is not None:
                deck.Close()


def _path(shape, canvas):
    cg=shape._element.find(qn("p:spPr")).find(qn("a:custGeom"))
    if cg is None:
        raise ValueError("Boolean output is not a persisted custom path")
    commands=[]
    sx=canvas["width"]/canvas["width_pt"]/12700
    sy=canvas["height"]/canvas["height_pt"]/12700
    for path in cg.find(qn("a:pathLst")):
        if path.get("fill","norm")!="norm" or path.get("stroke","1") not in {"1","true"}:
            raise ValueError("Unsupported compound path rendering semantics")
        w,h=int(path.get("w")),int(path.get("h"))
        if min(w,h)<=0:
            raise ValueError("Invalid actual path coordinate system")
        for node in path:
            tag=etree.QName(node).localname
            if tag=="close":
                commands.append(["Z"])
                continue
            if tag not in {"moveTo","lnTo","cubicBezTo","quadBezTo"}:
                raise ValueError("Unsupported actual path command "+tag)
            points=[[(shape.left+int(p.get("x"))*shape.width/w)*sx,
                     (shape.top+int(p.get("y"))*shape.height/h)*sy] for p in node]
            if tag=="quadBezTo":
                raise ValueError("Quadratic result requires explicit conversion; no approximation")
            commands.append([{"moveTo":"M","lnTo":"L","cubicBezTo":"C"}[tag],
                             *[v for p in points for v in p]])
    if len(commands)>512 or not commands or commands[-1]!=["Z"]:
        raise ValueError("Only bounded closed actual paths supported")
    return commands


def output_scene(scene, index, source, output, change, touched, prefix):
    before=properties(source)
    after=properties(output)
    old={o["id"]:o for o in scene["slides"][index]["objects"]}
    selected=change["inputs"]
    prs=Presentation(output)
    result=[]
    mapping={}
    act=change["action"]
    canvas=scene["canvas"]
    def relocated(o):
        o=copy.deepcopy(o)
        if o["kind"]=="group":
            o["children"]=[relocated(v) for v in o["children"]]
        else:
            box=after[f"{index}/{o['id']}"]["bbox_pt"]
            if o["kind"]=="path":
                shape=shapes_by_name(prs)[index,o["id"]][0]
                o["commands"]=_path(shape,canvas)
            else:
                o["bbox"]=[box[0]*canvas["width"]/canvas["width_pt"],box[1]*canvas["height"]/canvas["height_pt"],
                           box[2]*canvas["width"]/canvas["width_pt"],box[3]*canvas["height"]/canvas["height_pt"]]
        return o
    selected_children={o["id"]:o for n in selected for o in walk_objects([old[n]])}
    for shape in prs.slides[index].shapes:
        name=shape.name
        if name in old and name not in selected:
            result.append(copy.deepcopy(old[name]))
        elif act in MERGE:
            if not name.startswith(prefix+"-"):
                raise ValueError("Unidentified actual Boolean output")
            if shape.fill.type!=1 or shape.fill.fore_color.type!=1:
                raise ValueError("Actual Boolean fill must be explicit RGB")
            line=None
            if shape.line.fill.type==1:
                if shape.line.color.type!=1:
                    raise ValueError("Actual line color is not explicit RGB")
                line=str(shape.line.color.rgb)
            style={"fill":str(shape.fill.fore_color.rgb),"line":line,
                   "line_width_pt":float(shape.line.width.pt) if shape.line.width else 0}
            result.append({"id":name,"kind":"path","closed":True,"commands":_path(shape,canvas),
                           "native_path_frame":{"bbox":[shape.left/12700*canvas["width"]/canvas["width_pt"],
                               shape.top/12700*canvas["height"]/canvas["height_pt"],
                               shape.width/12700*canvas["width"]/canvas["width_pt"],
                               shape.height/12700*canvas["height"]/canvas["height_pt"]],
                               "word_wrap":shape.text_frame.word_wrap,
                               "auto_size":int(shape.text_frame.auto_size) if shape.text_frame.auto_size is not None else None},
                           "style":style,"editability":"path",
                           "evidence":{"status":"measured","note":"Actual Office MergeShapes derived geometry; no live constraints"}})
        elif act=="group":
            if shape.shape_type!=6 or name!=prefix+"-group":
                raise ValueError("Observed Group result differs from contract")
            children=[relocated(selected_children[s.name]) for s in shape.shapes]
            result.append({"id":name,"kind":"group","children":children,"editability":"group",
                           "evidence":{"status":"measured","note":"Actual Office Group; world coordinates, no scaling"}})
        elif act=="ungroup":
            if name not in selected_children or name in selected:
                raise ValueError("Unidentified Ungroup output")
            result.append(relocated(selected_children[name]))
        else:
            result.append(relocated(old[name]))
    updated=copy.deepcopy(scene)
    updated["slides"][index]["objects"]=result
    outputs=[o["id"] for o in result if o["id"] not in old or o["id"] in selected or o["id"] in touched]
    for name in selected:
        mapping[name]=outputs
    remaining={o["id"] for o in walk_objects(result)}
    requirements=updated["slides"][index].get("review_requirements",{})
    for field,names in requirements.items():
        requirements[field]=sorted({target for name in names
                                    for target in ([name] if name in remaining else outputs)
                                    if target in remaining})
    def order(pptx):
        return [{"name":name,"parent":parent,"relative_order":i,"shape_id":shape.shape_id}
                for i,((slide,name),(shape,parent)) in enumerate(shapes_by_name(Presentation(pptx)).items())
                if slide==index]
    return updated,outputs,{"inputs_to_outputs":mapping,"before":before,"after":after,
        "consumed":sorted(set(before)-set(after)),"added":sorted(set(after)-set(before)),
        "preserved":sorted(set(before)&set(after)),"output_names":outputs,
        "reparented":[{"name":key,"before":before[key]["parent"],"after":after[key]["parent"]}
                      for key in before.keys()&after.keys() if before[key]["parent"]!=after[key]["parent"]],
        "stacking_before":order(source),"stacking_after":order(output),
        "review_requirements_before":scene["slides"][index].get("review_requirements",{}),
        "review_requirements_after":requirements,
        "numeric_id_policy":"Read actual IDs; output stable names authoritative, numeric IDs may change",
        "backend":"PowerPoint.ShapeRange.Group/Ungroup/MergeShapes",
        "geometry_policy":"Persist actual closed M/L/C/Z compound paths; curves/holes retained without approximation",
        "request":change}


def protected_check(before, after, index, touched, outputs):
    old=semantic_nodes(before)
    new=semantic_nodes(after)
    old_rows=[(k,v) for k,v in old.items() if k not in {f"{index}/{n}" for n in touched}]
    new_rows=[(k,v) for k,v in new.items() if k not in {f"{index}/{n}" for n in outputs}]
    def normalized(rows):
        result=[]
        for key,value in rows:
            value={k:v for k,v in value.items() if k!="order"}
            root=etree.fromstring(value["xml"].encode())
            # Office adds an empty paragraph-end marker on a subsequent save.
            # Retain every attribute, child and nonempty marker as protected data.
            for node in root.iter(qn("a:endParaRPr")):
                if not node.attrib and not len(node) and not node.text:
                    node.getparent().remove(node)
            # PowerPoint may omit these schema defaults on the second save.
            # Compare effective defaults, while retaining non-default materials.
            for node in root.iter(qn('a:path')):
                for attr,default in [('fill','norm'),('stroke','1'),('extrusionOk','1')]:
                    effective=node.get(attr,default)
                    if attr!='fill' and effective=='true':effective='1'
                    if attr!='fill' and effective=='false':effective='0'
                    node.set(attr,effective)
            value["xml"]=etree.tostring(root,method="c14n").decode()
            result.append((key,value))
        return result
    # Global indices shift with legitimate topology; protected relative ordering must not.
    return normalized(old_rows)==normalized(new_rows)
