"""Render Agent-authored scene JSON to editable PPTX; no image understanding is implied."""
from __future__ import annotations
import argparse, io, math, sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
from pptx.util import Pt
from common import bbox_to_points, normalize_text, read_json, resolve_asset, sha256, walk_objects, write_json
from validate_scene import validate

GEOMETRY = {"rect":MSO_SHAPE.RECTANGLE,"round_rect":MSO_SHAPE.ROUNDED_RECTANGLE,
 "ellipse":MSO_SHAPE.OVAL,"triangle":MSO_SHAPE.ISOSCELES_TRIANGLE,"diamond":MSO_SHAPE.DIAMOND,
 "chevron":MSO_SHAPE.CHEVRON,"right_arrow":MSO_SHAPE.RIGHT_ARROW,
 "left_right_arrow":MSO_SHAPE.LEFT_RIGHT_ARROW,"star5":MSO_SHAPE.STAR_5_POINT,"arc":MSO_SHAPE.ARC}

def el(name, **attrs):
    e=OxmlElement(name)
    for k,v in attrs.items(): e.set(k,str(v))
    return e

def remove_theme_effect(shape):
    """Clear only inherited effects, not unrelated style/color references."""
    for ref in shape._element.findall('.//' + qn('a:effectRef')):
        ref.set('idx','0')

def set_gradient(shape, gradient, outline=False, alpha=1):
    sp=shape._element.find(qn('p:spPr'))
    if sp is None: raise ValueError('Shape has no spPr for gradient')
    if outline:
        sp=sp.get_or_add_ln()
    fill_tags={qn('a:'+x) for x in ['noFill','solidFill','gradFill','blipFill','pattFill','grpFill']}
    for child in list(sp):
        if child.tag in fill_tags: sp.remove(child)
    grad=el('a:gradFill',rotWithShape='1'); gslist=el('a:gsLst')
    for stop in sorted(gradient['stops'],key=lambda x:x['position']):
        gs=el('a:gs',pos=round(stop['position']*100000))
        rgb=el('a:srgbClr',val=stop['color'].upper())
        if 'alpha' in stop or alpha != 1:
            rgb.append(el('a:alpha',val=round(stop.get('alpha',1)*alpha*100000)))
        gs.append(rgb); gslist.append(gs)
    grad.append(gslist)
    if gradient.get('type','linear') == 'radial':
        cx,cy=gradient.get('center',[.5,.5])
        geometry=shape._element.find('./'+qn('p:spPr')+'/'+qn('a:prstGeom'))
        # Office's circle gradient uses a physical circle, so a wide ellipse
        # clips its vertical fade. Shape-relative gradients reach both edges.
        default_path='shape' if geometry is not None and geometry.get('prst')=='ellipse' else 'circle'
        path=el('a:path',path=gradient.get('path',default_path))
        path.append(el('a:fillToRect',l=round(cx*100000),t=round(cy*100000),
                       r=round((1-cx)*100000),b=round((1-cy)*100000)))
        grad.append(path)
    else:
        grad.append(el('a:lin',ang=round((gradient['angle_deg']%360)*60000),scaled='1'))
    # Insert after transform and geometry, before line/effects/3D/extensions.
    idx=0
    for i,child in enumerate(sp):
        if child.tag in {qn('a:xfrm'),qn('a:prstGeom'),qn('a:custGeom')}: idx=i+1
    sp.insert(idx,grad)

def apply_shape_style(shape, st):
    remove_theme_effect(shape)
    if not hasattr(shape, 'fill'):
        pass  # Line/Connector objects have no shape fill.
    elif 'gradient' in st:
        set_gradient(shape,st['gradient'],alpha=st.get('fill_alpha',1))
    elif st.get('fill') is not None:
        shape.fill.solid(); shape.fill.fore_color.rgb=RGBColor.from_string(st['fill'])
        if 'fill_alpha' in st:
            rgb=shape._element.find('.//' + qn('a:solidFill') + '/' + qn('a:srgbClr'))
            if rgb is not None: rgb.append(el('a:alpha',val=round(st['fill_alpha']*100000)))
    else: shape.fill.background()
    if st.get('line') is None and 'line_gradient' not in st: shape.line.fill.background()
    else:
        shape.line.color.rgb=RGBColor.from_string(st.get('line') or '000000')
        shape.line.width=Pt(st.get('line_width_pt',1))
        if st.get('dash')=='dash': shape.line.dash_style=MSO_LINE_DASH_STYLE.DASH
    if 'line_gradient' in st:
        set_gradient(shape,st['line_gradient'],outline=True,alpha=st.get('line_alpha',1))
    if st.get('line') is not None or 'line_gradient' in st:
        ln=shape._element.find('.//' + qn('a:ln'))
        if 'line_cap' in st: ln.set('cap',{'butt':'flat','round':'rnd','square':'sq'}[st['line_cap']])
        if 'line_alpha' in st and 'line_gradient' not in st:
            rgb=ln.find('.//' + qn('a:srgbClr'))
            if rgb is not None: rgb.append(el('a:alpha',val=round(st['line_alpha']*100000)))
        if 'line_join' in st:
            for child in list(ln):
                if child.tag in {qn('a:round'),qn('a:bevel'),qn('a:miter')}:ln.remove(child)
            join=el('a:'+st['line_join'])
            if st['line_join']=='miter':join.set('lim',str(round(st.get('miter_limit',4)*100000)))
            ln.append(join)
    if st.get('begin_arrow') or st.get('end_arrow'):
        ln=shape._element.find('.//' + qn('a:ln'))
        if ln is None: raise ValueError('No line node for arrow style')
        for key,tag in [('begin_arrow','headEnd'),('end_arrow','tailEnd')]:
            if key in st:
                for child in list(ln):
                    if child.tag==qn('a:'+tag):ln.remove(child)
                ln.append(el('a:'+tag,type=st[key],w=st.get(key+'_width','med'),len=st.get(key+'_length','med')))

def apply_font(run, st):
    run.font.name=st.get('font','Arial'); run.font.size=Pt(st.get('font_size_pt',18))
    run.font.bold=st.get('bold',False); run.font.italic=st.get('italic',False)
    run.font.color.rgb=RGBColor.from_string(st.get('color','222222'))
    rpr=run._r.get_or_add_rPr()
    ea=rpr.find(qn('a:ea'))
    if ea is None:
        ea=el('a:ea'); rpr.insert_element_before(ea,'a:cs','a:sym','a:hlinkClick','a:hlinkMouseOver','a:extLst')
    ea.set('typeface',st.get('font_east_asia',st.get('font','Arial')))
    if 'char_spacing_pt' in st:rpr.set('spc',str(round(st['char_spacing_pt']*100)))
    if 'text_outline' in st:
        # One editable text source, not duplicate offset text pretending to be a stroke.
        for old in list(rpr):
            if old.tag == qn('a:ln'): rpr.remove(old)
        outline = st['text_outline']
        ln = el('a:ln', w=int(Pt(outline['width_pt'])), cap='rnd', cmpd='sng', algn='ctr')
        sf = el('a:solidFill'); sf.append(el('a:srgbClr', val=outline['color'].upper()))
        ln.append(sf); ln.append(el('a:round'))
        rpr.insert(0, ln)
    if 'baseline' in st: rpr.set('baseline',{'normal':'0','superscript':'30000','subscript':'-25000'}[st['baseline']])

def set_text(tf, obj):
    st=obj.get('style',{}); tf.clear(); tf.word_wrap=st.get('wrap',False); tf.auto_size=MSO_AUTO_SIZE.NONE
    for side in ['left','right','top','bottom']:setattr(tf,'margin_'+side,Pt(st.get('margin_'+side+'_pt',st.get('margin_pt',0))))
    tf.vertical_anchor={'top':MSO_ANCHOR.TOP,'middle':MSO_ANCHOR.MIDDLE,'bottom':MSO_ANCHOR.BOTTOM}[st.get('valign','top')]
    if 'paragraphs' in obj:
        from native_text_range import paragraph_style
        for i,spec in enumerate(obj['paragraphs']):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
            paragraph_style(p,{"align":st.get("align","left"),"space_before_pt":st.get("space_before_pt",0),
                              "space_after_pt":st.get("space_after_pt",0),
                              **({"line_spacing_pt":st["line_spacing_pt"]} if "line_spacing_pt" in st else {}),
                              **spec["style"]})
            for rs in spec['runs']:
                r=p.add_run();r.text=rs['text'];apply_font(r,st|{k:v for k,v in rs.items() if k!='text'})
        return
    if st.get('warp')=='textPlain':
        bp=tf._txBody.bodyPr
        for old in list(bp):
            if old.tag==qn('a:prstTxWarp'):bp.remove(old)
        warp=el('a:prstTxWarp',prst='textPlain');warp.append(el('a:avLst'));bp.insert(0,warp)
    runs=obj.get('runs',[{'text':obj.get('text','')}]); p=tf.paragraphs[0]
    def para_style(p):
        p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[st.get('align','left')]
        p.space_before=Pt(st.get('space_before_pt',0));p.space_after=Pt(st.get('space_after_pt',0))
        if 'line_spacing_pt' in st:p.line_spacing=Pt(st['line_spacing_pt'])
    para_style(p)
    for spec in runs:
        for idx,part in enumerate(normalize_text(spec['text']).split('\n')):
            if idx:p=tf.add_paragraph();para_style(p)
            r=p.add_run();r.text=part;apply_font(r,st|{k:v for k,v in spec.items() if k!='text'})

def make_path(shapes, obj, canvas):
    commands=obj['commands']; sx=canvas['width_pt']/canvas['width'];sy=canvas['height_pt']/canvas['height']
    coords=[]
    for cmd in commands:
        coords.extend([(cmd[i]*sx,cmd[i+1]*sy) for i in range(1,len(cmd),2)])
    x=min(p[0] for p in coords);y=min(p[1] for p in coords)
    w=max(max(p[0] for p in coords)-x,.01);h=max(max(p[1] for p in coords)-y,.01)
    if 'native_path_frame' in obj:
        a,b,c,d=obj['native_path_frame']['bbox'];x,y,w,h=a*sx,b*sy,c*sx,d*sy
        if w<=0 or h<=0:raise ValueError('Positive actual path frame required')
    shape=shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(x),Pt(y),Pt(w),Pt(h));sp=shape._element.spPr
    if 'native_path_frame' in obj:
        shape.text_frame.word_wrap=obj['native_path_frame']['word_wrap']
        value=obj['native_path_frame']['auto_size']
        shape.text_frame.auto_size=None if value is None else MSO_AUTO_SIZE(value)
    old=sp.find(qn('a:prstGeom'));index=list(sp).index(old);sp.remove(old)
    cg=el('a:custGeom')
    for tag in ['avLst','gdLst','ahLst','cxnLst']:cg.append(el('a:'+tag))
    cg.append(el('a:rect',l='0',t='0',r='r',b='b'))
    paths=el('a:pathLst');path=el('a:path',w=int(Pt(w)),h=int(Pt(h)),fill='norm' if obj.get('closed') else 'none',stroke='1')
    for cmd in commands:
        op=cmd[0];node=el('a:'+{'M':'moveTo','L':'lnTo','C':'cubicBezTo','Z':'close'}[op])
        for i in range(1,len(cmd),2):node.append(el('a:pt',x=int(Pt(cmd[i]*sx-x)),y=int(Pt(cmd[i+1]*sy-y))))
        path.append(node)
    paths.append(path);cg.append(paths);sp.insert(index,cg)
    return shape

def build(scene_path: Path, out: Path, overwrite: bool=False) -> dict:
    scene_path=scene_path.resolve();base=scene_path.parent;scene=read_json(scene_path)
    issues=validate(scene,base)
    if issues:raise ValueError('\n'.join(issues))
    if out.exists() and not overwrite:raise FileExistsError(f'{out} exists; choose a new version or --overwrite')
    if out.resolve()==scene_path:raise ValueError('Output would overwrite input')
    out.parent.mkdir(parents=True,exist_ok=True)
    prs=Presentation();c=scene['canvas'];prs.slide_width=Pt(c['width_pt']);prs.slide_height=Pt(c['height_pt'])
    manifest=[]
    for si,sd in enumerate(scene['slides'],1):
        slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid()
        slide.background.fill.fore_color.rgb=RGBColor.from_string(c.get('background','FFFFFF'))
        lookup={};deferred=[]
        def create(shapes,o):
            kind=o['kind'];st=o.get('style',{})
            box=bbox_to_points(o['bbox'],c) if 'bbox' in o else None
            dims=[Pt(n) for n in box] if box else None
            if kind=='group':
                shape=shapes.add_group_shape()
                for child in o['children']:create(shape.shapes,child)
            elif kind=='text':
                shape=shapes.add_textbox(*dims);set_text(shape.text_frame,o)
            elif kind=='shape':
                shape=shapes.add_shape(GEOMETRY[o['geometry']],*dims)
                for idx,val in enumerate(o.get('adjustments',[])):shape.adjustments[idx]=val
            elif kind in {'line','connector'}:
                p1,p2=o['points'];sx=c['width_pt']/c['width'];sy=c['height_pt']/c['height']
                ct={'straight':MSO_CONNECTOR.STRAIGHT,'elbow':MSO_CONNECTOR.ELBOW,'curve':MSO_CONNECTOR.CURVE}[o.get('connector_type','straight')]
                shape=shapes.add_connector(ct,Pt(p1[0]*sx),Pt(p1[1]*sy),Pt(p2[0]*sx),Pt(p2[1]*sy))
                if kind=='connector':deferred.append((shape,o))
            elif kind=='path':shape=make_path(shapes,o,c)
            elif kind=='image':
                asset=resolve_asset(base,o['asset']);x,y,w,h=box
                with Image.open(asset) as im: iw,ih=im.size
                fit=o.get('fit','contain')
                if 'source_crop' in o:
                    cx,cy,cw,ch=o['source_crop']
                    shape=shapes.add_picture(str(asset),*dims)
                    shape.crop_left=cx/iw;shape.crop_top=cy/ih
                    shape.crop_right=(iw-cx-cw)/iw;shape.crop_bottom=(ih-cy-ch)/ih
                elif fit=='contain':
                    scale=min(w/iw,h/ih); nw,nh=iw*scale,ih*scale
                    shape=shapes.add_picture(str(asset),Pt(x+(w-nw)/2),Pt(y+(h-nh)/2),Pt(nw),Pt(nh))
                else:
                    shape=shapes.add_picture(str(asset),*dims)
                    if fit=='cover':
                        scale=max(w/iw,h/ih);vw,vh=w/scale,h/scale
                        shape.crop_left=shape.crop_right=(iw-vw)/(2*iw)
                        shape.crop_top=shape.crop_bottom=(ih-vh)/(2*ih)
                if 'mask' in o:
                    from native_extensions import apply_picture_mask
                    apply_picture_mask(shape,o['mask'])
            elif kind=='table':
                rows=o['rows'];shape=shapes.add_table(len(rows),len(rows[0]),*dims);table=shape.table
                for ri,row in enumerate(rows):
                    table.rows[ri].height=Pt(box[3]/len(rows))
                    for ci,text in enumerate(row):
                        cell=table.cell(ri,ci)
                        cell.margin_left=cell.margin_right=Pt(st.get('margin_pt',3));cell.margin_top=cell.margin_bottom=Pt(1)
                        cell.fill.solid();cell.fill.fore_color.rgb=RGBColor.from_string(o.get('header_fill','DDE7F2') if ri==0 else st.get('fill') or 'FFFFFF')
                        set_text(cell.text_frame,{'text':text,'style':st|{'margin_pt':st.get('margin_pt',3)}})
                        tcp=cell._tc.get_or_add_tcPr()
                        for tag in ['lnL','lnR','lnT','lnB']:
                            for old in list(tcp):
                                if old.tag==qn('a:'+tag):tcp.remove(old)
                            ln=el('a:'+tag,w=int(Pt(st.get('line_width_pt',.6))))
                            sf=el('a:solidFill');sf.append(el('a:srgbClr',val=st.get('line') or 'B8C6D6'));ln.append(sf)
                            tcp.insert_element_before(ln,'a:cell3D','a:noFill','a:solidFill','a:gradFill','a:blipFill','a:pattFill','a:grpFill','a:headers','a:extLst')
                for ci,width in enumerate(o.get('column_widths',[o['bbox'][2]/len(rows[0])]*len(rows[0]))):table.columns[ci].width=Pt(width*c['width_pt']/c['width'])
                from native_extensions import table_extensions
                table_extensions(shape,o,c,set_text)
            elif kind=='chart':
                if o['chart_type']=='xy':
                    data=XyChartData()
                    for spec in o['series']:
                        ser=data.add_series(spec['name'])
                        for x,y in spec['points']:ser.add_data_point(x,y)
                    ct=XL_CHART_TYPE.XY_SCATTER_LINES
                else:
                    data=CategoryChartData();data.categories=o['categories']
                    for spec in o['series']:data.add_series(spec['name'],spec['values'])
                    ct=XL_CHART_TYPE.COLUMN_CLUSTERED if o['chart_type']=='column' else XL_CHART_TYPE.LINE
                shape=shapes.add_chart(ct,*dims,data);chart=shape.chart;chart.has_title=False;chart.has_legend=o.get('legend',True)
                if chart.has_legend:chart.legend.position=XL_LEGEND_POSITION.BOTTOM;chart.legend.include_in_layout=False
                chart.font.name=st.get('font','Arial');chart.font.size=Pt(st.get('font_size_pt',12))
                for spec,ser in zip(o['series'],chart.series):
                    if 'color' in spec:
                        ser.format.line.color.rgb=RGBColor.from_string(spec['color'])
                        if o['chart_type']=='column':ser.format.fill.solid();ser.format.fill.fore_color.rgb=RGBColor.from_string(spec['color'])
                from native_extensions import chart_extensions
                chart_extensions(chart,o)
            else:raise ValueError(f'Unsupported kind: {kind}')
            shape.name=o['id'];lookup[o['id']]=shape
            if kind in {'text','shape','line','connector','path'}:apply_shape_style(shape,st)
            if 'native_format' in st:
                from native_format import apply as apply_native_format
                apply_native_format(shape, st['native_format'])
            if kind=='text':shape.left,shape.top,shape.width,shape.height=dims
            if 'rotation' in o:shape.rotation=o['rotation']
            manifest.append({'slide':si,'slide_id':sd['id'],'id':o['id'],'kind':kind,'editability':o['editability']})
            return shape
        for obj in sd['objects']:create(slide.shapes,obj)
        for con,o in deferred:
            for endpoint in ['begin','end']:
                if endpoint in o:
                    getattr(con,endpoint+'_connect')(lookup[o[endpoint]['object_id']],o[endpoint]['site'])
        notes=[sd.get('notes','')]+[f"{o['id']}: {o['data_provenance']}" for o in walk_objects(sd['objects']) if o['kind']=='chart']
        slide.notes_slide.notes_text_frame.text='\n'.join(n for n in notes if n)
    prs.save(out)
    report={'status':'passed','scope':'PPTX generation only; not an Office/visual acceptance','pptx_sha256':sha256(out),'scene_sha256':sha256(scene_path),'slide_count':len(scene['slides']),'objects':manifest}
    write_json(out.with_suffix('.build.json'),report)
    return report

def main():
    ap=argparse.ArgumentParser(description=__doc__);ap.add_argument('scene',type=Path);ap.add_argument('output',type=Path);ap.add_argument('--overwrite',action='store_true')
    a=ap.parse_args()
    try:r=build(a.scene,a.output,a.overwrite);print(f"Created {a.output}; {r['slide_count']} slides; Office review still required")
    except Exception as exc:print(f'Build failed: {exc}',file=sys.stderr);return 1
    return 0
if __name__=='__main__':raise SystemExit(main())
