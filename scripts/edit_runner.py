"""Plan and repeat bounded Office sample edits, preserving individual receipts."""
from __future__ import annotations

import copy
import hashlib
from pathlib import Path

if __package__:
    from .common import read_json, write_json, sha256, staged_directory
    from .review_impact import digest
    from .office_edit import validate_operations
    from .edit_reuse import equivalence
else:
    from common import read_json, write_json, sha256, staged_directory
    from review_impact import digest
    from office_edit import validate_operations
    from edit_reuse import equivalence

PLAN = 'office-edit-plan/1'
RUN = 'office-edit-run/1'


def ref(path):
    path = Path(path).resolve()
    return {'path': str(path), 'sha256': sha256(path)}


def bound(reference, check_path=lambda p: p):
    path = Path(check_path(Path(reference['path']))).resolve()
    if not path.is_file() or sha256(path) != reference['sha256']:
        raise ValueError('Edit runner evidence changed: '+str(path))
    return path


def inventory(source):
    from pptx import Presentation
    from lxml import etree
    deck = Presentation(source)
    rows = []
    def walk(shapes, slide, grouped=False):
        for shape in shapes:
            if shape.shape_type == 6:
                walk(shape.shapes, slide, True)
                continue
            rows.append({'slide': slide, 'name': shape.name, 'shape': shape, 'grouped': grouped,
                         'target_sha256': hashlib.sha256(etree.tostring(shape._element, method='c14n')).hexdigest()})
    for slide, item in enumerate(deck.slides, 1):
        walk(item.shapes, slide)
    return deck, rows


def suggested_operations(rows):
    """One explicit, modest edit per supported kind; always applied to a copy."""
    result, seen = [], set()
    counts = {}
    for row in rows:
        key = (row['slide'], row['name'])
        counts[key] = counts.get(key, 0)+1
    for row in rows:
        shape = row['shape']
        if not row['name'] or counts[(row['slide'], row['name'])] != 1:
            continue
        base = {'slide': row['slide'], 'name': row['name']}
        if shape.has_table:
            op = {**base, 'op': 'table.cell', 'row': 1, 'column': 1,
                  'text': (shape.table.cell(0, 0).text or '')+' [edit test]'}
        elif shape.shape_type == 13 and not row['grouped']:
            geometry = {k: float(getattr(shape, k))/12700 for k in ('left', 'top', 'width', 'height')}
            geometry['rotation'] = float(shape.rotation)
            op = {**base, 'op': 'picture.transform', 'units': 'slide_points', 'aspect_ratio': 'preserve',
                  'expected': geometry, 'value': {**geometry, 'left': geometry['left']+5,
                      'width': geometry['width']*1.05, 'height': geometry['height']*1.05}}
        elif shape.has_text_frame and shape.text.strip():
            op = {**base, 'op': 'text.set', 'text': shape.text+' [edit test]'}
        elif shape.shape_type in (1, 5):
            try:
                current = str(shape.fill.fore_color.rgb)
            except (AttributeError, ValueError, TypeError):
                current = None
            op = {**base, 'op': 'shape.fill', 'color': '278A69' if current != '278A69' else 'A74F39'}
        else:
            continue
        if op['op'] not in seen:
            result.append(op)
            seen.add(op['op'])
    if not result:
        raise ValueError('No unambiguous supported samples; supply explicit operations')
    return result


def environment(source):
    # Include theme font declarations too. An unresolved identity disables reuse.
    if __package__:
        from .text_calibration import environment as text_environment
    else:
        from text_calibration import environment as text_environment
    from lxml import etree
    from zipfile import ZipFile
    fonts = set()
    with ZipFile(source) as package:
        for name in package.namelist():
            if name.endswith('.xml') and name.startswith(('ppt/slides/', 'ppt/theme/')):
                root = etree.fromstring(package.read(name), etree.XMLParser(resolve_entities=False))
                fonts.update(value for value in root.xpath('//*[local-name()="latin" or local-name()="ea" or local-name()="cs"]/@typeface')
                             if value and not value.startswith('+'))
    env = text_environment(sorted(fonts))
    code = {name: sha256(Path(__file__).with_name(name)) for name in
            ('edit_runner.py', 'edit_reuse.py', 'office_edit.py', 'material_edit_review.py')}
    return {'available': env['available'], 'sha256': digest([env['sha256'], code]),
            'fonts': sorted(fonts), 'office': env['office'], 'code': code}


def evidence_check(evidence, operation, source_ref, check_path=lambda p: p):
    receipt_path = bound(evidence['receipt'], check_path)
    source = bound(evidence['input'], check_path)
    output = bound(evidence['output'], check_path)
    renders = [bound(item, check_path) for item in evidence['renders']]
    receipt = read_json(receipt_path)
    if (receipt.get('format') != 'ppt-edit-readback/1' or receipt.get('source_sha256') != source_ref['sha256']
            or evidence['input'] != source_ref or Path(receipt.get('source', '')).resolve() != source
            or receipt.get('edited_sha256') != evidence['output']['sha256']
            or (receipt_path.parent/receipt.get('edited_file', '')).resolve() != output
            or not renders or [(receipt_path.parent/r['file']).resolve() for r in receipt.get('renders', [])] != renders
            or [r['sha256'] for r in receipt['renders']] != [r['sha256'] for r in evidence['renders']]):
        raise ValueError('Edit runner receipt does not bind its files')
    index = evidence['check_index']
    if type(index) is not int or not 0 <= index < len(receipt.get('checks', [])):
        raise ValueError('Edit runner check index is invalid')
    check = receipt['checks'][index]
    if check.get('operation') != operation:
        raise ValueError('Edit runner operation differs from receipt')
    return (receipt.get('status') == 'passed' and check.get('matches') is True
            and ('font_size_pt' not in operation or check.get('font_size_matches') is True)
            and check.get('longer_text_layout', {}).get('status') != 'overflow_detected')


def replacement_inputs(source, operations, check_path=lambda p: p):
    refs = []
    for op in operations:
        if op['op'] == 'picture.replace':
            path = (Path(source).parent/op['image']).resolve()
            if not path.is_relative_to(Path(source).parent.resolve()):
                raise ValueError('Replacement image escaped candidate directory')
            item = {'path': str(path), 'sha256': op['image_sha256']}
            bound(item, check_path)
            refs.append(item)
    return refs


def prepare(source, outdir, operations=None, previous=None, check_path=lambda p: p, *, reuse_unchanged_slides=False):
    source = Path(check_path(Path(source))).resolve()
    source_ref = ref(source)
    _, rows = inventory(source)
    ops = validate_operations(copy.deepcopy(operations if operations is not None else suggested_operations(rows)))
    dependencies = replacement_inputs(source, ops, check_path)
    env = environment(source)
    old, reuse_reason, signature_cache = {}, 'not_requested', {}
    if previous:
        previous = Path(check_path(Path(previous))).resolve()
        prior = read_json(previous)
        if prior.get('format') != RUN:
            raise ValueError('Previous evidence must be an office-edit-run/1 manifest')
        bound(prior['source'], check_path)
        # Verify even evidence that will not be reused. Never silently conceal tampering.
        for case in prior['cases']:
            if case.get('evidence'):
                origin = case['evidence']['input']
                valid = evidence_check(case['evidence'], case['operation'], origin, check_path)
                equivalent = (origin == source_ref or (reuse_unchanged_slides and
                    equivalence(origin['path'], source, case['operation']['slide'], signature_cache)))
                if valid and equivalent:
                    old[digest(case['operation'])] = case['evidence']
        reuse_reason = 'exact_candidate_operation_and_environment'
        if reuse_unchanged_slides and prior['source'] != source_ref:
            reuse_reason = 'unchanged_slide_and_shared_package_parts'
        if ((prior['source'] != source_ref and not reuse_unchanged_slides) or not env['available']
                or prior['environment']['sha256'] != env['sha256']):
            old = {}
            reuse_reason = 'candidate_or_environment_changed_or_unresolved'
    cases = []
    for op in ops:
        matching = [row for row in rows if (row['slide'], row['name']) == (op['slide'], op['name'])]
        if len(matching) != 1:
            raise ValueError('Object must resolve uniquely in candidate: '+op['name'])
        identity = {k:op[k] for k in ('slide','name','op','row','column','edge','index','property') if k in op}
        cases.append({'id': digest(identity)[:20], 'operation': op, 'operation_sha256': digest(op),
                      'target_sha256': matching[0]['target_sha256'], 'reuse': old.get(digest(op))})
    record = {'format': PLAN, 'source': source_ref, 'environment': env, 'cases': cases,
              'reuse_unchanged_slides': reuse_unchanged_slides,
              'dependencies': dependencies, 'reuse_reason': reuse_reason,
              'previous': ref(previous) if previous else None, 'visual_status': 'pending',
              'scope': 'explicit sampled operations only', 'all_objects_tested': False}
    with staged_directory(outdir) as stage:
        write_json(stage/'operations.json', ops)
        record['operations_sha256'] = sha256(stage/'operations.json')
        write_json(stage/'edit-plan.json', record)
    return {'format': PLAN, 'status': 'planned', 'plan': str(Path(outdir).resolve()/'edit-plan.json'),
            'operations': str(Path(outdir).resolve()/'operations.json'), 'samples': len(cases),
            'run_count': sum(not c['reuse'] for c in cases), 'reuse_count': sum(bool(c['reuse']) for c in cases),
            'reuse_reason': reuse_reason, 'visual_status': 'pending'}


def load_plan(path, check_path=lambda p: p, *, check_environment=True):
    path = Path(check_path(Path(path))).resolve()
    plan = read_json(path)
    if plan.get('format') != PLAN:
        raise ValueError('Unknown Office edit plan format')
    source = bound(plan['source'], check_path)
    ops = validate_operations([case['operation'] for case in plan['cases']])
    operation_file = Path(check_path(path.parent/'operations.json'))
    if sha256(operation_file) != plan['operations_sha256'] or read_json(operation_file) != ops:
        raise ValueError('Operation file changed; regenerate edit plan')
    signature_cache = {}
    for case in plan['cases']:
        if digest(case['operation']) != case['operation_sha256']:
            raise ValueError('Planned operation changed; regenerate plan')
        if case['reuse']:
            origin = case['reuse']['input']
            if not evidence_check(case['reuse'], case['operation'], origin, check_path):
                raise ValueError('Failed or overflowing sample cannot be reused')
            if origin != plan['source'] and not (plan.get('reuse_unchanged_slides') is True and
                    equivalence(origin['path'], source, case['operation']['slide'], signature_cache)):
                raise ValueError('Changed slide or shared dependencies cannot reuse edit evidence')
    if plan.get('previous'):
        bound(plan['previous'], check_path)
    for item in plan['dependencies']:
        bound(item, check_path)
    if replacement_inputs(source, ops, check_path) != plan['dependencies']:
        raise ValueError('Replacement dependencies changed')
    if check_environment and environment(source)['sha256'] != plan['environment']['sha256']:
        raise ValueError('Office or font environment changed; regenerate edit plan')
    if any(c['reuse'] for c in plan['cases']) and not plan['environment']['available']:
        raise ValueError('Unresolved environment cannot reuse edit checks')
    return plan


def run(path, outdir, application_factory=None):
    path, outdir = Path(path).resolve(), Path(outdir).resolve()
    if outdir.exists():
        raise FileExistsError(outdir)
    plan = load_plan(path)
    pending = [case['operation'] for case in plan['cases'] if not case['reuse']]
    outdir.mkdir(parents=True, exist_ok=False)
    receipt, receipt_path = None, outdir/'changed/edit-readback.json'
    if pending:
        if __package__:
            from .office_edit import execute
        else:
            from office_edit import execute
        receipt = execute(plan['source']['path'], receipt_path.parent, pending, application_factory=application_factory)
    result = {'format': RUN, 'source': plan['source'], 'plan': ref(path), 'environment': plan['environment'],
              'status': 'passed', 'visual_status': 'pending', 'scope': plan['scope'],
              'all_objects_tested': False, 'cases': []}
    for case in plan['cases']:
        evidence = case['reuse']
        if not evidence and receipt and receipt.get('edited_file'):
            matches = [i for i,c in enumerate(receipt['checks']) if c['operation'] == case['operation']]
            if len(matches) == 1:
                evidence = {'receipt': ref(receipt_path), 'input': plan['source'],
                            'output': ref(receipt_path.parent/receipt['edited_file']), 'check_index': matches[0],
                            'renders': [ref(receipt_path.parent/r['file']) for r in receipt['renders']]}
        passed = bool(evidence) and evidence_check(evidence, case['operation'], evidence['input'])
        if not passed:
            result['status'] = 'failed'
        result['cases'].append({k:case[k] for k in ('id','operation','operation_sha256','target_sha256')} | {
            'reused': bool(case['reuse']), 'status': 'passed' if passed else 'failed', 'evidence': evidence,
            'candidate_binding': 'current' if not evidence or evidence['input'] == plan['source'] else 'equivalent_slide_only',
            'formal_submission': 'requires_current_candidate_evidence' if evidence and evidence['input'] != plan['source'] else 'existing_receipt_contract'})
    if receipt and receipt.get('error'):
        result['error'] = receipt['error']
    result.update(executed=len(pending), reused=len(plan['cases'])-len(pending))
    write_json(outdir/'edit-run.json', result)
    return {**result, 'manifest': str(outdir/'edit-run.json')}
