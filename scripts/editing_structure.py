"""Read-only editability hints and an explicit PowerPoint group movement roundtrip.

Hints are not visual verdicts. Office execution writes only a new test copy.
"""
from __future__ import annotations
import argparse
import json
import math
import os
import shutil
import sys
from pathlib import Path
from pptx import Presentation
from common import sha256, write_json
from office_edit import find_shape


def audit(source):
    source=Path(source).resolve(); prs=Presentation(source); slides=[]
    for index,slide in enumerate(prs.slides,1):
        groups=[];texts=[];leaves=[]
        def visit(shapes,parents):
            for s in shapes:
                if s.shape_type==6:
                    groups.append({'name':s.name,'parent_groups':parents,'children':len(s.shapes)})
                    visit(s.shapes,parents+[s.name])
                else:
                    leaves.append({'name':s.name,'parent_groups':parents})
                    if s.has_text_frame and s.text.strip():
                        texts.append({'name':s.name,'text':s.text,'parent_groups':parents,
                                      'bbox_emu':[s.left,s.top,s.width,s.height]})
        visit(slide.shapes,[])
        hints=[]
        for i,a in enumerate(texts):
            for b in texts[i+1:]:
                # Raw child coordinates can only be compared within the same group.
                if a['text']!=b['text'] or a['parent_groups']!=b['parent_groups']:continue
                x,y,w,h=a['bbox_emu'];u,v,p,q=b['bbox_emu']
                area=max(0,min(x+w,u+p)-max(x,u))*max(0,min(y+h,v+q)-max(y,v))
                ratio=area/max(1,min(w*h,p*q))
                if ratio>=.6:hints.append({'objects':[a['name'],b['name']], 'text':a['text'],
                                         'overlap_fraction':ratio,'reason':'Potential duplicate text backing; edit one source and inspect the saved render.'})
        slides.append({'slide':index,'groups':groups,'leaf_count':len(leaves),
                       'ungrouped_leaves':[o['name'] for o in leaves if not o['parent_groups']],
                       'duplicate_text_hints':hints})
    return {'format':'ppt-editing-structure/1','source_sha256':sha256(source),'slides':slides,
            'scope':'Object structure and same-parent duplicate-text hints only; no automatic grouping or visual acceptance.'}


def snapshot(group):
    out={}
    def visit(s,path):
        key=path+'/'+str(s.Name)
        out[key]={'box':[float(s.Left),float(s.Top),float(s.Width),float(s.Height)],'rotation':float(s.Rotation)}
        if int(s.Type)==6:
            for i in range(1,s.GroupItems.Count+1):visit(s.GroupItems.Item(i),key)
    visit(group,'')
    return out


def compare_snapshots(before,after,dx=0,dy=0,tolerance=.05):
    """PowerPoint child bounds are in slide points; compare all descendant deltas."""
    if set(before)!=set(after):return {'matches':False,'issues':['Group membership changed']}
    issues=[]
    for name,b in before.items():
        expected=[b['box'][0]+dx,b['box'][1]+dy,*b['box'][2:]];actual=after[name]
        if any(abs(x-y)>tolerance for x,y in zip(expected,actual['box'])) or abs(b['rotation']-actual['rotation'])>tolerance:
            issues.append(name)
    return {'matches':not issues,'issues':issues,'tolerance_pt':tolerance,'objects_checked':len(before)}


def roundtrip(source,outdir,slide,name,dx,dy):
    source=Path(source).resolve();outdir=Path(outdir).resolve()
    if not source.is_file() or source.suffix.lower()!='.pptx':raise ValueError('Existing PPTX required')
    if outdir.exists():raise ValueError('Output directory must be new')
    if slide<1 or not name.strip():raise ValueError('One-based slide and exact group name required')
    if not all(math.isfinite(v) for v in [dx,dy]) or (dx==0 and dy==0):raise ValueError('Finite nonzero displacement required')
    if os.name!='nt':raise RuntimeError('Windows PowerPoint required')
    from win32com.client import Dispatch
    app=Dispatch('PowerPoint.Application'); deck=None;source_hash=sha256(source)
    outdir.mkdir(parents=True);copy=outdir/'group-test-copy.pptx';shutil.copy2(source,copy)
    report={'format':'ppt-group-roundtrip/1','status':'failed','source_sha256':source_hash,
            'slide':slide,'group':name,'displacement_pt':[dx,dy],'visual_status':'not_run',
            'scope':'Explicit move/save/reopen and inverse move/save/reopen; not application Undo or arbitrary scaling coverage.'}
    try:
        deck=app.Presentations.Open(str(copy),0,0,0)
        group=find_shape(deck.Slides.Item(slide).Shapes,name)
        if int(group.Type)!=6:raise ValueError('Target must be a native group')
        before=snapshot(group);group.IncrementLeft(dx);group.IncrementTop(dy)
        deck.Save();deck.Close();deck=None
        deck=app.Presentations.Open(str(copy),0,0,0)
        group=find_shape(deck.Slides.Item(slide).Shapes,name);moved=snapshot(group)
        report['move_check']=compare_snapshots(before,moved,dx,dy)
        deck.Slides.Item(slide).Export(str(outdir/'moved.png'),'PNG',1600,round(1600*deck.PageSetup.SlideHeight/deck.PageSetup.SlideWidth))
        group.IncrementLeft(-dx);group.IncrementTop(-dy);deck.Save();deck.Close();deck=None
        deck=app.Presentations.Open(str(copy),-1,0,0)
        restored=snapshot(find_shape(deck.Slides.Item(slide).Shapes,name))
        report['restore_check']=compare_snapshots(before,restored)
        deck.Slides.Item(slide).Export(str(outdir/'restored.png'),'PNG',1600,round(1600*deck.PageSetup.SlideHeight/deck.PageSetup.SlideWidth))
        deck.Close();deck=None
        if sha256(source)!=source_hash:raise RuntimeError('Source changed')
        report.update(before=before,moved=moved,restored=restored,copy_sha256=sha256(copy))
        report['status']='passed' if report['move_check']['matches'] and report['restore_check']['matches'] else 'failed'
    except Exception as exc:report['error']=str(exc)
    finally:
        if deck is not None:
            try:deck.Saved=-1;deck.Close()
            except Exception as exc:report['cleanup_error']=str(exc)
        write_json(outdir/'group-roundtrip.json',report)
    return report


def main():
    p=argparse.ArgumentParser(description=__doc__);s=p.add_subparsers(dest='command',required=True)
    a=s.add_parser('audit');a.add_argument('--pptx',type=Path,required=True);a.add_argument('--out',type=Path,required=True)
    a=s.add_parser('group-roundtrip');a.add_argument('--pptx',type=Path,required=True);a.add_argument('--outdir',type=Path,required=True)
    a.add_argument('--slide',type=int,required=True);a.add_argument('--name',required=True);a.add_argument('--dx',type=float,default=12);a.add_argument('--dy',type=float,default=6)
    args=p.parse_args()
    try:
        if args.command=='audit':
            if args.out.exists():raise ValueError('Output must be new')
            result=audit(args.pptx);write_json(args.out,result)
        else:result=roundtrip(args.pptx,args.outdir,args.slide,args.name,args.dx,args.dy)
        summary={k:v for k,v in result.items() if k not in {'before','moved','restored'}}
        print(json.dumps(summary,ensure_ascii=False,indent=2));return 1 if result.get('status')=='failed' else 0
    except Exception as exc:print(json.dumps({'status':'failed','error':str(exc)},ensure_ascii=False));return 2


if __name__=='__main__':
    if hasattr(sys.stdout,'reconfigure'):sys.stdout.reconfigure(encoding='utf-8')
    raise SystemExit(main())
