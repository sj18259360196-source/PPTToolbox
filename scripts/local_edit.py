"""Finite incremental OOXML edits. Rebuilds are verification, never this backend."""
from __future__ import annotations

import copy
import hashlib
from pathlib import Path
from zipfile import ZipFile

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt

from common import bbox_to_points, read_json, resolve_asset, sha256, walk_objects, write_json
from patch_ops import _shift
from validate_scene import validate

STYLE_FIELDS = {"font_size_pt", "color", "align", "margin_left_pt", "margin_right_pt",
                "margin_top_pt", "margin_bottom_pt"}


def shapes_by_name(prs):
    result = {}
    def visit(shapes, slide, parent=None):
        for shape in shapes:
            key = (slide, shape.name)
            if key in result:
                raise ValueError("Duplicate stable object name")
            result[key] = (shape, parent)
            if shape.shape_type == 6:
                visit(shape.shapes, slide, shape.name)
    for index, slide in enumerate(prs.slides):
        visit(slide.shapes, index)
    return result


def properties(pptx):
    """Independent persisted-object readback, not generated patch expectations."""
    prs = Presentation(pptx)
    result = {}
    native = shapes_by_name(prs)
    for (slide, name), (shape, parent) in native.items():
        x, y, w, h = shape.left, shape.top, shape.width, shape.height
        ancestor = parent
        while ancestor:
            group, ancestor = native[slide, ancestor]
            tr = group._element.find(qn("p:grpSpPr")).find(qn("a:xfrm"))
            off, ext, co, ce = (tr.find(qn("a:"+t)) for t in ("off", "ext", "chOff", "chExt"))
            sx, sy = int(ext.get("cx"))/int(ce.get("cx")), int(ext.get("cy"))/int(ce.get("cy"))
            x, y = int(off.get("x"))+(x-int(co.get("x")))*sx, int(off.get("y"))+(y-int(co.get("y")))*sy
            w, h = w*sx, h*sy
        row = {"type": int(shape.shape_type), "parent": parent, "shape_id": shape.shape_id,
               "bbox_pt": [round(v/12700, 4) for v in (x, y, w, h)],
               "rotation": round(shape.rotation, 4)}
        if hasattr(shape, "line"):
            row["line_width_pt"] = float(shape.line.width.pt) if shape.line.width else 0.0
        if shape.has_text_frame:
            tf = shape.text_frame
            row["text"] = tf.text
            row["auto_size"] = int(tf.auto_size) if tf.auto_size is not None else None
            row["word_wrap"] = tf.word_wrap
            row["vertical_anchor"] = int(tf.vertical_anchor) if tf.vertical_anchor is not None else None
            row["margins_pt"] = [round(getattr(tf, "margin_"+s)/12700, 4) for s in ("left", "right", "top", "bottom")]
            row["paragraphs"] = [{
                "align": int(p.alignment) if p.alignment is not None else None,
                "runs": [{"text": r.text, "font": r.font.name,
                          "size_pt": r.font.size.pt if r.font.size is not None else None,
                          "bold": r.font.bold, "italic": r.font.italic,
                          "color": str(r.font.color.rgb) if r.font.color.type == 1 else str(r.font.color.type)}
                         for r in p.runs]} for p in tf.paragraphs]
        if hasattr(shape, "fill"):
            row["fill_type"] = str(shape.fill.type)
            if shape.fill.type == 1:
                row["fill_rgb"] = str(shape.fill.fore_color.rgb) if shape.fill.fore_color.type == 1 else str(shape.fill.fore_color.type)
        if shape.shape_type == 13:
            row["crop"] = [round(getattr(shape, "crop_"+s), 5) for s in ("left", "top", "right", "bottom")]
            row["media_sha256"] = hashlib.sha256(shape.image.blob).hexdigest()
        result[f"{slide}/{name}"] = row
    return result


def semantic_nodes(pptx):
    """Full per-object XML after comparable Office normalization; keep unknown content."""
    result = {}
    prs = Presentation(pptx)
    for order, (key, (shape, parent)) in enumerate(shapes_by_name(prs).items()):
        node = copy.deepcopy(shape._element)
        if shape.shape_type == 6:
            for child in list(node):
                if etree.QName(child).localname in {"sp", "pic", "grpSp", "graphicFrame", "cxnSp"}:
                    node.remove(child)
        # Office creation IDs are document metadata, not object content.
        for child in list(node.iter()):
            if etree.QName(child).localname in {"creationId", "modId"}:
                child.getparent().remove(child)
        result[f"{key[0]}/{key[1]}"] = {
            "parent": parent, "order": order,
            "xml": etree.tostring(node, method="c14n", exclusive=True).decode("utf-8")}
    return result


def part_diff(left, right):
    with ZipFile(left) as a, ZipFile(right) as b:
        aa, bb = set(a.namelist()), set(b.namelist())
        return {"added": sorted(bb-aa), "removed": sorted(aa-bb),
                "changed": sorted(n for n in aa & bb if a.read(n) != b.read(n))}


def prepare(scene_path, pptx, slide_id, allowed, changes, *, calibration=False):
    scene = read_json(scene_path)
    updated = copy.deepcopy(scene)
    si = next(i for i, s in enumerate(scene["slides"]) if s["id"] == slide_id)
    old = {o["id"]: o for o in walk_objects(scene["slides"][si]["objects"])}
    new = {o["id"]: o for o in walk_objects(updated["slides"][si]["objects"])}
    prs = Presentation(pptx)
    native = shapes_by_name(prs)
    canvas = scene["canvas"]
    touched = set()
    operations = set()
    for ch in changes:
        oid, op = ch["id"], ch["op"]
        if oid not in allowed or oid not in old or (si, oid) not in native:
            raise ValueError("Object outside existing authorized region")
        obj = old[oid]
        shape, parent = native[si, oid]
        if (oid, op) in operations or (oid in touched and not calibration):
            raise ValueError("Use one operation per object in a candidate")
        operations.add((oid, op))
        touched.add(oid)
        if op != "native.format" and shape._element.xpath(".//a:effectLst/*|.//a:effectDag/*"):
            raise ValueError("Effects/shadows are outside supported incremental bounds")
        if (parent and op != 'native.nodes') or shape.rotation or obj.get("rotation_deg", 0):
            raise ValueError("Only unrotated top-level targets supported; use group.translate for groups")
        if op == 'native.nodes':
            from native_nodes import apply as apply_nodes
            replacement=apply_nodes(obj,shape,ch,canvas,native,si,parent)
            new[oid].clear();new[oid].update(replacement)
        elif op == "text.range":
            from native_text_range import apply as apply_range
            replacement=apply_range(obj,shape,ch)
            new[oid].clear();new[oid].update(replacement)
        elif op == "native.format":
            from native_capabilities import validate_format
            from native_format import apply as apply_native, effect_margin
            validate_format(ch["format"], obj["kind"])
            if obj["kind"] not in {"shape", "text"} or "runs" in obj or "paragraphs" in obj:
                raise ValueError("Native format supports independent uniform text/autoshapes")
            if obj["kind"] == "shape" and obj.get("geometry") not in {"rect", "round_rect", "ellipse"}:
                raise ValueError("Native effect shape scope is rectangle/round rectangle/ellipse")
            if shape._element.xpath(".//a:effectDag|.//a:effectLst/*[not(self::a:outerShdw)]"):
                raise ValueError("Unrepresented native effect must be preserved; target unsupported")
            prior = obj.get("style", {}).get("native_format", {})
            if shape._element.xpath(".//a:outerShdw|.//a:sp3d") and not prior:
                raise ValueError("Existing effect has no tracked scene representation")
            if obj["kind"] == "text":
                if len(shape.text_frame.paragraphs) != 1 or len(shape.text_frame.paragraphs[0].runs) != 1:
                    raise ValueError("Native glyph format requires one paragraph/one run")
            if "text" in ch:
                if obj["kind"] != "text" or any(c in ch["text"] for c in "\r\n\v"):
                    raise ValueError("Native text replacement requires single-line text")
                shape.text_frame.paragraphs[0].runs[0].text = ch["text"]
                new[oid]["text"] = ch["text"]
            merged = {**prior, **ch["format"]}
            effect_margin(merged)
            new[oid].setdefault("style", {})["native_format"] = merged
            apply_native(shape, ch["format"])
        elif op in {"text.set", "text.style"}:
            if obj["kind"] != "text" or "runs" in obj or "paragraphs" in obj or not shape.has_text_frame:
                raise ValueError("Plain native text only; explicit rich runs are not supported")
            tf = shape.text_frame
            simple = all(len(p.runs) == 1 for p in tf.paragraphs)
            if not simple or (len(tf.paragraphs) != 1 and not calibration) or shape._element.xpath(".//a:fld"):
                raise ValueError("Single paragraph and single run only")
            if calibration:
                signatures = {etree.tostring(r._r.rPr) for p in tf.paragraphs for r in p.runs}
                if len(signatures) != 1 or op == "text.set":
                    raise ValueError("Calibration requires uniform immutable text")
            if op == "text.set":
                if any(c in ch["text"] for c in "\r\n\v"):
                    raise ValueError("Multiline content requires_rebuild")
                tf.paragraphs[0].runs[0].text = ch["text"]
                new[oid]["text"] = ch["text"]
            else:
                st = ch["style"]
                font = tf.paragraphs[0].runs[0].font
                if "font_size_pt" in st:
                    for paragraph in tf.paragraphs:
                        for run in paragraph.runs:
                            run.font.size = Pt(st["font_size_pt"])
                if "color" in st:
                    font.color.rgb = RGBColor.from_string(st["color"])
                if "align" in st:
                    tf.paragraphs[0].alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[st["align"]]
                for side in ("left", "right", "top", "bottom"):
                    if "margin_"+side+"_pt" in st:
                        setattr(tf, "margin_"+side, Pt(st["margin_"+side+"_pt"]))
                new[oid].setdefault("style", {}).update(st)
        elif op == "geometry.set":
            if obj["kind"] not in {"text", "shape", "image"} or shape.shape_type not in {1, 13, 17}:
                raise ValueError("Geometry supports ordinary textbox/autoshape/independent picture only")
            if obj["kind"] == "image" and "source_crop" not in obj and obj.get("fit", "contain") != "stretch":
                raise ValueError("Geometry of contain/cover images requires explicit source_crop")
            box = ch["bbox"]
            new[oid]["bbox"] = box
            for field, value in zip(("left", "top", "width", "height"), bbox_to_points(box, canvas)):
                setattr(shape, field, Pt(value))
        elif op == "fill.set":
            if obj["kind"] != "shape" or shape.shape_type != 1 or obj.get("style", {}).get("gradient") or obj.get("style", {}).get("fill_alpha", 1) != 1:
                raise ValueError("Opaque ordinary autoshape only; gradient/alpha requires_rebuild")
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor.from_string(ch["color"])
            new[oid].setdefault("style", {})["fill"] = ch["color"]
        elif op == "group.translate":
            if obj["kind"] != "group" or shape.shape_type != 6:
                raise ValueError("Existing native group required")
            x = shape._element.find(qn("p:grpSpPr")).find(qn("a:xfrm"))
            ext, child_ext = x.find(qn("a:ext")), x.find(qn("a:chExt"))
            off, child_off = x.find(qn("a:off")), x.find(qn("a:chOff"))
            if dict(ext.attrib) != dict(child_ext.attrib) or dict(off.attrib) != dict(child_off.attrib):
                raise ValueError("Scaled or previously transformed groups require_rebuild")
            if any(c.get("kind") == "group" or c.get("rotation_deg", 0) for c in obj["children"]):
                raise ValueError("Nested or rotated group members unsupported")
            dx, dy = ch["delta"]
            # Keep child local coordinates untouched; translate only parent off.
            off.set("x", str(int(off.get("x")) + round(dx*canvas["width_pt"]/canvas["width"]*12700)))
            off.set("y", str(int(off.get("y")) + round(dy*canvas["height_pt"]/canvas["height"]*12700)))
            _shift(new[oid], dx, dy)
            touched.update(o["id"] for o in walk_objects(obj["children"]))
        elif op == "picture.crop":
            if obj["kind"] != "image" or shape.shape_type != 13:
                raise ValueError("Independent native picture required")
            if "source_crop" not in obj and obj.get("fit", "contain") != "stretch":
                raise ValueError("Crop requires explicit source_crop or stretch fit")
            with Image.open(resolve_asset(scene_path.parent, obj["asset"])) as im:
                iw, ih = im.size
            x, y, w, h = ch["source_crop"]
            if not 0 <= x < x+w <= iw or not 0 <= y < y+h <= ih:
                raise ValueError("Crop outside original asset pixels")
            shape.crop_left, shape.crop_top = x/iw, y/ih
            shape.crop_right, shape.crop_bottom = (iw-x-w)/iw, (ih-y-h)/ih
            new[oid]["source_crop"] = [x, y, w, h]
        else:
            raise ValueError("Unsupported operation")
    errors = validate(updated, scene_path.parent)
    if errors:
        raise ValueError("Invalid updated scene: "+"; ".join(errors[:4]))
    return prs, updated, si, sorted(touched)


def equivalent(a, b, path=""):
    """Office quantizes coordinates; 0.01pt tolerance is not a visual tolerance."""
    if type(a) is not type(b):
        return False
    if isinstance(a, dict):
        return a.keys() == b.keys() and all(equivalent(a[k], b[k], path+"/"+k) for k in a)
    if isinstance(a, list):
        return len(a) == len(b) and all(equivalent(x, y, path) for x, y in zip(a, b))
    if isinstance(a, float):
        return abs(a-b) <= (0.01 if path.endswith("_pt") else 0.00001)
    return a == b


def surgical_write(source, destination, prs, slide_index):
    """Only the edited slide part is serialized. No Presentation.save/rebuild."""
    part = prs.slides[slide_index].part
    name = str(part.partname).lstrip("/")
    with ZipFile(source) as z, ZipFile(destination, "x") as out:
        for entry in z.infolist():
            out.writestr(copy.copy(entry), part.blob if entry.filename == name else z.read(entry.filename))
    diff = part_diff(source, destination)
    if diff["added"] or diff["removed"] or set(diff["changed"]) - {name}:
        raise ValueError("Surgical edit changed unrelated package parts")
    return diff


def office_roundtrip(source, output, receipt, *, native=False, targets=None, detailed=True, export=False):
    """Own only a new copy, never Quit PowerPoint or close another presentation."""
    import shutil
    from win32com.client import Dispatch
    from runtime_env import office_guard
    if output.exists():
        raise ValueError("Office copy output must be new")
    shutil.copyfile(source, output)
    with office_guard():
        app = Dispatch("PowerPoint.Application")
        deck = None
        try:
            deck = app.Presentations.Open(str(output.resolve()), 0, 0, 0)
            deck.Save()
            deck.Close()
            deck = app.Presentations.Open(str(output.resolve()), -1, 0, 0)
            rows = []
            def visit(items, slide, parent=None):
                for i in range(1, items.Count+1):
                    s = items.Item(i)
                    if targets is not None and (slide, str(s.Name)) not in targets:
                        if int(s.Type)==6: visit(s.GroupItems, slide, str(s.Name))
                        continue
                    row = {"slide": slide, "name": str(s.Name), "type": int(s.Type), "parent": parent,
                           "bbox_pt": [float(s.Left), float(s.Top), float(s.Width), float(s.Height)]}
                    if s.HasTextFrame:
                        tf = s.TextFrame
                        tr = tf.TextRange
                        row["text"] = str(tr.Text)
                        row["text_bounds_pt"] = [float(tr.BoundLeft), float(tr.BoundTop),
                                                  float(tr.BoundWidth), float(tr.BoundHeight)]
                        row["font_size_pt"] = float(tr.Font.Size)
                        row["font_name"] = str(tr.Font.Name)
                        row["margins_pt"] = [float(getattr(tf, "Margin"+side))
                                             for side in ("Left", "Right", "Top", "Bottom")]
                        row["autosize"] = int(tf.AutoSize)
                        row["wordwrap"] = int(tf.WordWrap)
                        row["anchor"] = int(tf.VerticalAnchor)
                        row["typographic_baseline"] = "unknown"
                        row["font_resolution"] = "unverified_per_glyph"
                    if int(s.Type) == 13:
                        row["crop_pt"] = {side: float(getattr(s.PictureFormat, "Crop"+side)) for side in ("Left", "Top", "Right", "Bottom")}
                    if native:
                        from native_office_readback import read_shape
                        row["native_format"] = read_shape(s)
                    rows.append(row)
                    if int(s.Type) == 6:
                        visit(s.GroupItems, slide, str(s.Name))
            for index in range(1, deck.Slides.Count+1):
                visit(deck.Slides.Item(index).Shapes, index)
            if export:
                for index in sorted({row["slide"] for row in rows}):
                    deck.Slides.Item(index).Export(str(output.parent/f"slide-{index:03}.png"), "PNG", 1200,
                        round(1200*float(deck.PageSetup.SlideHeight)/float(deck.PageSetup.SlideWidth)))
            write_json(receipt, {"status": "recorded", "renderer": "Microsoft PowerPoint",
                                "office_version": str(app.Version), "source_sha256": sha256(source),
                                "pptx_sha256": sha256(output), "objects": rows,
                                "scope": "Actual save/close/reopen readback; not visual approval"})
        finally:
            if deck is not None:
                deck.Close()
    return properties(output) if detailed else None
