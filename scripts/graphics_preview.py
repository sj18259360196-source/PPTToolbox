"""Local SVG/PNG preview of compiled primitives, never an Office acceptance signal."""
from __future__ import annotations
import math
import numpy as np
from lxml import etree as ET
from graphics_recipe import leafs
from graphics_geometry import as_path, flatten
from graphics_paint import mix

NS = "http://www.w3.org/2000/svg"


def _triangle_heads(root, path, obj, scale):
    """Office 16 triangle preview profile; the PPT keeps its native endpoint."""
    st = obj["style"]
    ends = [key for key in ("begin_arrow", "end_arrow")
            if st.get(key, "none") != "none"]
    if not ends:
        return
    if (obj["closed"] or not st.get("line") or "line_gradient" in st
            or any(st[key] != "triangle" for key in ends)):
        raise ValueError("Preview arrow profile supports solid open triangle arrows only")
    commands = obj["commands"]
    rings = flatten(commands, .03)
    if len(rings) != 1:
        raise ValueError("Arrow preview requires one continuous path")
    points = np.asarray(rings[0], dtype=float)
    distance = float(np.linalg.norm(np.diff(points, axis=0), axis=1).sum())
    # Measured against 27 real Office specimens at 1.5, 3 and 6 pt.
    unit = max(st.get("line_width_pt", 1), 2) / scale
    factors = {"sm": 2, "med": 3, "lg": 5}
    trims = [0., 0.]
    group = ET.SubElement(root, "g", opacity=str(st.get("line_alpha", 1)))
    root.remove(path)
    group.append(path)
    path.attrib.pop("stroke-opacity", None)
    for key in ends:
        first = key == "begin_arrow"
        tip = points[0 if first else -1]
        ordered = points[1:] if first else points[-2::-1]
        previous = next((p for p in ordered if np.linalg.norm(tip-p) > 1e-8), None)
        if previous is None:
            raise ValueError("Degenerate arrow direction")
        direction = (tip-previous)/np.linalg.norm(tip-previous)
        normal = np.array([-direction[1], direction[0]])
        width = unit*factors[st.get(key+"_width", "med")]
        length = unit*factors[st.get(key+"_length", "med")]
        base = tip-direction*length
        corners = [tip, base+normal*width/2, base-normal*width/2]
        ET.SubElement(group, "polygon",
                      points=" ".join(f"{p[0]},{p[1]}" for p in corners),
                      fill="#"+st["line"])
        trims[0 if first else 1] = length
    # Trim only the preview shaft, otherwise its flat cap pokes through the tip.
    visible = max(0., distance-sum(trims))
    if visible:
        path.set("stroke-dasharray", f"{visible} {distance*2+sum(trims)}")
        path.set("stroke-dashoffset", str(-trims[0]))
    else:
        path.set("stroke", "none")


def svg(compiled):
    recipe = compiled["recipe"]
    width, height = recipe["canvas"]
    scale = recipe.get("points_per_unit", 1)
    root = ET.Element("svg", nsmap={None: NS}, viewBox=f"0 0 {width} {height}",
                      width=str(width), height=str(height))
    defs = ET.SubElement(root, "defs")
    serial = 0
    for obj in leafs(compiled["objects"]):
        if obj["kind"] == "group":
            continue
        st = obj["style"]
        path = ET.SubElement(root, "path", id=obj["id"],
                             d=" ".join(" ".join(str(v) for v in c) for c in obj["commands"]))
        path.set("fill", "#"+st["fill"] if obj["closed"] and st.get("fill") else "none")
        path.set("stroke", "#"+st["line"] if st.get("line") else "none")
        path.set("stroke-width", str(st.get("line_width_pt", 1)/scale))
        for key, svg_key in (("fill_alpha", "fill-opacity"), ("line_alpha", "stroke-opacity"),
                             ("line_cap", "stroke-linecap"), ("line_join", "stroke-linejoin")):
            if key in st:
                path.set(svg_key, str(st[key]))
        for field, paint in (("gradient", "fill"), ("line_gradient", "stroke")):
            if field not in st or (paint == "fill" and not obj["closed"]):
                continue
            g = st[field]
            serial += 1
            gid = "paint"+str(serial)
            # Office paints custom paths against visible curve extents, not handles.
            x, y, right, bottom = as_path(obj["commands"]).bbox()
            w, h = max(right-x, .01), max(bottom-y, .01)
            if g["type"] == "linear":
                a = math.radians(g["angle_deg"])
                vector = np.array([math.cos(a)/w, math.sin(a)/h])
                vector /= np.linalg.norm(vector)
                corners = np.array([[x, y], [x+w, y], [x, y+h], [x+w, y+h]])
                p = corners@vector
                center = np.array([x+w/2, y+h/2])
                start, end = [center+vector*(v-center@vector) for v in (min(p), max(p))]
                grad = ET.SubElement(defs, "linearGradient", id=gid, gradientUnits="userSpaceOnUse",
                                     x1=str(start[0]), y1=str(start[1]), x2=str(end[0]), y2=str(end[1]))
            else:
                cx, cy = g.get("center", [.5, .5])
                grad = ET.SubElement(defs, "radialGradient", id=gid, cx=str(cx), cy=str(cy),
                                     r="0.5")
            stops = g["stops"]
            # SVG engines do not implement Office's measured transfer function.
            # Densify preview-only stops; the PPT retains its original 2..5 stops.
            ts = np.unique(np.concatenate([np.linspace(a["position"], b["position"], 65)
                                          for a, b in zip(stops, stops[1:])]))
            colors = np.array([[int(s["color"][j:j+2], 16)/255 for j in (0, 2, 4)] for s in stops])
            mixed = mix(ts, [s["position"] for s in stops], colors)
            alpha = np.interp(ts, [s["position"] for s in stops], [s.get("alpha", 1) for s in stops])
            for t, color, opacity in zip(ts, mixed, alpha):
                stop = {"position": t, "color": "".join(f"{v:02X}" for v in np.rint(color*255).astype(int)),
                        "alpha": opacity}
                ET.SubElement(grad, "stop", offset=str(stop["position"]),
                              **{"stop-color": "#"+stop["color"], "stop-opacity": str(stop.get("alpha", 1))})
            path.set(paint, "url(#"+gid+")")
        _triangle_heads(root, path, obj, scale)
    return ET.tostring(root, encoding="unicode")


def png(compiled, width=1000):
    import resvg_py
    return resvg_py.svg_to_bytes(svg_string=svg(compiled), width=width, skip_system_fonts=True)


def readback(pptx):
    """Conservative observed signature, independent of the recipe and its expected paths."""
    from pptx import Presentation
    from pptx.oxml.ns import qn
    result = {}
    def visit(shapes, parent=None):
        for index, shape in enumerate(shapes):
            if shape.name in result:
                raise ValueError("Ambiguous duplicate native object names")
            sp = shape._element
            # Normalize attribute order, namespaces, metadata IDs and equivalent numbers.
            def node(e):
                tag = ET.QName(e).localname
                attrs = {}
                for k, v in sorted(e.attrib.items()):
                    try:
                        v = round(float(v), 6)
                    except ValueError:
                        pass
                    attrs[ET.QName(k).localname] = v
                if tag == "path":
                    # PowerPoint removes these explicit DrawingML defaults on save.
                    attrs.setdefault("fill", "norm")
                    attrs.setdefault("stroke", 1.0)
                children = [v for c in e if (v := node(c)) is not None]
                return [tag, attrs, children]
            if shape.shape_type == 6:
                content = node(sp.find(qn("p:grpSpPr")))
            else:
                content = node(sp.find(qn("p:spPr")))
            theme = sp.find(qn("p:style"))
            result[shape.name] = {"parent": parent, "index": index, "type": int(shape.shape_type),
                                  "content": content, "theme": node(theme) if theme is not None else None,
                                  "text": shape.text if shape.has_text_frame else ""}
            if shape.shape_type == 6:
                visit(shape.shapes, shape.name)
    prs = Presentation(pptx)
    if len(prs.slides) != 1:
        raise ValueError("Graphics draft readback requires a single slide")
    background = prs.slides[0]._element.find(".//"+qn("p:bg"))
    result["$slide"] = {"width": int(prs.slide_width), "height": int(prs.slide_height),
                        "background": ET.tostring(background, method="c14n").decode()
                        if background is not None else None}
    visit(prs.slides[0].shapes)
    return result
