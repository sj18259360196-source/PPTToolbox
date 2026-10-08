"""Explicit code-point ranges, immutable text, bounded run/paragraph formatting."""
import copy
import hashlib
from pptx.oxml.ns import qn
from pptx.util import Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement


def text_hash(text):
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def paragraphs(obj):
    if "paragraphs" in obj:
        return copy.deepcopy(obj["paragraphs"])
    result=[{"runs":[],"style":{}}]
    for run in obj.get("runs",[{"text":obj.get("text","")}]):
        for i,text in enumerate(run["text"].split("\n")):
            if i:result.append({"runs":[],"style":{}})
            result[-1]["runs"].append({**run,"text":text})
    return result


def run_style(run,style):
    font=run.font
    if "font" in style:font.name=style["font"]
    if "font_east_asia" in style:
        rpr=run._r.get_or_add_rPr()
        ea=rpr.find(qn("a:ea"))
        if ea is None:
            ea=OxmlElement("a:ea")
            rpr.insert_element_before(ea,"a:cs","a:sym","a:hlinkClick","a:hlinkMouseOver","a:extLst")
        ea.set("typeface",style["font_east_asia"])
    if "font_size_pt" in style:font.size=Pt(style["font_size_pt"])
    if "bold" in style:font.bold=style["bold"]
    if "color" in style:font.color.rgb=RGBColor.from_string(style["color"])
    if "char_spacing_pt" in style:
        run._r.get_or_add_rPr().set("spc",str(round(style["char_spacing_pt"]*100)))


def paragraph_style(para,style):
    if "align" in style:
        para.alignment={"left":PP_ALIGN.LEFT,"center":PP_ALIGN.CENTER,"right":PP_ALIGN.RIGHT}[style["align"]]
    for key,attr in (("space_before_pt","space_before"),("space_after_pt","space_after"),("line_spacing_pt","line_spacing")):
        if key in style:setattr(para,attr,Pt(style[key]))


def apply(obj,shape,change):
    from jsonschema import Draft202012Validator
    from native_capabilities import TEXT_RANGE
    Draft202012Validator(TEXT_RANGE).validate(change)
    if obj["kind"]!="text" or not shape.has_text_frame:
        raise ValueError("Range formatting requires native text")
    if obj.get("style",{}).get("native_format") or shape._element.xpath(
            ".//a:fld|.//a:br|.//a:prstTxWarp|.//a:effectLst/*|.//a:effectDag|.//a:sp3d|.//a:hlinkClick"):
        raise ValueError("Fields, soft breaks, hyperlinks and advanced text effects are unsupported")
    ps=paragraphs(obj)
    text="\n".join("".join(r["text"] for r in p["runs"]) for p in ps)
    if any(ord(c)>0xFFFF or c in "\r\v" for c in text):
        raise ValueError("First range contract supports BMP text and paragraph LF only")
    if text!=shape.text_frame.text or text_hash(text)!=change["text_sha256"]:
        raise ValueError("Actual/source text or frozen hash differs")
    pi=change["paragraph"]
    if not 0<=pi<len(ps):raise ValueError("Paragraph index outside actual text")
    para=shape.text_frame.paragraphs[pi]
    content="".join(r["text"] for r in ps[pi]["runs"])
    start,end=change["start"],change["end"]
    if not 0<=start<end<=len(content) or content[start:end]!=change["expected_fragment"]:
        raise ValueError("Range/expected fragment mismatch; empty paragraphs are not selectable")
    if change["scope"]=="paragraph":
        if (start,end)!=(0,len(content)):
            raise ValueError("Paragraph style requires the entire selected paragraph")
        paragraph_style(para,change["style"])
        ps[pi].setdefault("style",{}).update(change["style"])
    else:
        # Split only intersected runs. Unselected XML and paragraph properties stay intact.
        cursor=0
        for run in list(para.runs):
            length=len(run.text);left=max(start-cursor,0);right=min(end-cursor,length)
            cursor+=length
            if left>=right:continue
            original=run._r
            parent=original.getparent();position=list(parent).index(original)
            pieces=[(run.text[:left],False),(run.text[left:right],True),(run.text[right:],False)]
            parent.remove(original)
            from pptx.text.text import _Run
            for value,selected in pieces:
                if not value:continue
                node=copy.deepcopy(original);node.find(qn("a:t")).text=value
                parent.insert(position,node);position+=1
                if selected:run_style(_Run(node,para),change["style"])
        cursor=0;updated=[]
        for run in ps[pi]["runs"]:
            n=len(run["text"]);left=max(start-cursor,0);right=min(end-cursor,n);cursor+=n
            if left>=right:updated.append(run);continue
            for value,selected in ((run["text"][:left],False),(run["text"][left:right],True),(run["text"][right:],False)):
                if value:updated.append({**run,**(change["style"] if selected else {}),"text":value})
        ps[pi]["runs"]=updated
    new=copy.deepcopy(obj)
    new.pop("text",None);new.pop("runs",None)
    new["paragraphs"]=ps
    return new


def read_com(tf):
    """BMP-only diagnostic; raw per-character values, independent of scene intent."""
    text=str(tf.TextRange.Text)
    if len(text)>2000 or any(ord(c)>0xFFFF for c in text):
        raise ValueError("Detailed character readback limited to 2000 BMP characters")
    from win32com.client import Dispatch
    # Office typelib exposes these indexed members as PROPERTYGET, not methods.
    # Fixed DISPIDs from TextRange2: Paragraphs=4, Characters=7. No caller-supplied members.
    def indexed(text_range,member,start,length):
        return Dispatch(text_range._oleobj_.Invoke(member,0,2,True,start,length))
    result=[]
    for i in range(1,int(indexed(tf.TextRange,4,-1,-1).Count)+1):
        para=indexed(tf.TextRange,4,i,1)
        pf=para.ParagraphFormat
        chars=[]
        raw=str(para.Text)
        for j in range(1,len(raw)+1):
            r=indexed(para,7,j,1)
            f=r.Font
            chars.append({"text":str(r.Text),"font":str(f.Name),"east_asia":str(f.NameFarEast),
                          "size_pt":float(f.Size),"bold":int(f.Bold),"color_bgr":int(f.Fill.ForeColor.RGB),
                          "spacing_pt":float(f.Spacing)})
        result.append({"text":raw,"characters":chars,"alignment":int(pf.Alignment),
                       "before_pt":float(pf.SpaceBefore),"after_pt":float(pf.SpaceAfter),
                       "within":float(pf.SpaceWithin),"line_rule_within":int(pf.LineRuleWithin),
                       "line_rule_before":int(pf.LineRuleBefore),"line_rule_after":int(pf.LineRuleAfter)})
    return result


def validate_readback(changes,before,after,rebuilt):
    """Compare independently read effective characters, never expected scene runs."""
    def table(receipt):
        return {r["name"]:r.get("native_format",{}).get("character_ranges",{}) for r in receipt["objects"]}
    a,b,c=map(table,(before,after,rebuilt))
    result=[]
    for change in changes:
        if change["op"]!="text.range":continue
        name=change["id"]
        rows=[t.get(name,{}) for t in (a,b,c)]
        if any(r.get("status")!="read" for r in rows):
            raise ValueError("Character style measurement unknown; requires review")
        old,new,rep=(r["value"] for r in rows)
        if new!=rep:raise ValueError("Effective character/paragraph styles differ from independent rebuild")
        if len(old)!=len(new):raise ValueError("Paragraph count changed")
        fields={"font":"font","font_east_asia":"east_asia","font_size_pt":"size_pt","bold":"bold",
                "color":"color_bgr","char_spacing_pt":"spacing_pt"}
        pfields={"align":"alignment","space_before_pt":"before_pt","space_after_pt":"after_pt",
                 "line_spacing_pt":"within"}
        for pi,(left,right) in enumerate(zip(old,new)):
            if left["text"]!=right["text"] or len(left["characters"])!=len(right["characters"]):
                raise ValueError("Actual paragraph content changed")
            for j,(lc,rc) in enumerate(zip(left["characters"],right["characters"])):
                allowed={fields[k] for k in change["style"]} if change["scope"]=="run" and pi==change["paragraph"] and change["start"]<=j<change["end"] else set()
                if {k:v for k,v in lc.items() if k not in allowed}!={k:v for k,v in rc.items() if k not in allowed}:
                    raise ValueError("Unselected effective character formatting changed")
                if allowed:
                    for key,value in change["style"].items():
                        wanted=(-1 if value else 0) if key=="bold" else (
                            int(value[0:2],16)+256*int(value[2:4],16)+65536*int(value[4:6],16)
                            if key=="color" else value)
                        observed=rc[fields[key]]
                        if (abs(observed-wanted)>.01 if isinstance(wanted,(float,int)) else observed!=wanted):
                            raise ValueError("Requested character format not observed")
            allowed={pfields[k] for k in change["style"]} if change["scope"]=="paragraph" and pi==change["paragraph"] else set()
            if "within" in allowed:allowed.add("line_rule_within")
            if "before_pt" in allowed:allowed.add("line_rule_before")
            if "after_pt" in allowed:allowed.add("line_rule_after")
            skip=allowed|{"characters"}
            if {k:v for k,v in left.items() if k not in skip}!={k:v for k,v in right.items() if k not in skip}:
                raise ValueError("Unselected paragraph formatting changed")
            if change["scope"]=="paragraph" and pi==change["paragraph"]:
                for key,value in change["style"].items():
                    wanted={"left":1,"center":2,"right":3}[value] if key=="align" else value
                    if abs(right[pfields[key]]-wanted)>.01:
                        raise ValueError("Requested paragraph format not observed")
                    rule={"space_before_pt":"line_rule_before","space_after_pt":"line_rule_after",
                          "line_spacing_pt":"line_rule_within"}.get(key)
                    if rule and right[rule]!=0:
                        raise ValueError("Paragraph spacing is not measured in requested points")
        result.append({"id":name,"status":"passed","scope":"Actual effective characters and paragraph formatting; layout reviewed separately"})
    return result
