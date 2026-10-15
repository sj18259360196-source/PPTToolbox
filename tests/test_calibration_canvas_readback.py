"""Evaluate-level canvas regressions. Synthetic receipts are not Office proof.

Accepted point values below come from the 2026-10-16 serial Office probes.
The tests rebuild fresh specimens, bind synthetic receipts and call evaluate;
they do not call or reproduce the physical-canvas validation helper.
"""
import copy
import json
from pathlib import Path
import sys
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from PIL import Image
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
from common import read_json, sha256, write_json
import text_calibration as calibration


@pytest.fixture
def calibration_environment(monkeypatch):
    identity = {'available': True, 'sha256': 'e' * 64, 'requested': ['Arial']}
    monkeypatch.setattr(calibration, 'environment', lambda fonts: copy.deepcopy(identity))
    return identity


def specimen(tmp_path, physical, office_physical):
    """Build real OOXML, then attach explicitly synthetic measurement evidence."""
    source = tmp_path / 'source.scene.json'
    canvas = {'width': 960, 'height': 540, 'width_pt': physical[0], 'height_pt': physical[1],
              'mapping': 'explicit_stretch', 'background': 'FFFFFF'}
    obj = {'id': 'sample', 'kind': 'text', 'role': 'content', 'bbox': [40, 40, 400, 70],
           'text': 'Canvas sample', 'style': {'font': 'Arial', 'font_size_pt': 20},
           'editability': 'text', 'evidence': {'status': 'synthetic', 'note': 'Controlled canvas regression'}}
    write_json(source, {'version': '1.0', 'canvas': canvas,
                        'slides': [{'id': 'page', 'objects': [obj]}]})
    plan_dir = tmp_path / 'plan'
    calibration.plan(source, plan_dir)
    plan = read_json(plan_dir / 'plan.json')
    office_dir = tmp_path / 'synthetic-office'
    office_dir.mkdir()
    slides, bounds = [], []
    presentation = Presentation(plan_dir / 'specimens.pptx')
    for target in plan['targets']:
        path = office_dir / f"slide-{target['index']:03}.png"
        Image.new('RGB', (960, 540), 'white').save(path)
        slides.append({'index': target['index'], 'file': path.name, 'sha256': sha256(path),
                       'width': 960, 'height': 540})
        shape = next(s for s in presentation.slides[target['index'] - 1].shapes if s.name == target['sample_id'])
        box = [shape.left / 12700, shape.top / 12700, shape.width / 12700, shape.height / 12700]
        text_bounds = [box[0], box[1], min(100, box[2] - 1), min(20, box[3] - 1)]
        bounds.append({'slide': target['index'], 'id': target['sample_id'], 'rotation': 0,
                       'box_pt': box, 'bound_pt': text_bounds,
                       'calibration': {'format': 'office-text-calibration/1', 'text': obj['text'],
                           'auto_size': 0, 'errors': [], 'truncated': False,
                           'runs': [{'text': obj['text'], 'font': 'Arial', 'office_font_size_pt': 20,
                                     'baseline_offset_fraction': 0, 'baseline_shift_pt': 0}],
                           'lines': [{'index': 1, 'bound_pt': text_bounds, 'bound_height_pt': text_bounds[3],
                                      'baseline_y_pt': None}]}})
    receipt = {'status': 'passed', 'renderer': 'Microsoft PowerPoint',
               'scope': 'SYNTHETIC TEST ONLY', 'probe_only': False,
               'pptx_sha256': sha256(plan_dir / 'specimens.pptx'), 'canvas_pt': office_physical,
               'slides': slides, 'text_bounds': bounds}
    # Negative cases deliberately inject NaN/Infinity tokens accepted by the JSON
    # reader. The normal write_json helper correctly refuses to author them.
    (office_dir / 'office-render.json').write_text(json.dumps(receipt, allow_nan=True), encoding='utf-8')
    return source, plan_dir, office_dir, tmp_path / 'measurement.json'


def evaluate(paths):
    _, plan, office, output = paths
    return calibration.evaluate(plan, office, output)


@pytest.mark.parametrize(('physical', 'readback'), [
    ((720, 405.21531100478467), (720, 405.25)),
    ((960, 540.2870813397129), (960, 540.25)),
    ((960, 540.5405405405405), (960, 540.5)),
    ((960.0624, 540.1249), (960, 540.125)),
    ((960.0626, 540.1251), (960.125, 540.125)),
    ((960.1874, 540.2499), (960.125, 540.25)),
    ((960.1876, 540.2501), (960.25, 540.25)),
    ((960.1249, 540.0624), (960.125, 540)),
    ((960.1251, 540.0626), (960.125, 540.125)),
    ((960.3751, 540.3751), (960.375, 540.375)),
    # Nearest integer-EMU inputs on opposing sides of 1/8-point midpoints.
    ((12192793 / 12700, 540.06252), (960, 540.125)),
    ((12194382 / 12700, 6860381 / 12700), (960.25, 540.125)),
])
def test_evaluate_accepts_recorded_office_canvas_values(tmp_path, calibration_environment, physical, readback):
    paths = specimen(tmp_path, list(physical), list(readback))
    outcome = evaluate(paths)
    report = read_json(paths[3])
    assert outcome['status'] == 'advisory'
    assert report['visual_review'] == 'pending'
    assert report['automatic_changes'] is False
    assert report['binding']['receipt_sha256'] == sha256(paths[2] / 'office-render.json')
    validation = report['canvas_validation']
    assert outcome['canvas_validation'] == validation
    assert validation['expected_pt'] == list(physical)
    assert validation['office_pt'] == list(readback)
    assert validation['package_binding'] == 'exact_source_emu'
    assert validation['quantization_step_pt'] == .125
    assert validation['numeric_tolerance_pt'] == .001
    presentation = Presentation(paths[1] / 'specimens.pptx')
    assert validation['package_emu'] == [presentation.slide_width, presentation.slide_height]


def test_evaluate_keeps_direct_precision_readback_compatible(tmp_path, calibration_environment):
    paths = specimen(tmp_path, [960, 540], [960.00003, 540])
    assert evaluate(paths)['status'] == 'advisory'


@pytest.mark.parametrize('readback', [
    [960.02, 540], [960, 540.02], [961, 540], [960, 541],
    [float('nan'), 540], [float('inf'), 540], [960, float('-inf')],
    [True, 540], [960, False], [960, 0], [960, -1], [960], None,
])
def test_evaluate_rejects_invalid_or_unexplained_canvas_readback(tmp_path, calibration_environment, readback):
    paths = specimen(tmp_path, [960, 540], readback)
    with pytest.raises(ValueError, match='canvas'):
        evaluate(paths)
    assert not paths[3].exists()


@pytest.mark.parametrize(('physical', 'wrong_readback'), [
    ((960.0626, 540.1251), (960, 540.125)),
    ((960.1874, 540.2499), (960.25, 540.25)),
    ((12192793 / 12700, 540.06252), (960, 540)),
    ((12194382 / 12700, 6860381 / 12700), (960.25, 540.25)),
])
def test_evaluate_rejects_wrong_adjacent_grid_point(tmp_path, calibration_environment, physical, wrong_readback):
    paths = specimen(tmp_path, list(physical), list(wrong_readback))
    with pytest.raises(ValueError, match='canvas'):
        evaluate(paths)
    assert not paths[3].exists()


def test_evaluate_rejects_true_package_resizing_even_when_receipt_hashes_match(tmp_path, calibration_environment):
    paths = specimen(tmp_path, [960, 540], [960, 540.125])
    _, plan_dir, office_dir, output = paths
    pptx = plan_dir / 'specimens.pptx'
    with ZipFile(pptx) as package:
        entries = [(entry, package.read(entry.filename)) for entry in package.infolist()]
    with ZipFile(pptx, 'w') as package:
        for entry, data in entries:
            if entry.filename == 'ppt/presentation.xml':
                root = ET.fromstring(data)
                size = root.find('{http://schemas.openxmlformats.org/presentationml/2006/main}sldSz')
                size.set('cy', str(int(size.attrib['cy']) + 1588))
                data = ET.tostring(root, encoding='utf-8', xml_declaration=True)
            package.writestr(entry, data)
    plan = read_json(plan_dir / 'plan.json')
    plan['specimens_sha256'] = sha256(pptx)
    write_json(plan_dir / 'plan.json', plan)
    receipt = read_json(office_dir / 'office-render.json')
    receipt['pptx_sha256'] = sha256(pptx)
    write_json(office_dir / 'office-render.json', receipt)
    with pytest.raises(ValueError, match='canvas'):
        evaluate(paths)
    assert not output.exists()


def test_evaluate_rejects_plan_canvas_edit_that_disagrees_with_frozen_source(tmp_path, calibration_environment):
    paths = specimen(tmp_path, [960, 540], [960, 540.125])
    plan_path = paths[1] / 'plan.json'
    plan = read_json(plan_path)
    plan['canvas']['height_pt'] = 540.125
    write_json(plan_path, plan)
    with pytest.raises(ValueError, match='canvas'):
        evaluate(paths)
    assert not paths[3].exists()
