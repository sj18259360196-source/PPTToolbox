"""Geometry and policy contracts; real Office evidence is tested separately."""
import copy
import json
from pathlib import Path
import pytest
from PIL import Image,ImageDraw
from test_tools import ROOT
from shape_construction import construct,compare_contours
from graphics_recipe import compile_recipe
from native_nodes import topology
from local_edit import prepare,surgical_write,semantic_nodes
from common import write_json
from build_pptx import build

ROOT=Path(__file__).resolve().parents[1]

def polygon():
    return dict(mode='rounded_polygon',id='demo',canvas=[200,200],style={'fill':'69ACED','line':None},
                points=[[20,20],[160,20],[160,160],[20,160]],corner_inset=10)

def test_parameter_recipe_retains_cubic_editability():
    a=polygon();old=copy.deepcopy(a);result=construct(a)
    assert a==old
    commands=result['objects'][0]['commands']
    assert sum(c[0]=='C' for c in commands)==4
    assert topology(commands)['holes']==0
    assert result['office']=='not_run'
    assert compile_recipe(result['recipe'])['objects']==result['objects']

@pytest.mark.parametrize('points,inset',[
    ([[10,10],[10,10],[90,90]],1),
    ([[10,10],[90,90],[90,10],[10,90]],0),
    ([[10,10],[10.01,10],[90,90]],0),
    ([[10,10],[90,10],[50,40],[90,90],[10,90]],1),
    ([[10,10],[90,10],[90,90],[10,90]],40)])
def test_bad_polygon_rejected(points,inset):
    a=polygon();a.update(points=points,corner_inset=inset)
    with pytest.raises(ValueError):construct(a)

def test_concave_polygon_can_remain_sharp():
    a=polygon();a.update(points=[[10,10],[90,10],[50,40],[90,90],[10,90]],corner_inset=0)
    assert construct(a)['objects']

def masks(dx=0):
    r=Image.new('L',(80,80));c=r.copy()
    ImageDraw.Draw(r).rectangle([20,20,40,40],fill=255)
    ImageDraw.Draw(c).rectangle([20+dx,20,40+dx,40],fill=255)
    return r,c

def test_contour_translation_report_not_hidden_by_alignment():
    r,c=masks(5);out=compare_contours(r,c,r,c)
    assert out['raw']['max_px']==5
    assert out['centroid_translation_candidate_to_reference_px']==[-5,0]
    assert out['centroid_aligned']['max_px']==0
    assert out['visual_review']=='not_assessed' and 'score' not in out

def test_masks_define_objects_independently_of_text_and_color():
    r,c=masks();a=r.convert('RGB');b=c.convert('RGB')
    ImageDraw.Draw(b).text((45,10),'OTHER',fill='red')
    assert compare_contours(a,b,r,c)['raw']['max_px']==0
    gray=r.copy();gray.putpixel((0,0),128)
    with pytest.raises(ValueError,match='binary'):compare_contours(a,b,gray,c)
    with pytest.raises(ValueError,match='frame'):compare_contours(a,b.resize((40,40)),r,c)
    with pytest.raises(ValueError,match='Insufficient'):compare_contours(a,b,r,c,Image.new('L',r.size,255))

def node_fixture(tmp_path,grouped=True):
    Image.new('RGB',(200,200),'white').save(tmp_path/'ref.png')
    path={'id':'page.body.curve','kind':'path','closed':True,
          'commands':[['M',20,20],['C',50,20,70,40,100,20],['L',100,100],['L',20,100],['Z']],
          'style':{'fill':'82ACF4','line':None,'fill_alpha':.6,'gradient':{'type':'linear','angle_deg':90,
                   'stops':[{'position':0,'color':'82ACF4','alpha':1},{'position':1,'color':'AB83F7','alpha':.8}]}}}
    protected={'id':'page.body.protected','kind':'shape','geometry':'rect','bbox':[120,20,30,30],'style':{'fill':'F09824','line':None}}
    objects=[{'id':'page.body.group','kind':'group','children':[path,protected]}] if grouped else [path,protected]
    scene={'version':'1.0','canvas':{'width':200,'height':200,'width_pt':200,'height_pt':200,'mapping':'uniform'},
           'slides':[{'id':'page','reference':'ref.png','objects':objects}]}
    from workflow_scene import EDIT
    from common import walk_objects
    for o in walk_objects(objects):
        o['editability']=EDIT[o['kind']]
        o['evidence']={'status':'inferred','note':'Synthetic geometry fixture'}
    sp=tmp_path/'scene.json';ppt=tmp_path/'source.pptx';write_json(sp,scene);build(sp,ppt)
    change={'op':'native.nodes','id':path['id'],'max_displacement':5,
            'edits':[{'command_index':1,'point_index':1,'expected':[70,40],'delta':[0,3]}]}
    return sp,ppt,change

def test_grouped_node_control_edit_protects_neighbors_and_fill(tmp_path):
    sp,ppt,ch=node_fixture(tmp_path)
    before=semantic_nodes(ppt)
    prs,scene,si,touched=prepare(sp,ppt,'page',{ch['id']},[ch])
    out=tmp_path/'edited.pptx';surgical_write(ppt,out,prs,si)
    after=semantic_nodes(out)
    assert touched==[ch['id']]
    assert before['0/page.body.protected']==after['0/page.body.protected']
    assert before['0/page.body.group']==after['0/page.body.group']
    p=scene['slides'][0]['objects'][0]['children'][0]
    assert p['commands'][1][3:5]==[70,43]
    assert p['style']['fill_alpha']==.6
    assert topology(p['commands'])['rings']==1

@pytest.mark.parametrize('mutation',['stale','large','outside','duplicate','unauthorized'])
def test_node_guard(tmp_path,mutation):
    sp,ppt,ch=node_fixture(tmp_path)
    allowed={ch['id']}
    if mutation=='stale':ch['edits'][0]['expected']=[71,40]
    if mutation=='large':ch['edits'][0]['delta']=[0,6]
    if mutation=='outside':ch['edits'][0].update(command_index=0,point_index=0,expected=[20,20],delta=[-1,0])
    if mutation=='duplicate':ch['edits'].append(copy.deepcopy(ch['edits'][0]))
    if mutation=='unauthorized':allowed=set()
    with pytest.raises(ValueError):prepare(sp,ppt,'page',allowed,[ch])

def test_new_tools_discovered_and_override_respected(tmp_path):
    from toolbox_manager.service import Manager
    from toolbox_manager.mcp import MCP
    from toolbox_manager.policy import PolicyDenied
    from toolbox_manager.contracts import schemas
    m=Manager(tmp_path/'data',root=ROOT,home=tmp_path/'home')
    names={n for n,_,_ in MCP(tmp_path/'rpc',root=ROOT,home=tmp_path/'rh').specs()}
    assert {'graphics_construct','graphics_compare_contours','graphics_read_properties','graphics_boolean_trials'}<=names
    assert 'native.nodes' in json.dumps(schemas(ROOT)['patch'])
    with pytest.raises(PolicyDenied):m.graphics('construct',polygon())
    assert m.graphics('construct',polygon(),source='owner')['objects']
    m.toggle('tool','graphics.construct',False)
    with pytest.raises(PolicyDenied):m.graphics('construct',polygon(),source='owner')


def test_mixed_office_transparency_is_unknown():
    from native_office_readback import read_fill
    from types import SimpleNamespace
    result=read_fill(SimpleNamespace(Type=-2,Transparency=-2147483648))
    assert result['alpha']['status']=='unknown'
    assert result['transparency']['status']=='unknown'


def test_topology_defaults_equivalent_but_material_changes_protected(tmp_path):
    from native_topology import protected_check
    from pptx import Presentation
    sp,ppt,ch=node_fixture(tmp_path)
    prs=Presentation(ppt);path=prs.slides[0].shapes[0].shapes[0]._element.xpath('.//a:path')[0]
    path.attrib.pop('fill',None);path.attrib.pop('stroke',None)
    out=tmp_path/'default.pptx';prs.save(out)
    assert protected_check(ppt,out,0,set(),set())
    path.set('fill','none');prs.save(out)
    assert not protected_check(ppt,out,0,set(),set())


def test_readback_policy_and_hash_rejection_without_office(tmp_path):
    from toolbox_manager.service import Manager
    from toolbox_manager.policy import PolicyDenied
    from common import sha256
    from shape_evidence import read_properties
    project=tmp_path/'project';project.mkdir()
    sp,ppt,ch=node_fixture(project)
    m=Manager(tmp_path/'data',root=ROOT,home=tmp_path/'home')
    a=dict(project=str(project),pptx='source.pptx',pptx_sha256=sha256(ppt),targets=[{'slide':1,'id':ch['id']}])
    with pytest.raises(PolicyDenied):m.graphics('read_properties',a,source='owner')
    m.authorize_project(project,[])
    m.toggle('tool','office.edit-readback',False)
    with pytest.raises(PolicyDenied):m.graphics('read_properties',a,source='owner')
    with pytest.raises(ValueError,match='Stale'):read_properties(project,{**a,'pptx_sha256':'0'*64})
    with pytest.raises(ValueError):read_properties(project,{**a,'pptx':'../outside.pptx'})
    with pytest.raises(ValueError,match='Duplicate'):read_properties(project,{**a,'targets':a['targets']*2})
    assert not (project/'assets/native-readback').exists()


def test_transformed_group_and_hole_topology_are_guarded(tmp_path):
    from pptx import Presentation
    sp,ppt,ch=node_fixture(tmp_path)
    prs=Presentation(ppt);prs.slides[0].shapes[0].rotation=10;prs.save(ppt)
    with pytest.raises(ValueError,match='Transformed group'):prepare(sp,ppt,'page',{ch['id']},[ch])
    outer=[['M',0,0],['L',100,0],['L',100,100],['L',0,100],['Z']]
    inner=[['M',20,20],['L',20,80],['L',80,80],['L',80,20],['Z']]
    assert topology(outer+inner)['holes']==1
    inner[0]=['M',0,0]
    with pytest.raises(ValueError):topology(outer+inner)
