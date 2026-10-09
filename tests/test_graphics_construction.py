"""Construction, managed access and native package tests; actual Office is separate."""
import base64
import copy
import io
import json
import sys
from pathlib import Path

import numpy as np
import pytest
from jsonschema import ValidationError
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT), str(ROOT/"scripts")]
from graphics_geometry import flatten, offset, reverse_edge, transform
from graphics_recipe import compile_recipe, SCHEMA, digest
from graphics_gradient import fit_gradient, predict
from graphics_preview import png, readback


def recipe():
    return {"format": "graphics-recipe/1", "id": "test", "canvas": [400, 300],
            "paths": [
                {"id": "a", "commands": [["M", 40, 100], ["C", 60, 20, 200, 20, 240, 100]],
                 "closed": False, "visible": False, "style": {"fill": None, "line": "23765A", "line_width_pt": 2}},
                {"id": "b", "commands": [["M", 40, 180], ["C", 60, 100, 200, 100, 240, 180]],
                 "closed": False, "visible": False, "style": {"fill": None, "line": "23765A"}}],
            "curve_groups": [{"id": "curves", "mode": "interpolate", "source": "a", "target": "b", "count": 8}],
            "order": ["curves"]}


def shared():
    return {"format": "graphics-recipe/1", "id": "river", "canvas": [400, 300],
            "edges": [
                {"id": "bank", "commands": [["M", 20, 150], ["C", 100, 80, 220, 220, 380, 150]]},
                {"id": "top", "commands": [["M", 380, 150], ["L", 380, 20], ["L", 20, 20], ["L", 20, 150]]},
                {"id": "bottom", "commands": [["M", 20, 150], ["L", 20, 280], ["L", 380, 280], ["L", 380, 150]]}],
            "faces": [
                {"id": "land", "loops": [[{"edge": "bank"}, {"edge": "top"}]], "style": {"fill": "23765A", "line": None}},
                {"id": "water", "loops": [[{"edge": "bank", "reverse": True}, {"edge": "bottom"}]],
                 "style": {"fill": "ABD5B8", "line": None}}],
            "instances": [{"id": "reflection", "source": "land", "matrix": [1, 0, 0, -.3, 0, 290],
                           "style_override": {"fill": "23765A", "fill_alpha": .25}}],
            "order": ["land", "water", "reflection"]}


def test_interpolation_native_and_deterministic():
    r = recipe()
    a = compile_recipe(r)
    b = compile_recipe(r)
    assert a == b
    assert len(a["objects"][0]["children"]) == 8
    assert a["objects"][0]["children"][0]["commands"] == r["paths"][0]["commands"]
    assert a["objects"][0]["children"][-1]["commands"] == r["paths"][1]["commands"]
    assert len(png(a)) > 500


def test_schema_rejects_unknown_and_inapplicable_fields():
    from jsonschema import Draft202012Validator
    r = recipe()
    r["anything"] = True
    assert not Draft202012Validator(SCHEMA).is_valid(r)
    r = recipe()
    r["curve_groups"][0]["matrix_step"] = [1, 0, 0, 1, 0, 0]
    with pytest.raises(ValueError, match="not applicable"):
        compile_recipe(r)
    del r["curve_groups"][0]["target"]
    assert not Draft202012Validator(SCHEMA).is_valid(r)


@pytest.mark.parametrize("case", ["nan", "duplicate", "order", "cycle", "missing", "singular", "closure"])
def test_negative_recipe_cases(case):
    r = shared()
    if case == "nan":
        r["edges"][0]["commands"][0][1] = float("nan")
    elif case == "duplicate":
        r["edges"][0]["id"] = "land"
    elif case == "order":
        r["order"].append("land")
    elif case == "cycle":
        r["instances"][0]["source"] = "reflection"
    elif case == "missing":
        r["instances"][0]["source"] = "absent"
    elif case == "singular":
        r["instances"][0]["matrix"] = [0, 0, 0, 0, 0, 0]
    elif case == "closure":
        r["edges"][1]["commands"][-1][2] += 1
    with pytest.raises((ValueError, ValidationError)):
        compile_recipe(r)


def test_reversing_cubic_swaps_handles():
    c = [["M", 0, 0], ["C", 10, 20, 30, 40, 50, 60]]
    rev = reverse_edge(c)
    assert rev == [["M", 50, 60], ["C", 30, 40, 10, 20, 0, 0]]
    assert reverse_edge(rev) == c


def test_shared_boundary_propagates_and_instances_lock():
    r = shared()
    before = compile_recipe(r)
    r["edges"][0]["commands"][1][2] += 10
    after = compile_recipe(r, before, before["objects"])
    assert set(after["changed"]) == {"land", "water", "reflection"}
    assert len(after["shared_edges"]["bank"]) == 2
    assert after["shared_edges"]["bank"][0][1] != after["shared_edges"]["bank"][1][1]
    r["instances"][0]["state"] = "locked"
    locked = compile_recipe(r, before, before["objects"])
    assert locked["objects"][-1] == before["objects"][-1]
    assert set(locked["changed"]) == {"land", "water"}


def test_detached_instance_survives_source_removal():
    r = shared()
    old = compile_recipe(r)
    r["faces"] = []
    r["instances"][0]["state"] = "detached"
    r["order"] = ["reflection"]
    out = compile_recipe(r, old, old["objects"])
    assert out["objects"][0] == old["objects"][-1]
    assert out["dependencies"]["reflection"] == []


def test_manual_scene_edits_refuse_regeneration():
    r = shared()
    old = compile_recipe(r)
    edited = copy.deepcopy(old["objects"])
    edited[0]["style"]["fill"] = "FF0000"
    with pytest.raises(ValueError, match="Manual edit"):
        compile_recipe(r, old, edited)
    with pytest.raises(ValueError, match="observed"):
        compile_recipe(r, old)


def test_offset_line_and_cubic_distance():
    line, report = offset([["M", 20, 20], ["L", 120, 20]], 10)
    assert np.allclose(flatten(line)[0], [[20, 30], [120, 30]])
    curve, report = offset(recipe()["paths"][0]["commands"], 10, .2)
    assert report["hausdorff_error"] <= .2
    assert report["segments"] <= 128
    assert any(c[0] == "C" for c in curve)


@pytest.mark.parametrize("join", ["round", "bevel", "miter"])
def test_offset_corner_policies(join):
    commands, report = offset([["M", 20, 20], ["L", 80, 20], ["L", 80, 80]], -10, .2, join)
    assert report["hausdorff_error"] <= .2


def test_affine_repeat_and_matrix_validation():
    r = recipe()
    r["curve_groups"] = [{"id": "curves", "mode": "affine_repeat", "source": "a",
                          "count": 3, "matrix_step": [1, 0, 0, 1, 0, 20]}]
    out = compile_recipe(r)
    assert out["objects"][0]["children"][2]["commands"][0] == ["M", 40, 140]


def test_gradient_fit_on_held_out_nonsquare_data():
    w, h = 180, 90
    y, x = np.mgrid[:h, :w]
    xy = np.column_stack(((x.ravel()+.5)/w, (y.ravel()+.5)/h))
    gradient = {"type": "linear", "angle_deg": 37, "stops": [
        {"position": 0, "color": "21654B"}, {"position": 1, "color": "CBDEAC"}]}
    data = np.rint(predict(xy, w, h, gradient)*255).astype("uint8").reshape(h, w, 3)
    image = Image.fromarray(data)
    out = fit_gradient(image, max_stops=2)
    assert out["selected"]["model"] == "linear"
    assert out["selected"]["validation_mae"] < .005
    assert out["validation_samples"] > 100
    assert "opaque_effective_color_only" in out["warnings"]


def test_gradient_mask_and_flat_color_ambiguity():
    image = Image.new("RGB", (40, 40), "#336655")
    out = fit_gradient(image, max_stops=2, max_nfev=5)
    assert out["selected"]["model"] == "solid"
    assert "gradient_direction_not_identifiable" in out["warnings"]
    with pytest.raises(ValueError, match="samples"):
        fit_gradient(image, Image.new("L", (40, 40), 0))


def test_native_pptx_line_gradient_and_shared_paths(tmp_path):
    from build_pptx import build
    from pptx import Presentation
    r = recipe()
    r["paths"][0]["style"]["line_gradient"] = {"type": "linear", "angle_deg": 90, "stops": [
        {"position": 0, "color": "23765A", "alpha": 1}, {"position": 1, "color": "23765A", "alpha": 0}]}
    out = compile_recipe(r)
    source = tmp_path/"scene.json"
    source.write_text(json.dumps(out["scene"]), encoding="utf-8")
    target = tmp_path/"draft.pptx"
    build(source, target)
    prs = Presentation(target)
    paths = list(prs.slides[0].shapes[0].shapes)
    assert len(paths) == 8
    assert all(p._element.xpath(".//a:ln/a:gradFill") for p in paths)
    assert all(p._element.xpath(".//a:cubicBezTo") for p in paths)
    before = readback(target)
    paths[0].left += 12700
    prs.save(target)
    assert readback(target) != before


def manager_fixture(tmp_path):
    from toolbox_manager.service import Manager
    root = tmp_path/"app"
    root.mkdir()
    (root/"SKILL.md").write_text('version: "test"', encoding="utf-8")
    m = Manager(tmp_path/"data", root=root)
    project = tmp_path/"project"
    project.mkdir()
    m.authorize_project(project, [])
    return m, project


def test_managed_drafts_and_native_conflicts(tmp_path):
    from toolbox_manager.policy import PolicyDenied
    from pptx import Presentation
    m, project = manager_fixture(tmp_path)
    request = {"project": str(project), "recipe": shared()}
    with pytest.raises(PolicyDenied, match="disabled"):
        m.graphics("compile", request)
    result = m.graphics("compile", request, source="owner")
    assert Path(result["pptx"]).is_file()
    assert result["office"] == "not_run"
    assert m.graphics("validate", {"project": str(project), "version": result["version"]})["structure"] == "unchanged"
    same = m.graphics("compile", request, source="owner")
    assert same["version"] == result["version"]
    updated = shared()
    updated["edges"][0]["commands"][1][2] += 5
    new = m.graphics("regenerate", {"project": str(project), "version": result["version"],
                                    "recipe": updated}, source="owner")
    assert new["version"] != result["version"]
    prs = Presentation(result["pptx"])
    prs.slides[0].shapes[0].left += 12700
    prs.save(result["pptx"])
    with pytest.raises(ValueError, match="Manual PowerPoint"):
        m.graphics("regenerate", {"project": str(project), "version": result["version"],
                                 "recipe": updated}, source="owner")


def test_governance_and_partial_draft_protection(tmp_path):
    from toolbox_manager.policy import PolicyDenied
    m, project = manager_fixture(tmp_path)
    with pytest.raises(PolicyDenied):
        m.graphics("compile", {"project": str(tmp_path/"unapproved"), "recipe": shared()}, source="owner")
    package = m.package()
    with m.store.db() as c:
        c.execute("INSERT INTO overrides VALUES(?,?,?,?)",
                  (package["id"], "tool", "graphics.preview", "false"))
    with pytest.raises(PolicyDenied):
        m.graphics("preview", {"recipe": recipe()}, source="owner")
    out = m.graphics("compile", {"project": str(project), "recipe": shared()}, source="owner")
    Path(out["directory"], "preview.svg").write_text("changed", encoding="utf-8")
    with pytest.raises(ValueError, match="Immutable"):
        m.graphics("validate", {"project": str(project), "version": out["version"]})


def test_builtin_example_and_locked_preview(tmp_path):
    example = json.loads((ROOT/"examples/graphics/capabilities.json").read_text(encoding="utf-8"))
    assert len(compile_recipe(example)["objects"]) == 6
    m, project = manager_fixture(tmp_path)
    out = m.graphics("compile", {"project": str(project), "recipe": shared()}, source="owner")
    r = shared()
    r["instances"][0]["state"] = "locked"
    preview = m.graphics("preview", {"project": str(project), "version": out["version"],
                                     "recipe": r}, source="owner")
    assert preview["preview"].startswith("data:image/png")
    assert preview["objects"][-1] == out["objects"][-1]


def test_native_reordering_is_a_conflict(tmp_path):
    from build_pptx import build
    from pptx import Presentation
    out = compile_recipe(shared())
    path = tmp_path/"scene.json"
    path.write_text(json.dumps(out["scene"]), encoding="utf-8")
    target = tmp_path/"draft.pptx"
    build(path, target)
    before = readback(target)
    prs = Presentation(target)
    tree = prs.slides[0].shapes._spTree
    tree.append(prs.slides[0].shapes[0]._element)
    prs.save(target)
    assert readback(target) != before


def test_manifest_identity_cannot_be_relabelled(tmp_path):
    m, project = manager_fixture(tmp_path)
    out = m.graphics("compile", {"project": str(project), "recipe": shared()}, source="owner")
    path = Path(out["directory"])/"manifest.json"
    manifest = json.loads(path.read_text(encoding="utf-8"))
    manifest["parent"] = "0"*64
    path.write_text(json.dumps(manifest), encoding="utf-8")
    with pytest.raises(ValueError, match="identity"):
        m.graphics("validate", {"project": str(project), "version": out["version"]})


def test_native_omitted_path_defaults_are_equivalent(tmp_path):
    from build_pptx import build
    from pptx import Presentation
    out = compile_recipe(shared())
    path = tmp_path/"scene.json"
    path.write_text(json.dumps(out["scene"]), encoding="utf-8")
    target = tmp_path/"draft.pptx"
    build(path, target)
    before = readback(target)
    prs = Presentation(target)
    for shape in prs.slides[0].shapes:
        for path in shape._element.xpath(".//a:path"):
            if path.get("stroke") == "1":
                del path.attrib["stroke"]
            if path.get("fill") == "norm":
                del path.attrib["fill"]
    prs.save(target)
    assert readback(target) == before


def test_compiled_graphics_enter_existing_region_workflow(tmp_path):
    from workflow_scene import compile_fragment
    out = compile_recipe(recipe())
    page = {"id": "page", "mapping": {"scale_x": 1, "scale_y": 1,
                                     "offset_x": 0, "offset_y": 0}, "fragments": {}}
    region = {"id": "region", "bbox": [0, 0, 400, 300]}
    fragment = compile_fragment(page, region, {"objects": out["objects"],
                                "source_notes": "Synthetic construction test"},
                                out["scene"]["canvas"], tmp_path)
    assert fragment["objects"][0]["id"] == "page.region.curves"
    assert len(fragment["objects"][0]["children"]) == 8


def test_shared_face_hole_has_opposite_winding():
    from shapely import Polygon
    r = {"format": "graphics-recipe/1", "id": "hole", "canvas": [300, 200],
         "edges": [
             {"id": "outer", "commands": [["M", 10, 10], ["L", 250, 10], ["L", 250, 180],
                                          ["L", 10, 180], ["L", 10, 10]]},
             {"id": "inner", "commands": [["M", 50, 50], ["L", 100, 50], ["L", 100, 100],
                                          ["L", 50, 100], ["L", 50, 50]]}],
         "faces": [{"id": "face", "loops": [[{"edge": "outer"}], [{"edge": "inner"}]]}],
         "order": ["face"]}
    out = compile_recipe(r)
    rings = flatten(out["objects"][0]["commands"])
    assert len(rings) == 2
    assert Polygon(rings[0]).exterior.is_ccw != Polygon(rings[1]).exterior.is_ccw
    assert Polygon(rings[0], rings[1:]).is_valid


def test_empirical_gradient_profile_monotone_and_nonsquare_projection():
    from graphics_paint import profile, transfer
    from graphics_gradient import positions
    p = profile()
    assert p["positions"][0] == p["weights"][0] == 0
    assert p["positions"][-1] == p["weights"][-1] == 1
    assert np.all(np.diff(p["weights"]) >= 0)
    xy = np.array([[.2, .3], [.4, .8]])
    assert np.allclose(positions(xy, 400, 160, "linear", [37]),
                       positions(xy, 160, 280, "linear", [37]))
    assert transfer(.5) == pytest.approx(.5, abs=.02)


def test_gradient_excludes_palette_transparency():
    image = Image.new("P", (40, 40), 0)
    image.putpalette([51, 102, 85]*256)
    image.info["transparency"] = 0
    with pytest.raises(ValueError, match="samples"):
        fit_gradient(image)
