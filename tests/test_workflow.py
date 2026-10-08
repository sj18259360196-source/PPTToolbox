"""Phase-2 tests. Synthetic Office receipts/assertions ONLY exercise orchestration.
Actual python-pptx builds/edits are real; NO Windows/host model/visual certification.
"""
from __future__ import annotations
import copy
import json
import os
import shutil
import subprocess
import sys
from pathlib import Path
import pytest
from PIL import Image
from pptx import Presentation
from test_tools import ROOT
from common import read_json,write_json,sha256
import workflow as wf
from workflow_store import load,locked,unlock,WorkflowError,save
from workflow_scene import compile_fragment,valid_plan,input_mapping,assemble,auto_regions
from tool_registry import entries,listing,show,validate_registry,run as tool_run


def ref(tmp,size=(640,360),name='reference.png'):
    p=tmp/name;Image.new('RGB',size,'white').save(p);return p


def respond(project,payload,t=None):
    t=t or wf.next_task(project);f=Path(t['task']['response_template']);r=read_json(f)
    r['result']=payload(t) if callable(payload) else payload;write_json(f,r)
    return wf.submit(project,f)


def assertion(t):
    return {'status':'passed','note':'SYNTHETIC TEST ASSERTION; not a real Office or human review',
            'reviewer':'test-fixture','viewed_files':t['task']['must_view_files']}


def region_payload(label='TEST',with_path=False):
    objects=[{'id':'label','kind':'text','text':label,'bbox':[12,12,260,32],'style':{'font_size_pt':20}}]
    if with_path:objects.append({'id':'connection','kind':'line','points':[[20,70],[190,70]],'style':{'line':'112233','line_width_pt':2,'begin_arrow':'triangle','end_arrow':'triangle'}})
    return {'objects':objects,'source_notes':'Known synthetic source, controlled by test','relationship_ids':['connection'] if with_path else []}


def make_ready(tmp,local=False,two_regions=False):
    project=tmp/'project';wf.start(project,[ref(tmp)])
    regions=[{'id':'one','bbox':[0,0,320 if two_regions else 640,360],'role':'content','summary':'Known synthetic region','local_review':local}]
    if two_regions:regions.append({'id':'two','bbox':[320,0,640,360],'role':'content','summary':'Second region'})
    respond(project,{'regions':regions});respond(project,region_payload('ONE',local))
    if two_regions:respond(project,region_payload('TWO'))
    respond(project,assertion)
    task=wf.next_task(project);assert task['task']['kind']=='office_render'
    return project,task


def fake_office(tmp,project):
    s=load(project);candidate=project/s['run']['candidate']['file'];folder=tmp/'SYNTHETIC-OFFICE';folder.mkdir()
    rows=[]
    for page in s['pages']:
        name=f'slide-{page["index"]:03}.png';shutil.copyfile(project/page['reference'],folder/name)
        rows.append({'index':page['index'],'file':name,'sha256':sha256(folder/name),'width':page['size'][0],'height':page['size'][1]})
    path=folder/'office-render.json';write_json(path,{'status':'passed','renderer':'Microsoft PowerPoint','pptx_sha256':sha256(candidate),
        'slides':rows,'scope':'SYNTHETIC TEST FIXTURE, NOT REAL OFFICE EXECUTION'})
    return path


def bring_to_review(tmp,local=True):
    project,t=make_ready(tmp,local);receipt=fake_office(tmp,project)
    respond(project,{'receipt':str(receipt)},t)
    task=wf.next_task(project);assert task['task']['kind']=='review_full'
    return project,task


def complete_review_tasks(project):
    for _ in range(50):
        task=wf.next_task(project)
        if not task.get('task'):return task
        kind=task['task']['kind'];r=assertion(task);packet=read_json(task['task']['packet'])
        if kind=='review_full':
            r={'reviewer':r['reviewer'],'viewed_files':r['viewed_files'],'checks':{k:{'status':'passed','note':'SYNTHETIC TEST REVIEW ONLY'} for k in ['visual_full','raster_scope','text_geometry']}}
        if kind=='review_local':r['relationships']=[{'object_id':x,'note':'TEST ONLY: both endpoints declared, not independently observed'} for x in packet['context']['relationships']]
        if kind in {'editable_behavior','reproducibility'}:
            current=project/load(project)['run']['candidate']['file']
            out=project/(kind+'-test.pptx')
            if kind=='editable_behavior':
                deck=Presentation(current);shape=next(x for x in deck.slides[0].shapes if x.has_text_frame);shape.text='ACTUAL PYTHON EDIT';deck.save(out)
            else:
                from build_pptx import build
                build(project/load(project)['run']['dir']/'scene.json',out)
            r['files']=[str(out)];r['actions']=[{'kind':k,'note':'Controlled Python operation; no Office test'} for k in packet['context']['required_edit_actions']]
        respond(project,r,task)
    raise AssertionError('Unexpected workflow loop')


def test_start_copies_input_and_never_overwrites(tmp_path):
    source=ref(tmp_path);original=sha256(source);p=tmp_path/'project';wf.start(p,[source])
    assert sha256(source)==original and sha256(p/'input/slide-001.png')==original
    with pytest.raises(WorkflowError,match='new project'):wf.start(p,[source])


def test_start_missing_input_leaves_no_partial_project(tmp_path):
    with pytest.raises(Exception):wf.start(tmp_path/'project',[tmp_path/'missing.png'])
    assert not (tmp_path/'project').exists()


@pytest.mark.parametrize('ratio',['0:9','-16:9','inf:9','nan:9','16:0','bad'])
def test_invalid_ratio_is_atomic(tmp_path,ratio):
    with pytest.raises(Exception):wf.start(tmp_path/'p',[ref(tmp_path)],ratio=ratio)
    assert not (tmp_path/'p').exists()


def test_size_mismatch_requires_explicit_mapping(tmp_path):
    source=ref(tmp_path,(640,480))
    with pytest.raises(ValueError,match='aspect ratio'):wf.start(tmp_path/'p',[source],ratio='16:9')
    wf.start(tmp_path/'q',[source],ratio='16:9',mapping='explicit_stretch')
    m=load(tmp_path/'q')['pages'][0]['mapping'];assert m['scale_x']!=m['scale_y']


@pytest.mark.parametrize('regions',[
    [],[{'id':'full','bbox':[0,0,640,360],'role':'content','summary':'x'}],
    [{'id':'same','bbox':[0,0,20,20],'role':'content','summary':'x'},{'id':'SAME','bbox':[20,0,40,20],'role':'content','summary':'x'}],
    [{'id':'bad','bbox':[0,0,641,360],'role':'content','summary':'x'}],
    [{'id':'bad','bbox':[0.0,0,20,20],'role':'content','summary':'x'}],
    [{'id':'bad','bbox':[0,0,20,20],'role':'invented','summary':'x'}],
    [{'id':'bad','bbox':[0,0,20,20],'role':'content','summary':''}],
])
def test_bad_plan_preserves_active_task(tmp_path,regions):
    p=tmp_path/'p';wf.start(p,[ref(tmp_path)]);task=wf.next_task(p);before=sha256(p/'workflow/state.json')
    with pytest.raises(WorkflowError):respond(p,{'regions':regions},task)
    assert sha256(p/'workflow/state.json')==before and wf.next_task(p)['task']['id']==task['task']['id']


def test_repeated_next_is_idempotent(tmp_path):
    p=tmp_path/'p';wf.start(p,[ref(tmp_path)]);a=wf.next_task(p);before=sha256(p/'workflow/state.json');b=wf.next_task(p)
    assert a['task']==b['task'];assert sha256(p/'workflow/state.json')==before


def test_invalid_envelope_and_replay_rejected(tmp_path):
    p=tmp_path/'p';wf.start(p,[ref(tmp_path)]);t=wf.next_task(p);f=Path(t['task']['response_template']);r=read_json(f);r['token']='wrong';write_json(f,r)
    with pytest.raises(WorkflowError,match='token'):wf.submit(p,f)
    r['token']=read_json(t['task']['packet'])['token'];r['result']['regions'][0]['summary']='Known synthetic source'
    # Exercise legacy contract compatibility; new templates also scaffold scope.
    r['result'].pop('element_scope', None);write_json(f,r)
    wf.submit(p,f)
    with pytest.raises(WorkflowError,match='No active task'):wf.submit(p,f)


def test_source_change_rejected_not_silently_resigned(tmp_path):
    p=tmp_path/'p';wf.start(p,[ref(tmp_path)]);wf.next_task(p);Image.new('RGB',(640,360),'black').save(p/'input/slide-001.png')
    with pytest.raises(WorkflowError,match='Immutable input'):wf.status(p)


def test_packet_change_rejected(tmp_path):
    p=tmp_path/'p';wf.start(p,[ref(tmp_path)]);t=wf.next_task(p);Path(t['task']['packet']).write_text('{}')
    with pytest.raises(WorkflowError,match='packet'):wf.next_task(p)


def test_two_regions_get_local_coordinates_and_prefixed_ids(tmp_path):
    p,_=make_ready(tmp_path,two_regions=True);state=load(p);scene=read_json(p/state['run']['dir']/'scene.json')
    a,b=scene['slides'][0]['objects'];assert a['id']=='slide-001.one.label';assert b['id']=='slide-001.two.label'
    assert a['bbox'][0]==12 and b['bbox'][0]==332
    assert b['style']['font_size_pt']==20 and b['editability']=='text'


def test_one_fragment_not_whole_scene_in_packet(tmp_path):
    p=tmp_path/'p';wf.start(p,[ref(tmp_path)]);respond(p,{'regions':[{'id':'one','bbox':[0,0,320,360],'role':'content','summary':'Left'},{'id':'two','bbox':[320,0,640,360],'role':'content','summary':'Right'}]})
    t=wf.next_task(p);packet=read_json(t['task']['packet'])
    assert packet['context']['local_size']==[320,360] and 'slides' not in packet['context']
    with Image.open(t['task']['must_view_files'][0]) as im:assert im.size==(320,360)
    assert 'source_notes' in read_json(t['task']['response_template'])['result']


@pytest.mark.parametrize('mutate',[
    lambda r:r['objects'][0].update(kind='unimplemented'),
    lambda r:r['objects'][0].update(geometry='unknown'),
    lambda r:r['objects'][0]['style'].update(unknown_field=3),
    lambda r:r['objects'][0].update(bbox=[1,1,-3,4]),
    lambda r:r.update(relationship_ids=['missing']),
    lambda r:r['objects'][0].update(editability='image_replace'),
    lambda r:r['objects'][0].update(evidence={'status':'certain'}),
])
def test_bad_fragment_field_rejected_without_losing_plan(tmp_path,mutate):
    p=tmp_path/'p';wf.start(p,[ref(tmp_path)]);respond(p,{'regions':[{'id':'one','bbox':[0,0,640,360],'role':'content','summary':'Synthetic'}]});t=wf.next_task(p)
    payload=region_payload();mutate(payload);before=sha256(p/'workflow/state.json')
    with pytest.raises(WorkflowError):respond(p,payload,t)
    assert sha256(p/'workflow/state.json')==before
    assert load(p)['pages'][0]['plan'] is not None


def test_source_review_requires_actual_image_path(tmp_path):
    p=tmp_path/'p';wf.start(p,[ref(tmp_path)]);respond(p,{'regions':[{'id':'one','bbox':[0,0,640,360],'role':'content','summary':'x'}]});respond(p,region_payload())
    t=wf.next_task(p);r=assertion(t);r['viewed_files']=[]
    with pytest.raises(WorkflowError,match='viewed_files'):respond(p,r,t)
    assert not list((p/'runs').iterdir())


def test_waiting_office_does_not_rebuild(tmp_path):
    p,t=make_ready(tmp_path);s=load(p);pptx=p/s['run']['candidate']['file'];h=sha256(pptx);rev=s['revision']
    for _ in range(3):assert wf.next_task(p)['task']['id']==t['task']['id']
    assert sha256(pptx)==h and load(p)['revision']==rev and len(list((p/'runs').glob('run-*')))==1


def test_cli_waiting_is_exit_zero(tmp_path):
    p,_=make_ready(tmp_path)
    proc=subprocess.run([sys.executable,str(ROOT/'toolbox.py'),'next','--project',str(p)],capture_output=True,text=True,encoding='utf-8')
    assert proc.returncode==0 and json.loads(proc.stdout)['workflow_status']=='awaiting_office'


def test_scene_import_with_candidate_preserves_bytes(tmp_path):
    source=tmp_path/'source';source.mkdir();shutil.copyfile(ROOT/'examples/sample-alpha.png',source/'sample-alpha.png')
    scene=read_json(ROOT/'examples/three-page-scene.json')
    from common import walk_objects
    for slide in scene['slides']:
        for obj in walk_objects(slide['objects']):
            if obj['kind']=='image':
                obj.update(asset_role='photo',raster_content='texture',
                           raster_reason='Synthetic raster texture fixture; not a native icon or visual acceptance')
    for idx,slide in enumerate(scene['slides']):
        rp=ref(source,(960,540),f'ref-{idx}.png');slide['reference']=rp.name
    write_json(source/'scene.json',scene)
    from build_pptx import build
    candidate=source/'candidate.pptx';build(source/'scene.json',candidate);original=sha256(candidate)
    p=tmp_path/'p';wf.start(p,scene_path=source/'scene.json',candidate=candidate)
    for _ in scene['slides']:assert wf.next_task(p)['task']['kind']=='source_review';respond(p,assertion)
    t=wf.next_task(p);assert t['task']['kind']=='office_render'
    s=load(p);assert s['builder']=='external' and sha256(p/s['run']['candidate']['file'])==original and sha256(candidate)==original


def test_external_builder_does_not_implicitly_build(tmp_path,monkeypatch):
    p=tmp_path/'p';wf.start(p,[ref(tmp_path)],builder='external');respond(p,{'regions':[{'id':'one','bbox':[0,0,640,360],'role':'content','summary':'x'}]});respond(p,region_payload());respond(p,assertion)
    monkeypatch.setattr(wf,'build',lambda *a:pytest.fail('external builder must not invoke Python build'))
    t=wf.next_task(p);assert t['task']['kind']=='candidate';assert not list((p/'runs').rglob('*.pptx'))


def test_external_bad_then_good_submission_atomic(tmp_path):
    p=tmp_path/'p';wf.start(p,[ref(tmp_path)],builder='external');respond(p,{'regions':[{'id':'one','bbox':[0,0,640,360],'role':'content','summary':'x'}]});respond(p,region_payload());respond(p,assertion);t=wf.next_task(p)
    bad=tmp_path/'bad.pptx';Presentation().save(bad)
    with pytest.raises(WorkflowError):respond(p,{'file':str(bad)},t)
    assert not list((p/'runs').glob('*/external-*'))
    from build_pptx import build
    good=tmp_path/'good.pptx';build(p/load(p)['run']['dir']/'scene.json',good);respond(p,{'file':str(good)},t)
    assert wf.next_task(p)['task']['kind']=='office_render'


def test_source_approval_not_presumed_from_scene_import(tmp_path):
    data=read_json(ROOT/'examples/three-page-scene.json');shutil.copyfile(ROOT/'examples/sample-alpha.png',tmp_path/'sample-alpha.png')
    for i,s in enumerate(data['slides']):s['reference']=ref(tmp_path,(960,540),f'r{i}.png').name
    write_json(tmp_path/'scene.json',data);p=tmp_path/'p';wf.start(p,scene_path=tmp_path/'scene.json')
    t=wf.next_task(p);assert t['task']['kind']=='source_review';assert not list((p/'runs').iterdir())


def test_auto_comparison_covers_paths_and_relations(tmp_path):
    p,_=bring_to_review(tmp_path,True);s=load(p);base=p/s['run']['dir'];regions=read_json(base/'regions.json')
    ids={o for r in regions['slides'][0]['regions'] for o in r['object_ids']}
    assert 'slide-001.one.connection' in ids
    report=read_json(p/s['run']['comparison']['file']);assert report['slides'][0]['required_review']['relationship_objects']==['slide-001.one.connection']
    assert list(base.rglob('full-side-by-side.png')) and list(base.rglob('roi-001-side-by-side.png'))


def test_old_or_fake_bad_receipt_rejected(tmp_path):
    p,t=make_ready(tmp_path);receipt=fake_office(tmp_path,p);d=read_json(receipt);d['pptx_sha256']='0'*64;write_json(receipt,d)
    with pytest.raises(WorkflowError):respond(p,{'receipt':str(receipt)},t)
    assert load(p)['run'].get('render') is None


def test_local_relationship_notes_cannot_be_skipped(tmp_path):
    p,t=bring_to_review(tmp_path,True)
    respond(p,{'reviewer':'test','viewed_files':t['task']['must_view_files'],'checks':{k:{'status':'passed','note':'SYNTHETIC'} for k in ['visual_full','raster_scope','text_geometry']}},t)
    t=wf.next_task(p);assert t['task']['kind']=='review_local';r=assertion(t);r['relationships']=[]
    with pytest.raises(WorkflowError,match='relationship'):respond(p,r,t)


def test_full_checks_must_be_separate(tmp_path):
    p,t=bring_to_review(tmp_path,False);r=assertion(t)
    with pytest.raises(WorkflowError):respond(p,r,t)
    assert not load(p)['run']['reviews']


def test_edit_rebuild_files_and_gate_positive(tmp_path):
    p,_=bring_to_review(tmp_path,True);result=complete_review_tasks(p)
    assert result['workflow_status']=='ready_to_deliver'
    s=load(p);gate=read_json(p/s['run']['gate']['file']);assert gate['status']=='passed'
    review=read_json(p/s['run']['review']['file'])
    assert review['checks']['relationships']['objects'][0]['object_id']=='slide-001.one.connection'
    assert 'synthetic' in review['reviewer'] or 'test' in review['reviewer']
    delivered=wf.complete(p);assert delivered['delivery_created'];out=Path(delivered['delivery'])/'editable.pptx'
    assert sha256(out)==s['run']['candidate']['sha256']


def test_simple_page_not_forced_into_useless_local_review(tmp_path):
    p,_=bring_to_review(tmp_path,False);complete_review_tasks(p);s=load(p);r=read_json(p/s['run']['review']['file'])
    assert r['checks']['visual_local']['status']=='not_applicable' and r['checks']['relationships']['status']=='not_applicable'


def test_missing_edit_evidence_cannot_pass(tmp_path):
    p,t=bring_to_review(tmp_path,False)
    respond(p,{'reviewer':'test','viewed_files':t['task']['must_view_files'],'checks':{k:{'status':'passed','note':'SYNTHETIC'} for k in ['visual_full','raster_scope','text_geometry']}},t)
    t=wf.next_task(p);r=assertion(t);r.update(files=[],actions=[{'kind':'text','note':'Claim only'}])
    with pytest.raises(WorkflowError,match='output files'):respond(p,r,t)


def test_failed_review_stops_progress_and_can_be_reopened(tmp_path):
    p,t=bring_to_review(tmp_path,False);checks={k:{'status':'passed','note':'SYNTHETIC'} for k in ['visual_full','raster_scope','text_geometry']};checks['visual_full']['status']='needs_changes'
    respond(p,{'reviewer':'test','viewed_files':t['task']['must_view_files'],'checks':checks},t)
    assert wf.next_task(p)['workflow_status']=='awaiting_revision'
    before=load(p)['run']['candidate']['sha256'];wf.review_again(p,t['task']['id'],'Test a second review on the same unchanged file')
    t2=wf.next_task(p);assert t2['task']['kind']=='review_full' and t2['task']['id']!=t['task']['id'];assert load(p)['run']['candidate']['sha256']==before


def test_missing_comparison_after_generation_is_not_accepted(tmp_path):
    p,t=bring_to_review(tmp_path,True);s=load(p);comp=read_json(p/s['run']['comparison']['file']);page=comp['slides'][0]
    image=(p/s['run']['comparison']['file']).parent/page['full_side_by_side'];image.unlink()
    with pytest.raises(Exception):wf.next_task(p)


def test_revision_preserves_siblings_old_run_and_invalidates_task(tmp_path):
    p,t=make_ready(tmp_path,two_regions=True);s=load(p);second=copy.deepcopy(s['pages'][0]['fragments']['two']);old=s['run']['dir'];old_task=read_json(t['task']['response_template'])
    wf.revise(p,'slide-001','one','Only first region changes')
    s=load(p);assert s['pages'][0]['fragments']['two']==second and (p/old).is_dir() and s['run'] is None
    new=wf.next_task(p);assert new['task']['kind']=='region_objects';assert new['task']['target']['region_id']=='one'
    f=tmp_path/'old.json';write_json(f,old_task)
    with pytest.raises(WorkflowError,match='token'):wf.submit(p,f)


def test_reopening_source_does_not_rebuild_unchanged_candidate(tmp_path):
    p,_=make_ready(tmp_path);s=load(p);refrow=s['pages'][0]['source_review'];record=read_json(p/refrow['file']);h=s['run']['candidate']['sha256']
    wf.review_again(p,record['task_id'],'Independent source recheck');t=wf.next_task(p)
    assert t['task']['kind']=='source_review';respond(p,assertion,t);assert wf.next_task(p)['task']['kind']=='office_render'
    assert load(p)['run']['candidate']['sha256']==h


def test_replace_candidate_keeps_input_and_creates_new_run(tmp_path):
    p,_=make_ready(tmp_path);s=load(p);old=s['run']['dir'];candidate=p/s['run']['candidate']['file'];new=tmp_path/'new.pptx';shutil.copyfile(candidate,new)
    wf.replace_candidate(p,new,'Same scene new external candidate');t=wf.next_task(p)
    assert t['task']['kind']=='office_render' and load(p)['run']['dir']!=old and (p/old).is_dir()
    assert sha256(new)==load(p)['run']['candidate']['sha256']


def test_host_asset_handoff_and_true_alpha(tmp_path):
    p=tmp_path/'p';wf.start(p,[ref(tmp_path)]);respond(p,{'regions':[{'id':'one','bbox':[0,0,640,360],'role':'diagram','summary':'Complex icon'}]})
    respond(p,{'action':'request_asset','request':{'id':'icon','purpose':'Synthetic complex icon','prompt':'TEST request only, no actual model call','transparent':True,'allowed_approximation':True}})
    t=wf.next_task(p);assert t['task']['kind']=='asset_material'
    bad=ref(tmp_path,(64,64),'rgb.png')
    with pytest.raises(WorkflowError,match='transparency'):respond(p,{'file':str(bad),'tool_used':'test-fixture','provenance':'synthetic'},t)
    good=tmp_path/'alpha.png';im=Image.new('RGBA',(64,64),(0,0,0,0));im.putpixel((30,30),(255,0,0,255));im.save(good)
    respond(p,{'file':str(good),'tool_used':'Pillow synthetic fixture, not ImageGen','provenance':'Created solely by this test','model_reported':None},t)
    new=wf.next_task(p);packet=read_json(new['task']['packet']);assert new['task']['kind']=='region_objects';assert packet['context']['available_assets'][0]['asset'].startswith('assets/')


def test_cancelled_asset_request_does_not_deadlock_reopened_region(tmp_path):
    p=tmp_path/'p';wf.start(p,[ref(tmp_path)]);respond(p,{'regions':[{'id':'one','bbox':[0,0,640,360],'role':'content','summary':'Synthetic'}]})
    respond(p,{'action':'request_asset','request':{'id':'a','purpose':'icon','prompt':'x','transparent':True,'allowed_approximation':True}})
    assert wf.next_task(p)['task']['kind']=='asset_material'
    wf.revise(p,'slide-001','one','Use native path instead of missing host capability')
    assert wf.next_task(p)['task']['kind']=='region_objects'


def test_asset_import_refreshes_token_without_erasing_plan(tmp_path):
    p=tmp_path/'p';wf.start(p,[ref(tmp_path)]);t=wf.next_task(p);asset=ref(tmp_path,(32,32),'asset.png')
    wf.add_asset(p,asset,'User supplied test asset');new=wf.next_task(p);assert new['task']['id']!=t['task']['id']
    assert load(p)['asset_catalog'][0]['asset'].startswith('assets/')


def test_second_writer_blocked_and_explicit_unlock_token(tmp_path):
    p=tmp_path/'p';wf.start(p,[ref(tmp_path)])
    with locked(p):
        lock=read_json(p/'workflow/writer.lock')
        with pytest.raises(WorkflowError,match='writer lock'):wf.next_task(p)
        with pytest.raises(WorkflowError):unlock(p,lock['token'],False)
        with pytest.raises(WorkflowError):unlock(p,'WRONG',True)
    assert not (p/'workflow/writer.lock').exists()


def test_interrupted_step_needs_explicit_retry(tmp_path):
    p=tmp_path/'p';wf.start(p,[ref(tmp_path)]);s=load(p);s['operation']={'kind':'render'};save(p,s,'TEST_INTERRUPTION')
    assert wf.next_task(p)['workflow_status']=='interrupted_operation'
    wf.retry(p,'Confirmed no active renderer; test recovery');assert wf.next_task(p)['task']['kind']=='page_plan'


def test_successful_work_not_restarted_by_retry(tmp_path):
    p,_=make_ready(tmp_path)
    with pytest.raises(WorkflowError,match='No failed'):wf.retry(p,'Unnecessary retry')


def test_project_relocation_reissues_absolute_paths_without_old_root(tmp_path):
    p=tmp_path/'p';wf.start(p,[ref(tmp_path)]);t=wf.next_task(p);other=tmp_path/'moved';shutil.copytree(p,other)
    moved=wf.next_task(other);assert moved['task']['id']==t['task']['id']
    assert all(str(other) in x for x in moved['task']['must_view_files'])
    respond(other,{'regions':[{'id':'one','bbox':[0,0,640,360],'role':'content','summary':'After moving'}]},moved)
    assert wf.next_task(other)['task']['kind']=='region_objects'


def test_accepted_record_mutation_is_rejected(tmp_path):
    p,_=make_ready(tmp_path);s=load(p);(p/s['accepted'][0]['file']).write_text('{}')
    with pytest.raises(WorkflowError,match='Accepted task'):wf.status(p)


def test_frozen_input_change_is_rejected(tmp_path):
    p,_=make_ready(tmp_path);s=load(p);(p/s['run']['dir']/'scene.json').write_text('{}')
    with pytest.raises(WorkflowError,match='Frozen'):wf.status(p)


def test_imported_source_replacement_requires_identical_reference(tmp_path):
    p,_=make_ready(tmp_path);scene=assemble(load(p));scene['slides'][0]['reference']=ref(tmp_path,(640,360),'different.png').name
    Image.new('RGB',(640,360),'black').save(tmp_path/'different.png');write_json(tmp_path/'replacement.json',scene)
    with pytest.raises(ValueError,match='reference differs'):wf.replace_scene(p,tmp_path/'replacement.json','Bad source')


def test_tool_registry_complete_and_machine_readable():
    assert validate_registry()['status']=='passed'
    import jsonschema
    jsonschema.validate(read_json(ROOT/'toolbox/registry.json'),read_json(ROOT/'assets/schemas/tool-registry.schema.json'))
    assert len(entries())>=49  # v1.3 adds task-level tools; all 49 legacy entries remain


def test_registry_host_module_not_falsely_executed():
    for name in ['host.image-generation','office.native','manual.geometry']:
        result,code=tool_run(name,[])
        assert code==0 and result['command_status']=='action_required' and 'argv' not in result


def test_registry_search_supports_native_chinese_words():
    results=listing('异形');assert any(r['id']=='office.add-pptpath' for r in results)
    assert show('office.merge-pptshapes')['signature'].startswith('param(')


def test_registry_script_help_uses_real_entry_without_shell():
    result,code=tool_run('assets.crop',['--help']);assert code==0 and 'box' in result['stdout'];assert result['argv'][0]==sys.executable


@pytest.mark.parametrize('name',['missing','../escape','office;rm -rf /'])
def test_unknown_tool_does_not_execute(name):
    with pytest.raises(ValueError,match='Unknown tool'):tool_run(name,[])


def test_root_help_and_task_schema():
    proc=subprocess.run([sys.executable,str(ROOT/'toolbox.py'),'--help'],capture_output=True,text=True,encoding='utf-8')
    assert proc.returncode==0 and 'next' in proc.stdout and 'review-again' in proc.stdout
    import jsonschema
    jsonschema.Draft202012Validator.check_schema(read_json(ROOT/'assets/schemas/workflow-response.schema.json'))


def test_missing_optional_office_is_action_not_false_success():
    if os.name!='nt':
        result,code=tool_run('office.render',['-ProbeOnly','-OutputDir','not-created'])
        assert code==0 and result['command_status']=='awaiting_capability'


def test_short_skill_routes_into_real_toolbox():
    skill=(ROOT/'SKILL.md').read_text(encoding='utf-8');assert len(skill.splitlines())<=90 and 'toolbox.py' in skill
    assert 'asset-policy.md' in skill and '跨模型实测' in skill


def test_office_task_returns_fully_resolved_argv(tmp_path):
    p,t=make_ready(tmp_path)
    cmd=t['tool_commands'][0]['argv'];assert cmd[:2]==[sys.executable,str(ROOT/'toolbox.py')]
    for flag in ['-Pptx','-OutputDir','-SizesJson']:
        assert Path(cmd[cmd.index(flag)+1]).is_absolute()
    assert t['submit_argv'][-1]==t['task']['response_template']


def test_compare_failure_retry_keeps_built_candidate(tmp_path,monkeypatch):
    p,t=make_ready(tmp_path);receipt=fake_office(tmp_path,p);respond(p,{'receipt':str(receipt)},t)
    previous=load(p)['run']['candidate'].copy();real=wf.compare_deck
    def fail(*args,**kwargs):raise OSError('Synthetic comparison failure')
    monkeypatch.setattr(wf,'compare_deck',fail)
    result=wf.next_task(p);assert result['workflow_status']=='tool_failed'
    assert load(p)['run']['candidate']==previous
    monkeypatch.setattr(wf,'compare_deck',real);wf.retry(p,'Removed synthetic test failure')
    result=wf.next_task(p);assert result['task']['kind']=='review_full'
    assert load(p)['run']['candidate']==previous and len(list((p/'runs').rglob('candidate.pptx')))==1


def test_build_failure_retry_publishes_no_passed_candidate(tmp_path,monkeypatch):
    p=tmp_path/'p';wf.start(p,[ref(tmp_path)]);respond(p,{'regions':[{'id':'one','bbox':[0,0,640,360],'role':'content','summary':'test'}]});respond(p,region_payload());respond(p,assertion)
    real=wf.build
    def fail(*args,**kwargs):raise RuntimeError('Synthetic build failure')
    monkeypatch.setattr(wf,'build',fail);assert wf.next_task(p)['workflow_status']=='tool_failed'
    assert not load(p)['run'].get('candidate')
    monkeypatch.setattr(wf,'build',real);wf.retry(p,'Removed injected failure')
    assert wf.next_task(p)['task']['kind']=='office_render'


def test_valid_scene_replacement_rechecks_source(tmp_path):
    p,_=make_ready(tmp_path);state=load(p);scene_path=p/state['run']['dir']/'scene.json'
    copyroot=tmp_path/'replacement';shutil.copytree(scene_path.parent,copyroot)
    incoming=read_json(copyroot/'scene.json');incoming['slides'][0]['objects'][0]['text']='CHANGED'
    write_json(copyroot/'scene.json',incoming)
    wf.replace_scene(p,copyroot/'scene.json','Explicit content revision for test')
    assert load(p)['pages'][0]['native_first'] is True
    task=wf.next_task(p);assert task['task']['kind']=='source_review'
    assert 'CHANGED' in json.dumps(read_json(task['task']['packet'])['context'])
