"""Measured radius recipes and bounded edits; Office evidence remains separate."""
import copy
from pathlib import Path
import sys

from pptx import Presentation
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT/"scripts"))
from build_pptx import GEOMETRY, build
from common import read_json, sha256, write_json
from components import compile_component
from native_practice import CANVAS, make_cases
from patch_ops import patch_scene, normalize_scene_change
from rebuild_assistance import production_guidance
from validate_scene import ADJUSTMENT_COUNTS, validate


@pytest.mark.parametrize("geometry", list(GEOMETRY))
def test_adjustment_count_matches_actual_builder(geometry):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_shape(GEOMETRY[geometry], 0, 0, 1000000, 1000000)
    assert len(shape.adjustments) == ADJUSTMENT_COUNTS[geometry]


@pytest.mark.parametrize("box,radius,wanted", [
    ([10, 20, 240, 120], 6, .05), ([10, 20, 240, 60], 30, .5),
    ([10, 20, 70, 210], 7, .1), ([10, 20, 80, 80], 0, 0)])
def test_measured_radius_is_native_and_in_logical_units(box, radius, wanted):
    obj = compile_component({"id": "measured", "recipe": "rounded_rect",
                             "bbox": box, "params": {"radius": radius}})
    assert obj["kind"] == "shape"
    assert obj["geometry"] == "round_rect"
    assert obj["bbox"] == box
    assert obj["adjustments"] == [wanted]


@pytest.mark.parametrize("radius", [-1, 31, True, float("inf"), "6"])
def test_invalid_radius_rejected(radius):
    with pytest.raises(ValueError):
        compile_component({"id": "measured", "recipe": "rounded_rect",
                           "bbox": [10, 20, 240, 60], "params": {"radius": radius}})


@pytest.mark.parametrize("adjustments,substring", [
    ([6000], "not raw OOXML"), ([.1, .2], "at most 1"),
    ([.51], "0..0.5"), ([True], "not of type"), ([float("nan")], "finite")])
def test_unsupported_adjustments_fail_before_build(tmp_path, adjustments, substring):
    slide = make_cases()[0][0]
    slide["objects"][2]["adjustments"] = adjustments
    errors = validate({"version": "1.0", "canvas": CANVAS, "slides": [slide]}, tmp_path)
    assert any(substring in error for error in errors)


def test_adjustment_patch_changes_no_other_object_or_field(tmp_path):
    slide = make_cases()[0][0]
    scene = {"version": "1.0", "canvas": CANVAS, "slides": [slide]}
    source = tmp_path/"scene.json"
    write_json(source, scene)
    before = sha256(source)
    spec = {"base_sha256": before, "allowed_ids": ["basic-shapes/basic-1"],
            "reason": "Test native corner only", "changes": [
                {"slide": "basic-shapes", "id": "basic-1",
                 "op": "set_adjustments", "value": [.08]}]}
    patch_scene(source, spec, tmp_path/"patched")
    updated = read_json(tmp_path/"patched/scene.json")
    expected = copy.deepcopy(scene)
    expected["slides"][0]["objects"][2]["adjustments"] = [.08]
    assert updated == expected
    assert sha256(source) == before
    build(tmp_path/"patched/scene.json", tmp_path/"candidate.pptx")
    assert Presentation(tmp_path/"candidate.pptx").slides[0].shapes[2].adjustments[0] == .08
    spec["allowed_ids"] = ["basic-shapes/basic-0"]
    with pytest.raises(ValueError, match="outside permitted"):
        patch_scene(source, spec, tmp_path/"denied")
    assert not (tmp_path/"denied").exists()


def test_unit_guidance_only_in_authoring_packets():
    assert "native_units" in production_guidance("region_objects")
    assert "native_units" not in production_guidance("review_full")
    assert "6000" in production_guidance("region_objects")["native_units"]


@pytest.mark.parametrize("op,alias,value", [
    ("set_runs", "runs", [{"text": "Exact"}]), ("move", "delta", [0, -5])])
def test_patch_alias_is_exact_and_conflicts_rejected(op, alias, value):
    change = {"slide": "s", "id": "o", "op": op, alias: value}
    before = copy.deepcopy(change)
    assert normalize_scene_change(change) == {"slide": "s", "id": "o", "op": op, "value": value}
    assert change == before
    with pytest.raises(ValueError, match="Conflicting"):
        normalize_scene_change({**change, "value": value})
    with pytest.raises(ValueError, match="Change fields"):
        normalize_scene_change({**change, "arbitrary": "rejected"})
