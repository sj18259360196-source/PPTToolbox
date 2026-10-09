"""No Office claims: manifest, schema, native serialization and scope tests."""
import copy
import json
import sys
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT/"scripts"))
from common import write_json, sha256
from build_pptx import build
from validate_scene import validate
from local_edit import prepare, surgical_write
from native_capabilities import FORMAT, validate_format, search
from native_format import read, apply, effect_margin

FILL = {"mode": "linear", "angle_deg": 30, "rotate_with_shape": False,
        "stops": [{"position": 0, "color": "227755", "alpha": 1},
                  {"position": .5, "color": "22AA88", "alpha": .5},
                  {"position": 1, "color": "BBEECC", "alpha": 1}]}
SHADOW = {"enabled": True, "color": "223344", "alpha": .4, "blur_pt": 4,
          "dx_pt": 3, "dy_pt": 5}


def fixture(tmp_path, kind="shape", native=None):
    Image.new("RGB", (640, 360), "white").save(tmp_path/"ref.png")
    o = {"id": "page.body.target", "kind": kind, "bbox": [100, 120, 260, 80],
         "style": {"fill": "DDEEFF", "line": "224455"},
         "editability": "shape" if kind == "shape" else "text",
         "evidence": {"status": "measured", "note": "program fixture"}}
    o.update({"geometry": "rect"} if kind == "shape" else {"text": "Test 42 单位"})
    if native:
        o["style"]["native_format"] = native
    scene = {"version": "1.0", "canvas": {"width": 640, "height": 360, "width_pt": 960,
             "height_pt": 540, "mapping": "uniform"},
             "slides": [{"id": "page", "reference": "ref.png", "objects": [o]}]}
    write_json(tmp_path/"scene.json", scene)
    return scene


@pytest.mark.parametrize("query,wanted", [
    ("卡片需要柔和投影", "shape.shadow"), ("把文字本身渐变，不是文本框底色", "text.fill"),
    ("用三色渐变带半透明色标", "shape.fill"), ("只保留重叠部分", "boolean.intersect"),
    ("从A中挖掉B", "boolean.subtract"), ("按交叠区域拆成可单独改色的块", "boolean.fragment"),
    ("可改字的弧形标题", "text.warp"), ("图形有厚度和倒角", "shape.three_d"),
    ("文字有阴影和描边", "text.outline"), ("组合后整体移动", "group.create")])
def test_search_intent(query, wanted):
    assert wanted in [r["capability_id"] for r in search(query)]


def test_unimplemented_is_not_callable():
    assert all(not r["callable"] for r in search("SVG矢量转换"))
    assert all(r["validation"]=="code_only" for r in search("只保留重叠部分"))
    assert {"group.ungroup", "boolean.fragment", "vector.convert"} <= {
        r["capability_id"] for r in search("拆分")}


@pytest.mark.parametrize("bad", [
    {"version": "1", "shell": "bad"},
    {"version": "1", "fill": {"mode": "solid", "color": "FFFFFF", "alpha": 50}},
    {"version": "1", "shadow": {**SHADOW, "blur_pt": -1}},
    {"version": "1", "warp": "arbitrary"},
    {"version": "1", "fill": {**FILL, "stops": list(reversed(FILL["stops"]))}},
])
def test_schema_rejects(bad):
    with pytest.raises(Exception):
        validate_format(bad, "shape")


def test_schema_generated_matches_manifest():
    schema = json.loads((ROOT/"assets/schemas/scene.schema.json").read_text())
    assert schema["$defs"]["style"]["properties"]["native_format"] == FORMAT


@pytest.mark.parametrize("kind,value", [
    ("shape", {"version": "1", "fill": FILL, "shadow": SHADOW}),
    ("text", {"version": "1", "text_fill": FILL, "text_shadow": SHADOW,
              "text_outline": {"enabled": True, "color": "112233", "width_pt": 1},
              "warp": "textArchUp"}),
    ("shape", {"version": "1", "three_d": {"enabled": True, "depth_pt": 12,
              "color": "22AA88", "bevel": "circle", "bevel_width_pt": 3,
              "bevel_height_pt": 2, "material": "matte", "camera": "orthographicFront",
              "lighting": "threePt"}})
])
def test_builder_and_incremental_share_format(tmp_path, kind, value):
    scene = fixture(tmp_path, kind)
    src = tmp_path/"source.pptx"
    build(tmp_path/"scene.json", src)
    before = sha256(src)
    prs, updated, si, touched = prepare(tmp_path/"scene.json", src, "page",
        {"page.body.target"}, [{"id": "page.body.target", "op": "native.format", "format": value}])
    surgical_write(src, tmp_path/"patched.pptx", prs, si)
    write_json(tmp_path/"updated.json", updated)
    build(tmp_path/"updated.json", tmp_path/"rebuilt.pptx")
    assert read(Presentation(tmp_path/"patched.pptx").slides[0].shapes[0]) == read(
        Presentation(tmp_path/"rebuilt.pptx").slides[0].shapes[0])
    assert sha256(src) == before
    assert validate(updated, tmp_path) == []


def test_fill_preserves_shadow_and_text_edit_preserves_glyph_effect(tmp_path):
    fixture(tmp_path, "text", {"version": "1", "text_shadow": SHADOW, "text_fill": FILL})
    build(tmp_path/"scene.json", tmp_path/"source.pptx")
    p, updated, _, _ = prepare(tmp_path/"scene.json", tmp_path/"source.pptx", "page",
        {"page.body.target"}, [{"id": "page.body.target", "op": "native.format",
                               "text": "Changed 56 单位",
                               "format": {"version": "1", "text_fill": FILL}}])
    target = p.slides[0].shapes[0]
    assert target.text == "Changed 56 单位"
    assert target._element.xpath(".//a:outerShdw")
    apply(target, {"version": "1", "text_shadow": {"enabled": False}})
    assert not target._element.xpath(".//a:outerShdw")
    assert target._element.xpath(".//a:gradFill")


def test_disabled_dependencies_are_not_callable(tmp_path):
    from toolbox_manager.service import Manager
    m = Manager(tmp_path/"manager")
    m.toggle("tool", "office.render", False)
    assert not m.native_card("shape.shadow")["callable"]
    assert not m.native_card("boolean.fragment")["callable"]


def test_3d_oblique_scope_rejected():
    with pytest.raises(ValueError, match="build-only"):
        effect_margin({"three_d": {"enabled": True, "camera": "isometricTopUp"}})


@pytest.mark.parametrize("name", ["native_folded_banner", "native_overlap_regions",
                                 "native_effect_title", "native_arch_title", "native_depth_label"])
def test_named_recipe_compiles_to_native_scene(tmp_path, name):
    from components import compile_component
    scene = fixture(tmp_path)
    scene["slides"][0]["objects"] = [compile_component({
        "id": "recipe", "recipe": name, "bbox": [80, 100, 450, 150]})]
    write_json(tmp_path/"scene.json", scene)
    assert validate(scene, tmp_path) == []
    build(tmp_path/"scene.json", tmp_path/"recipe.pptx")
    assert len(Presentation(tmp_path/"recipe.pptx").slides[0].shapes) == 1


def test_operation_card_examples_validate_actual_contract():
    from native_capabilities import entries
    from toolbox_manager.contracts import validate
    for r in entries():
        if r.get("operation") != "native.format":
            continue
        args = {"project": "p", "operation_id": "sample", "base_revision": 1,
                "scene_sha256": "a"*64, "pptx_sha256": "b"*64, "slide": "slide-001",
                "region": "body", "reason": "schema example",
                "changes": [r["minimal_example"]]}
        validate("patch", args, ROOT)


def test_recipe_search_is_not_learning_activation():
    found = search("两端折回去的横幅")
    assert found[0]["recipe"] == "native_folded_banner"
    assert found[0]["validation"] == "code_only"


def test_operation_card_points_to_its_real_entry(tmp_path):
    from toolbox_manager.service import Manager
    m = Manager(tmp_path/"manager")
    recipe = m.native_card("recipe.native_folded_banner")
    atom = m.native_card("shape.fill")
    assert "rebuild_probe" in recipe["entry_note"] and "result.components" in recipe["entry_note"]
    assert "rebuild_patch" not in recipe["entry_note"]
    assert "rebuild_patch" in atom["entry_note"] and "native.format" in atom["entry_note"]
