"""Fractional logical canvases with real PPTX packages and synthetic receipts.

No test launches Office. Synthetic raster/readback evidence exercises validation,
not a claim about PowerPoint's visual output or font rendering.
"""
import copy
import json
import sys
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Pt

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from common import bbox_to_points, read_json, sha256, write_json
import text_calibration as calibration
import text_fit


def _scene():
    return {
        "version": "1.0",
        "canvas": {
            "width": 1672,
            "height": 940.5,
            "width_pt": 960,
            "height_pt": 540,
            "mapping": "uniform",
        },
        "slides": [{"id": "page", "objects": [{
            "id": "title", "kind": "text", "role": "header",
            "bbox": [40, 40, 400, 70], "text": "Measured sample",
            "style": {"font": "Arial", "font_size_pt": 20, "margin_pt": 0},
            "editability": "text",
            "evidence": {"status": "synthetic", "note": "Controlled test fixture"},
        }]}],
    }


@pytest.fixture
def environment(monkeypatch):
    value = {"available": True, "sha256": "e" * 64, "requested": ["Arial"]}
    monkeypatch.setattr(calibration, "environment", lambda fonts: copy.deepcopy(value))
    return value


def _plan(tmp_path, data=None):
    source = tmp_path / "scene.json"
    write_json(source, data if data is not None else _scene())
    plan_dir = tmp_path / "plan"
    result = calibration.plan(source, plan_dir)
    return source, plan_dir, result


def _receipt(plan_dir, office_dir, *, size=(1672, 941), physical=None):
    plan = read_json(plan_dir / "plan.json")
    office_dir.mkdir()
    pages, bounds = [], []
    for target in plan["targets"]:
        image = office_dir / f"slide-{target['index']:03}.png"
        Image.new("RGB", size, "white").save(image)
        pages.append({"index": target["index"], "file": image.name,
                      "sha256": sha256(image), "width": size[0], "height": size[1]})
        obj = target["object"]
        x, y, w, h = bbox_to_points(obj["bbox"], plan["canvas"])
        bounds.append({
            "slide": target["index"], "id": target["sample_id"], "rotation": 0,
            "box_pt": [x, y, w, h], "bound_pt": [x, y, w - 2, h - 2],
            "calibration": {
                "format": "office-text-calibration/1", "text": obj["text"],
                "auto_size": 0, "errors": [], "truncated": False,
                "runs": [{"text": obj["text"], "font": "Arial",
                          "office_font_size_pt": obj["style"]["font_size_pt"],
                          "baseline_offset_fraction": 0, "baseline_shift_pt": 0}],
                "lines": [{"index": 1, "bound_pt": [x, y, w - 2, h - 2],
                           "bound_height_pt": h - 2, "baseline_y_pt": None}],
            },
        })
    write_json(office_dir / "office-render.json", {
        "status": "passed", "renderer": "Microsoft PowerPoint",
        "scope": "SYNTHETIC TEST ONLY", "pptx_sha256": sha256(plan_dir / "specimens.pptx"),
        "slides": pages, "text_bounds": bounds,
        "canvas_pt": physical if physical is not None else [960, 540],
    })
    return office_dir


def _snapshot(*roots):
    files = []
    for root in roots:
        files.extend(root.rglob("*") if root.is_dir() else [root])
    return {str(path): (sha256(path), path.stat().st_mtime_ns)
            for path in files if path.is_file()}


def test_fractional_canvas_preserves_scene_geometry_font_and_real_pptx(tmp_path, environment):
    data = _scene()
    source = tmp_path / "scene.json"
    write_json(source, data)
    original = _snapshot(source)
    plan_dir = tmp_path / "plan"
    result = calibration.plan(source, plan_dir)
    plan = read_json(plan_dir / "plan.json")
    specimen = read_json(plan_dir / "specimens.scene.json")

    assert plan["canvas"] == specimen["canvas"] == data["canvas"]
    assert plan["export_size_px"] == result["export_size_px"] == [1672, 941]
    assert result["office_arguments"] == ["-TextCalibration", "-Width", "1672", "-Height", "941"]
    expected = copy.deepcopy(data["slides"][0]["objects"][0])
    expected["id"] = plan["targets"][0]["sample_id"]
    assert specimen["slides"][0]["objects"][0] == expected
    deck = Presentation(plan_dir / "specimens.pptx")
    assert (deck.slide_width, deck.slide_height) == (Pt(960), Pt(540))
    text = deck.slides[0].shapes[0]
    assert [text.left.pt, text.top.pt, text.width.pt, text.height.pt] == pytest.approx(
        [v * 960 / 1672 for v in expected["bbox"]], abs=1 / 12700)
    assert text.text == expected["text"]
    assert text.text_frame.paragraphs[0].runs[0].font.size.pt == 20
    assert _snapshot(source) == original
    assert result["visual_review"] == "pending" and result["automatic_changes"] is False


@pytest.mark.parametrize("height,pixel_height", [(940.49, 940), (940.5, 941), (940.51, 941)])
def test_export_size_uses_explicit_half_up_rounding(tmp_path, environment, height, pixel_height):
    data = _scene()
    data["canvas"].update(height=height, mapping="explicit_stretch")
    _, plan_dir, result = _plan(tmp_path, data)
    assert read_json(plan_dir / "plan.json")["canvas"]["height"] == height
    assert result["export_size_px"] == [1672, pixel_height]
    assert result["office_arguments"][-1] == str(pixel_height)


def test_fractional_receipt_evaluates_without_rewriting_any_input(tmp_path, environment):
    source, plan_dir, _ = _plan(tmp_path)
    office = _receipt(plan_dir, tmp_path / "synthetic-office")
    before = _snapshot(source, plan_dir, office)
    report_path = tmp_path / "measurement.json"
    result = calibration.evaluate(plan_dir, office, report_path)
    report = read_json(report_path)
    assert result["status"] == report["status"] == "advisory"
    assert report["export_size_px"] == [1672, 941]
    assert isinstance(report["canvas_validation"], dict) and report["canvas_validation"]
    assert report["targets"][0]["confidence"] == "high"
    assert report["targets"][0]["runs"][0]["effective_font_size_pt"] == 20
    assert report["visual_review"] == "pending" and report["automatic_changes"] is False
    cached = calibration.cache(report_path)
    assert isinstance(cached, tuple) and len(cached) == 3
    assert _snapshot(source, plan_dir, office) == before


@pytest.mark.parametrize("size", [(1672, 940), (1671, 941)])
def test_wrong_integer_png_size_is_rejected_even_when_receipt_hashes_match(tmp_path, environment, size):
    source, plan_dir, _ = _plan(tmp_path)
    office = _receipt(plan_dir, tmp_path / "synthetic-office", size=size)
    before = _snapshot(source, plan_dir, office)
    output = tmp_path / "rejected.json"
    with pytest.raises(ValueError, match="(?i)(export|pixel|dimension)"):
        calibration.evaluate(plan_dir, office, output)
    assert not output.exists()
    assert _snapshot(source, plan_dir, office) == before


@pytest.mark.parametrize("physical", [[961, 540], [960, 541]])
def test_fractional_raster_does_not_allow_one_point_physical_rescaling(tmp_path, environment, physical):
    _, plan_dir, _ = _plan(tmp_path)
    office = _receipt(plan_dir, tmp_path / "synthetic-office", physical=physical)
    output = tmp_path / "rejected.json"
    with pytest.raises(ValueError, match="(?i)(physical|canvas)"):
        calibration.evaluate(plan_dir, office, output)
    assert not output.exists()


@pytest.mark.parametrize("export_size", [[1672, 940], [1672.0, 941], [True, 941], [1672, 941, 1], None])
def test_export_metadata_must_be_the_exact_integer_derivation(tmp_path, environment, export_size):
    _, plan_dir, _ = _plan(tmp_path)
    office = _receipt(plan_dir, tmp_path / "synthetic-office")
    path = plan_dir / "plan.json"
    plan = read_json(path)
    plan["export_size_px"] = export_size
    write_json(path, plan)
    output = tmp_path / "rejected.json"
    with pytest.raises(ValueError):
        calibration.evaluate(plan_dir, office, output)
    assert not output.exists()


def test_plan_canvas_cannot_drift_from_hash_bound_frozen_source(tmp_path, environment):
    _, plan_dir, _ = _plan(tmp_path)
    office = _receipt(plan_dir, tmp_path / "synthetic-office")
    path = plan_dir / "plan.json"
    plan = read_json(path)
    # Both dimensions still round to the same pixels and retain uniform mapping.
    plan["canvas"].update(width=1672.1, height=1672.1 * 9 / 16)
    write_json(path, plan)
    output = tmp_path / "rejected.json"
    with pytest.raises(ValueError, match="(?i)(canvas|source)"):
        calibration.evaluate(plan_dir, office, output)
    assert not output.exists()


def test_specimen_package_canvas_is_bound_to_the_frozen_source(tmp_path, environment):
    _, plan_dir, _ = _plan(tmp_path)
    pptx = plan_dir / "specimens.pptx"
    deck = Presentation(pptx)
    deck.slide_width = Pt(961)
    deck.save(pptx)
    path = plan_dir / "plan.json"
    plan = read_json(path)
    plan["specimens_sha256"] = sha256(pptx)
    write_json(path, plan)
    office = _receipt(plan_dir, tmp_path / "synthetic-office")
    output = tmp_path / "rejected.json"
    with pytest.raises(ValueError, match="(?i)(canvas|specimen|physical)"):
        calibration.evaluate(plan_dir, office, output)
    assert not output.exists()


def test_legacy_integer_plan_without_export_metadata_remains_readable(tmp_path, environment):
    data = _scene()
    data["canvas"].update(height=941, mapping="explicit_stretch")
    source, plan_dir, _ = _plan(tmp_path, data)
    path = plan_dir / "plan.json"
    plan = read_json(path)
    del plan["export_size_px"]
    write_json(path, plan)
    office = _receipt(plan_dir, tmp_path / "synthetic-office")
    before = _snapshot(source, plan_dir, office)
    report = tmp_path / "measurement.json"
    calibration.evaluate(plan_dir, office, report)
    assert read_json(report)["export_size_px"] == [1672, 941]
    assert len(calibration.cache(report)) == 3
    assert _snapshot(source, plan_dir, office) == before


def test_legacy_fractional_plan_requires_regeneration(tmp_path, environment):
    _, plan_dir, _ = _plan(tmp_path)
    path = plan_dir / "plan.json"
    plan = read_json(path)
    del plan["export_size_px"]
    write_json(path, plan)
    office = _receipt(plan_dir, tmp_path / "synthetic-office")
    output = tmp_path / "rejected.json"
    with pytest.raises(ValueError, match="(?i)(export|legacy|fraction|regenerat|canvas)"):
        calibration.evaluate(plan_dir, office, output)
    assert not output.exists()


@pytest.mark.parametrize("change", ["none", "logical_same_pixel_bucket", "physical", "environment"])
def test_reuse_requires_exact_logical_canvas_physical_canvas_and_environment(tmp_path, environment, change):
    source, plan_dir, _ = _plan(tmp_path)
    office = _receipt(plan_dir, tmp_path / "synthetic-office")
    report = tmp_path / "measurement.json"
    calibration.evaluate(plan_dir, office, report)
    data = read_json(source)
    if change == "logical_same_pixel_bucket":
        data["canvas"].update(width=1672.1, height=1672.1 * 9 / 16)
    elif change == "physical":
        data["canvas"].update(width_pt=720, height_pt=405)
    elif change == "environment":
        environment["sha256"] = "f" * 64
    write_json(source, data)
    before = _snapshot(source, plan_dir, office, report)
    output = tmp_path / "reuse"
    result = calibration.plan(source, output, reuse=report)
    assert result["export_size_px"] == [1672, 941]
    if change == "none":
        assert result["status"] == "reuse_only" and result["reused"] == 1
        assert result["samples"] == 0 and result["office_arguments"] == []
        assert not (output / "specimens.pptx").exists()
    else:
        assert result["status"] == "awaiting_office" and result["reused"] == 0
        assert result["samples"] == 1
    assert _snapshot(source, plan_dir, office, report) == before


@pytest.mark.parametrize("height", [True, 0, 4096.5, float("nan"), float("inf")])
def test_invalid_logical_canvas_fails_before_output(tmp_path, environment, height):
    data = _scene()
    data["canvas"].update(height=height, mapping="explicit_stretch")
    source = tmp_path / "scene.json"
    # Deliberately serialize malformed NaN/Infinity input to test the boundary.
    source.write_text(json.dumps(data), encoding="utf-8")
    before = _snapshot(source)
    output = tmp_path / "invalid"
    with pytest.raises(ValueError):
        calibration.plan(source, output)
    assert not output.exists() and _snapshot(source) == before


def test_reference_pixel_fit_modes_keep_their_integer_canvas_requirement(tmp_path):
    source = tmp_path / "scene.json"
    write_json(source, _scene())
    fit_output = tmp_path / "fit"
    with pytest.raises(ValueError, match="(?i)integer"):
        text_fit.plan(source, {}, fit_output, deltas=[0])
    request = tmp_path / "request.json"
    write_json(request, {"source_sha256": sha256(source), "targets": [{"slide": "page", "id": "title"}]})
    refine_output = tmp_path / "refine"
    with pytest.raises(ValueError, match="(?i)integer"):
        text_fit.refined_plan(source, request, refine_output)
    assert not fit_output.exists() and not refine_output.exists()
