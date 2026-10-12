"""Real PPTX picture edits, no Office/UI certification."""
import hashlib
import sys
from pathlib import Path
import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT/'scripts')]
from material_edit_review import validate_image_edits


def source(tmp_path, grouped=False):
    Image.new('RGB', (80, 60), 'red').save(tmp_path/'red.png')
    Image.new('RGB', (80, 60), 'blue').save(tmp_path/'blue.png')
    deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[6])
    shapes = slide.shapes.add_group_shape().shapes if grouped else slide.shapes
    pic = shapes.add_picture(str(tmp_path/'red.png'), Inches(1), Inches(1), Inches(2), Inches(1.5))
    pic.name = 'slide1.art'
    slide.shapes.add_textbox(Inches(0), Inches(0), Inches(2), Inches(.5)).text = 'original'
    path = tmp_path/'current.pptx'; deck.save(path)
    return path


def picture(deck, grouped=False):
    return deck.slides[0].shapes[0].shapes[0] if grouped else deck.slides[0].shapes[0]


@pytest.mark.parametrize('change', ['identical', 'text', 'resize', 'crop'])
def test_unrelated_or_incomplete_changes_cannot_pass(tmp_path, change):
    current = source(tmp_path); deck = Presentation(current)
    pic = picture(deck)
    if change == 'text': deck.slides[0].shapes[1].text = 'changed only native text'
    if change == 'resize': pic.width += Inches(.4)
    if change == 'crop': pic.crop_left = .2
    edited = tmp_path/'edited.pptx'; deck.save(edited)
    with pytest.raises(ValueError, match='image_move' if change == 'resize' else 'actual picture changes'):
        validate_image_edits(current, [edited])


def test_actual_move_resize_replace_across_independent_copies(tmp_path):
    current = source(tmp_path); before = hashlib.sha256(current.read_bytes()).hexdigest()
    copies = []
    for change in ['move', 'resize', 'replace']:
        deck = Presentation(current); pic = picture(deck)
        if change == 'move': pic.left += Inches(.5)
        if change == 'resize': pic.width += Inches(.5)
        if change == 'replace':
            # Retain the stable object name; Office replacement may change ID.
            old = pic._element; name = pic.name
            old.getparent().remove(old)
            replacement = deck.slides[0].shapes.add_picture(str(tmp_path/'blue.png'), Inches(1), Inches(1), Inches(2), Inches(1.5))
            replacement.name = name
        path = tmp_path/(change+'.pptx'); deck.save(path); copies.append(path)
    result = validate_image_edits(current, copies)
    assert all(result['verified_actions'].values())
    assert len(result['verified_actions']['image_move']) == 1
    assert result['office_ui_verification'] == 'not_run'
    assert hashlib.sha256(current.read_bytes()).hexdigest() == before


def test_group_move_resize_and_embedded_replacement(tmp_path):
    current = source(tmp_path, True); deck = Presentation(current)
    group = deck.slides[0].shapes[0]
    group.left += Inches(.5); group.width += Inches(.5)
    pic = picture(deck, True)
    image = deck.part.package.get_or_add_image_part(str(tmp_path/'blue.png'))
    rid = deck.slides[0].part.relate_to(image, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/image')
    pic._element.blipFill.blip.rEmbed = rid
    edited = tmp_path/'group-edited.pptx'; deck.save(edited)
    result = validate_image_edits(current, [edited])
    assert all(result['verified_actions'].values())


def test_source_candidate_cannot_be_its_own_edit_copy(tmp_path):
    current = source(tmp_path)
    with pytest.raises(ValueError, match='separate PPTX copy'):
        validate_image_edits(current, [current])
