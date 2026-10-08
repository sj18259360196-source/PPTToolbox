"""Synthetic regression tests. Mock Office receipts never count as real Office validation."""
from __future__ import annotations
import copy, io, json, re, shutil, subprocess, sys, zipfile
from pathlib import Path
import pytest
from PIL import Image
from lxml import etree as ET
from pptx import Presentation
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'scripts'))
from common import bbox_to_points, normalize_text, read_json, resolve_asset, sha256, walk_objects, write_json
from validate_scene import validate
from build_pptx import build
from inspect_pptx import inspect,NS
from compare_images import compare
from compare_deck import compare_deck
from assets_tool import audit,crop,chroma_key,trim_alpha,matte
from init_project import initialize
from verify_delivery import verify,REQUIRED
from run_pipeline import run

@pytest.fixture
def scene(tmp_path):
    data=read_json(ROOT/'examples/three-page-scene.json')
    shutil.copyfile(ROOT/'examples/sample-alpha.png',tmp_path/'sample-alpha.png')
    p=tmp_path/'scene.json';write_json(p,data)
    return data,p

@pytest.fixture
def deck(scene,tmp_path):
    data,path=scene;out=tmp_path/'candidate.pptx';build(path,out)
    result=inspect(out,path);audit_path=tmp_path/'audit.json';write_json(audit_path,result)
    return out,audit_path,path,result

def test_valid_scene(scene):
    data,path=scene;assert validate(data,path.parent)==[]

@pytest.mark.parametrize('mutation',[
    'duplicate_id','unknown_field','zero_width','outside','bad_mapping','invalid_kind','missing_evidence',
    'bad_color','image_provenance','fullpage_reference','invalid_path','table_ragged','chart_length',
    'chart_provenance','endpoint_missing','empty_objects','negative_font','unused_style','wrong_editability'
])
def test_bad_scene_rejected(scene,mutation):
    d,p=scene;d=copy.deepcopy(d);objects=d['slides'][0]['objects']
    if mutation=='duplicate_id':objects[1]['id']=objects[0]['id']
    elif mutation=='unknown_field':objects[0]['something_unsupported']=1
    elif mutation=='zero_width':objects[0]['bbox'][2]=0
    elif mutation=='outside':objects[0]['bbox'][0]=-5
    elif mutation=='bad_mapping':d['canvas']['height_pt']=500
    elif mutation=='invalid_kind':objects[0]['kind']='bitmap_everything'
    elif mutation=='missing_evidence':del objects[0]['evidence']
    elif mutation=='bad_color':objects[0]['style']['color']='NOTHEX'
    elif mutation=='image_provenance':o=objects[-2];o['source_kind']='external';del o['provenance']
    elif mutation=='fullpage_reference':objects[-2]['asset_role']='reference_fullpage'
    elif mutation=='invalid_path':d['slides'][2]['objects'][-3]['children'][0]['commands'][1]=['C',1,2]
    elif mutation=='table_ragged':d['slides'][1]['objects'][2]['rows'][1].append('extra')
    elif mutation=='chart_length':d['slides'][1]['objects'][3]['series'][0]['values']=[1]
    elif mutation=='chart_provenance':del d['slides'][1]['objects'][3]['data_provenance']
    elif mutation=='endpoint_missing':d['slides'][2]['objects'][4]['begin']['object_id']='missing'
    elif mutation=='empty_objects':d['slides'][0]['objects']=[]
    elif mutation=='negative_font':objects[0]['style']['font_size_pt']=-1
    elif mutation=='unused_style':d['slides'][1]['objects'][3]['style']['gradient']={'angle_deg':0,'stops':[{'position':0,'color':'FFFFFF'},{'position':1,'color':'000000'}]}
    elif mutation=='wrong_editability':objects[-2]['editability']='path'
    assert validate(d,p.parent),mutation

def test_units():
    c={'width':1920,'height':1080,'width_pt':960,'height_pt':540,'mapping':'uniform'}
    assert bbox_to_points([100,200,300,400],c)==[50,100,150,200]
    c['height_pt']=500
    with pytest.raises(ValueError):bbox_to_points([1,2,3,4],c)

def test_text_normalization_is_not_rewriting():
    assert normalize_text('  H₂O\r\n±10%  ')== '  H₂O\n±10%  '

def test_asset_traversal_rejected(tmp_path):
    with pytest.raises(ValueError):resolve_asset(tmp_path,'../outside.png')

def test_alpha_not_assumed(tmp_path):
    rgb=tmp_path/'rgb.png';Image.new('RGB',(3,3),'white').save(rgb)
    assert audit(rgb)['has_transparent_pixels'] is False
    alpha=tmp_path/'alpha.png';Image.new('RGBA',(3,3),(0,0,0,255)).save(alpha)
    assert audit(alpha)['has_transparent_pixels'] is False
    Image.new('RGBA',(3,3),(0,0,0,0)).save(alpha)
    assert audit(alpha)['visible_subject_exists'] is False

def test_palette_transparency(tmp_path):
    p=tmp_path/'pal.png';im=Image.new('P',(3,3));im.info['transparency']=0;im.save(p)
    assert audit(p)['has_transparent_pixels'] is True

@pytest.mark.parametrize('box',[[-1,0,2,2],[0,0,5,2],[0,0,0,2],[0,0,2,5]])
def test_crop_outside_rejected(tmp_path,box):
    source=tmp_path/'source.png';Image.new('RGB',(4,4)).save(source)
    with pytest.raises(ValueError):crop(source,tmp_path/'out.png',box)

def test_crop_does_not_overwrite_source(tmp_path):
    p=tmp_path/'source.png';Image.new('RGB',(4,4)).save(p);old=sha256(p)
    with pytest.raises(ValueError):crop(p,p,[0,0,2,2])
    assert sha256(p)==old

def test_crop_valid(tmp_path):
    p=tmp_path/'source.png';Image.new('RGB',(4,4)).save(p)
    assert crop(p,tmp_path/'out.png',[1,1,3,4])['size']==[2,3]


def test_chroma_trim_and_matte(tmp_path):
    src=tmp_path/'solid.png';im=Image.new('RGB',(10,10),(255,0,255))
    for x in range(3,7):
        for y in range(2,8):im.putpixel((x,y),(0,120,255))
    im.save(src)
    keyed=tmp_path/'keyed.png';r=chroma_key(src,keyed,'FF00FF',tolerance=10,feather=0)
    assert r['has_transparent_pixels'] is True
    trimmed=tmp_path/'trimmed.png';r=trim_alpha(keyed,trimmed,padding=1)
    assert r['size']==[6,8]
    prev=tmp_path/'preview.png';matte(trimmed,prev,'DCEBFA')
    assert Image.open(prev).mode=='RGB'

def test_native_deck_roundtrip(deck):
    pptx,_,_,report=deck
    prs=Presentation(pptx);assert len(prs.slides)==3
    assert report['checks']['package']['status']=='passed'
    assert report['checks']['native_content']['status']=='passed'
    assert report['checks']['chart_cache']['status']=='passed'
    assert sum(s['tables'] for s in report['slides'])==1
    assert sum(s['charts'] for s in report['slides'])==1
    assert report['checks']['geometry']['status']=='blocked' # group transforms are honestly deferred

def test_two_arrowheads_and_attachment(deck):
    pptx,*_=deck
    with zipfile.ZipFile(pptx) as z:
        r=ET.fromstring(z.read('ppt/slides/slide3.xml'))
        con=r.find('.//p:cxnSp',NS)
        assert con.find('.//a:headEnd',NS).get('type')=='triangle'
        assert con.find('.//a:tailEnd',NS).get('type')=='triangle'
        assert con.find('.//a:stCxn',NS) is not None
        assert con.find('.//a:endCxn',NS) is not None

def test_gradient_order_and_local_path(deck):
    pptx,*_=deck
    with zipfile.ZipFile(pptx) as z:
        r=ET.fromstring(z.read('ppt/slides/slide1.xml'))
        sp=r.xpath('.//p:sp[p:nvSpPr/p:cNvPr[@name="p01.gradient"]]/p:spPr',namespaces=NS)[0]
        tags=[ET.QName(x).localname for x in sp]
        assert tags.index('prstGeom')<tags.index('gradFill')<tags.index('ln')
        r=ET.fromstring(z.read('ppt/slides/slide3.xml'))
        sp=r.xpath('.//p:sp[p:nvSpPr/p:cNvPr[@name="p03.half"]]',namespaces=NS)[0]
        assert sp.find('.//a:cubicBezTo',NS) is not None
        assert sp.find('.//a:close',NS) is not None
        assert int(sp.find('.//a:xfrm/a:ext',NS).get('cx'))<960*12700

def mutate_zip(src,dest,part,edit):
    with zipfile.ZipFile(src) as zin,zipfile.ZipFile(dest,'w',zipfile.ZIP_DEFLATED) as zout:
        for name in zin.namelist():
            blob=zin.read(name)
            if name==part:blob=edit(blob)
            zout.writestr(name,blob)

def test_changed_text_fails_content(deck,tmp_path):
    src,_,scene,_=deck;dest=tmp_path/'changed.pptx'
    mutate_zip(src,dest,'ppt/slides/slide1.xml',lambda b:b.replace('普通信息页'.encode(), '被修改标题'.encode()))
    assert inspect(dest,scene)['checks']['native_content']['status']=='failed'

def test_cache_mismatch_detected(deck,tmp_path):
    src,*_=deck;dest=tmp_path/'cache.pptx'
    def edit(b):
        r=ET.fromstring(b);r.find('.//c:numCache/c:pt/c:v',NS).text='999';return ET.tostring(r)
    mutate_zip(src,dest,'ppt/charts/chart1.xml',edit)
    assert inspect(dest)['checks']['chart_cache']['status']=='failed'

def test_missing_relationship_target_detected(deck,tmp_path):
    src,*_=deck;dest=tmp_path/'badrel.pptx'
    mutate_zip(src,dest,'ppt/slides/_rels/slide1.xml.rels',lambda b:b.replace(b'image1.png',b'not-there.png'))
    assert inspect(dest)['checks']['package']['status']=='failed'

def test_overwrite_guard(deck):
    pptx,_,scene,_=deck
    with pytest.raises(FileExistsError):build(scene,pptx)

@pytest.mark.parametrize('kind',['column','xy'])
def test_chart_types(scene,tmp_path,kind):
    d,path=scene;o=d['slides'][1]['objects'][3];o['chart_type']=kind
    if kind=='xy':
        o.pop('categories')
        for spec in o['series']:spec['points']=[[i,v] for i,v in enumerate(spec.pop('values'))]
    write_json(path,d);ppt=tmp_path/'test.pptx';build(path,ppt)
    assert inspect(ppt,path)['checks']['chart_cache']['status']=='passed'

def test_compare_identity_and_local(tmp_path):
    a=tmp_path/'a.png';b=tmp_path/'b.png';Image.new('RGB',(10,10),'white').save(a);shutil.copyfile(a,b)
    r=compare(a,b,tmp_path/'cmp',[{'id':'part','bbox':[1,1,5,5]}]);assert r['regions'][0]['mae_rgb']==0
    assert len(r['regions'])==2;assert (tmp_path/'cmp/part-side-by-side.png').is_file()

def test_compare_size_not_silent(tmp_path):
    a=tmp_path/'a.png';b=tmp_path/'b.png';Image.new('RGB',(10,10)).save(a);Image.new('RGB',(5,5)).save(b)
    with pytest.raises(ValueError):compare(a,b,tmp_path/'cmp')
    r=compare(a,b,tmp_path/'cmp',resize_render=True);assert r['resized_render'] is True

def test_invalid_region(tmp_path):
    a=tmp_path/'a.png';Image.new('RGB',(10,10)).save(a)
    with pytest.raises(ValueError):compare(a,a,tmp_path/'cmp',[{'id':'x','bbox':[0,0,11,3]}])


def test_compare_deck_receipt(tmp_path):
    ref=tmp_path/'ref.png';Image.new('RGB',(10,10),'white').save(ref)
    scene=tmp_path/'scene.json';write_json(scene,{'slides':[{'id':'slide-001','reference':'ref.png'}]})
    render_dir=tmp_path/'office';render_dir.mkdir();rend=render_dir/'slide-001.png';Image.new('RGB',(10,10),'white').save(rend)
    receipt=render_dir/'office-render.json';write_json(receipt,{'status':'passed','renderer':'Microsoft PowerPoint','probe_only':False,'pptx_sha256':'a'*64,'slides':[{'index':1,'file':rend.name,'sha256':sha256(rend),'width':10,'height':10}]})
    candidate=tmp_path/'comparison-only.pptx';candidate.write_bytes(b'SYNTHETIC FILE ID ONLY')
    rc=read_json(receipt);rc['pptx_sha256']=sha256(candidate);write_json(receipt,rc)
    out=tmp_path/'deckcmp';r=compare_deck(scene,receipt,out,pptx_path=candidate)
    assert r['status']=='passed';assert (out/'slide-001/full-side-by-side.png').is_file();assert (out/'deck-comparison.json').is_file()

def test_init_requires_real_analysis(tmp_path):
    a=tmp_path/'r.png';Image.new('RGB',(1671,941)).save(a)
    p=tmp_path/'new';d=initialize(p,[a],'16:9')
    assert read_json(p/'input/references.json')[0]['size']==[1671,941]
    assert validate(d,p)
    with pytest.raises(FileExistsError):initialize(p,[a])

def test_empty_review_cannot_pass(deck):
    p,a,_,_=deck;r=verify(p,a,None,None)
    assert r['status']=='blocked';assert r['pending']

def synthetic_receipts(deck,tmp_path):
    # Synthetic receipts exercise the code only; NOT real Office or visual acceptance.
    from evidence_contract import required_review, logical_bounds
    import math
    p,a,scene_path,_=deck;scene=read_json(scene_path);images=[];region_slides=[]
    for i,slide in enumerate(scene['slides'],1):
        ref=tmp_path/f'ref-{i}.png';Image.new('RGB',(960,540),(240,i*30,230)).save(ref)
        slide['reference']=ref.name
        f=tmp_path/f'fake-{i}.png';Image.new('RGB',(960,540),(i*30,200,230)).save(f)
        images.append({'index':i,'file':f.name,'sha256':sha256(f),'width':960,'height':540})
        by_id={o['id']:o for o in walk_objects(slide['objects'])};rr=[]
        for j,oid in enumerate(required_review(slide)['local_objects'],1):
            l,t,r,b=logical_bounds(by_id[oid])
            rr.append({'id':f'object-{j}','bbox':[max(0,math.floor(l)-3),max(0,math.floor(t)-3),min(960,math.ceil(r)+3),min(540,math.ceil(b)+3)],'object_ids':[oid]})
        region_slides.append({'index':i,'regions':rr})
    write_json(scene_path,scene);write_json(a,inspect(p,scene_path))
    render=tmp_path/'fake-office.json';write_json(render,{'status':'passed','renderer':'Microsoft PowerPoint','pptx_sha256':sha256(p),'slides':images,'probe_only':False})
    region_path=tmp_path/'fake-regions.json';write_json(region_path,{'slides':region_slides})
    compdir=tmp_path/'fake-comparisons';report=compare_deck(scene_path,render,compdir,region_path,pptx_path=p)
    comparison=compdir/'deck-comparison.json'
    ev=tmp_path/'fake-review-note.txt';ev.write_text('SYNTHETIC TEST ONLY; NO OFFICE OR VISUAL REVIEW',encoding='utf-8')
    checks={n:{'status':'passed','note':'SYNTHETIC TEST FIXTURE ONLY','slides':[1,2,3],
        'evidence':[{'file':ev.name,'sha256':sha256(ev)}]} for n in REQUIRED}
    full=[];local=[];region_reviews=[];rel_reviews=[]
    for page in report['slides']:
        full_file=compdir/page['full_side_by_side'];full.append({'file':full_file.relative_to(tmp_path).as_posix(),'sha256':sha256(full_file)})
        for region in page['local_regions']:
            f=compdir/f"slide-{page['index']:03}"/f"{region['id']}-side-by-side.png"
            local.append({'file':f.relative_to(tmp_path).as_posix(),'sha256':sha256(f)})
            region_reviews.append({'index':page['index'],'id':region['id'],'object_ids':region['object_ids'],'note':'Synthetic observation, not actual review'})
        for oid in page['required_review']['relationship_objects']:
            rel_reviews.append({'index':page['index'],'object_id':oid,'note':'Synthetic direction assertion only'})
    checks['visual_full']['evidence']=full
    checks['visual_local'].update(evidence=local,regions=region_reviews)
    checks['relationships'].update(evidence=local,objects=rel_reviews)
    review=tmp_path/'fake-review.json';write_json(review,{'pptx_sha256':sha256(p),'scene_sha256':sha256(scene_path),
        'comparison_sha256':sha256(comparison),'reviewer':'SYNTHETIC_TEST','checks':checks})
    return render,review,comparison

def test_gate_fixture_validates_identity_only(deck,tmp_path):
    render,review,comparison=synthetic_receipts(deck,tmp_path);p,a,_,_=deck
    assert verify(p,a,render,review,comparison,deck[2])['status']=='passed'

@pytest.mark.parametrize('tamper',['hash','png','missing_page','probe','review_hash','review_evidence','review_coverage','comparison_hash','comparison_file'])
def test_gate_rejects_bad_evidence(deck,tmp_path,tamper):
    render,review,comparison=synthetic_receipts(deck,tmp_path);p,a,_,_=deck
    r=read_json(render);v=read_json(review);c=read_json(comparison)
    if tamper=='hash':r['pptx_sha256']='0'*64
    if tamper=='png':r['slides'][0]['sha256']='0'*64
    if tamper=='missing_page':r['slides'].pop()
    if tamper=='probe':r['probe_only']=True
    if tamper=='review_hash':v['pptx_sha256']='0'*64
    if tamper=='review_evidence':v['checks']['content_source']['evidence']=[]
    if tamper=='review_coverage':v['checks']['visual_full']['slides']=[1]
    if tamper=='comparison_hash':c['pptx_sha256']='0'*64
    if tamper=='comparison_file':c['slides'][0]['full_side_by_side']='missing.png'
    write_json(render,r);write_json(review,v);write_json(comparison,c)
    assert verify(p,a,render,review,comparison,deck[2])['status']!='passed'

def test_pipeline_leaves_review_not_run(scene,tmp_path):
    _,path=scene;out=tmp_path/'run'
    assert run(path,out,office=False)==3
    review=read_json(out/'review.json');assert all(c['status']=='not_run' for c in review['checks'].values())

def test_skill_structure_and_links():
    import yaml,jsonschema
    text=(ROOT/'SKILL.md').read_text(encoding='utf-8');fm=yaml.safe_load(text.split('---',2)[1])
    assert fm['name']==read_json(ROOT/'MANIFEST.json')['name']  # Portable app folder is intentionally relocatable.
    assert re.fullmatch('[a-z0-9]+(?:-[a-z0-9]+)*',fm['name'])
    assert 0<len(fm['description'])<=1024;assert len(text.splitlines())<500
    for path in re.findall(r'\]\(([^)]+)\)',text):assert (ROOT/path).is_file(),path
    jsonschema.Draft202012Validator.check_schema(read_json(ROOT/'assets/schemas/scene.schema.json'))
    assert len(read_json(ROOT/'references/experience-ledger.json'))==22
    assert len({x['sha256'] for x in read_json(ROOT/'references/experience-ledger.json')})==22

def test_no_font_or_credentials_bundled():
    from package_skill import file_list
    assert not [p for p in file_list(ROOT) if p.suffix.lower() in {'.ttf','.otf','.woff','.woff2','.ttc','.key','.pem'}]


def test_soft_linebreak_is_preserved():
    from inspect_pptx import text_body
    body=ET.fromstring(('<p:txBody xmlns:p="'+NS['p']+'" xmlns:a="'+NS['a']+'"><a:p><a:r><a:t>甲</a:t></a:r><a:br/><a:r><a:t>乙</a:t></a:r></a:p></p:txBody>').encode())
    assert text_body(body)=='甲\n乙'


def test_grouped_table_count_is_not_duplicated(deck):
    # Synthetic XML audit test, not support for building grouped tables.
    from inspect_pptx import flatten
    p,*_=deck
    with zipfile.ZipFile(p) as z:
        tree=ET.fromstring(z.read('ppt/slides/slide2.xml')).find('p:cSld/p:spTree',NS)
        graphic=tree.find('p:graphicFrame',NS)
        tree.remove(graphic)
        group=ET.SubElement(tree,'{'+NS['p']+'}grpSp');group.append(graphic)
        items=list(flatten(tree))
        assert sum(bool(x['table_rows']) for x in items)==1


def test_grouped_data_objects_require_adapter(scene):
    data,path=scene;table=data['slides'][1]['objects'].pop(2)
    data['slides'][1]['objects'].append({'id':'p02.table.group','kind':'group','role':'group','editability':'group',
       'evidence':{'status':'inferred','note':'synthetic regression'},'children':[table]})
    assert any('table/chart at slide root only' in x for x in validate(data,path.parent))


def test_real_library_edit_save_reopen(deck,tmp_path):
    # Actual python-pptx edits, not a mock. This is not an Office editing test.
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    src,*_=deck;p=Presentation(src)
    target=next(s for s in p.slides[0].shapes if s.has_text_frame and s.text.startswith('普通信息页'))
    target.text='编辑往返测试';target_name=target.name
    card=next(s for s in p.slides[0].shapes if s.name=='p01.gradient')
    card.fill.solid();card.fill.fore_color.rgb=RGBColor(255,0,0)
    table=next(s.table for s in p.slides[1].shapes if s.has_table);table.cell(1,1).text='99.9'
    chart=next(s.chart for s in p.slides[1].shapes if s.has_chart)
    d=CategoryChartData();d.categories=[0,1,2,3];d.add_series('Changed',[2,4,6,8]);chart.replace_data(d)
    out=tmp_path/'edited.pptx';p.save(out);q=Presentation(out)
    assert next(s for s in q.slides[0].shapes if s.name==target_name).text=='编辑往返测试'
    assert next(s for s in q.slides[0].shapes if s.name=='p01.gradient').fill.fore_color.rgb==RGBColor(255,0,0)
    assert next(s.table for s in q.slides[1].shapes if s.has_table).cell(1,1).text=='99.9'
    c=next(s.chart for s in q.slides[1].shapes if s.has_chart)
    assert list(c.series[0].values)==[2,4,6,8]
    assert inspect(out)['checks']['chart_cache']['status']=='passed'


@pytest.mark.parametrize('command',[[[],1,2],['M','not-a-number',2],['UNKNOWN',1,2]])
def test_malformed_path_produces_diagnostics_not_exception(scene,command):
    d,p=scene;d['slides'][2]['objects'][-3]['children'][0]['commands'][1]=command
    assert validate(d,p.parent)
